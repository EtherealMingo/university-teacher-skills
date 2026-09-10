#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""实验进度跟踪表「状态色标」后处理 —— 必须跑在 office-layer xlsx_kit 之后

流程（唯一通道，不许绕过）：

    PY="vendor/office-layer/.venv/bin/python"
    $PY vendor/office-layer/scripts/xlsx_kit.py build --spec 进度spec.json --out 进度跟踪表.xlsx
    $PY skills/lab-meeting/scripts/status_colorize.py \
        --in 进度跟踪表.xlsx --sheet 进度跟踪 --status-col 4

色标口径（全组统一，不许各自解释）：
    🟢 正常 / 已完成   绿底 C6EFCE
    🟡 滞后            黄底 FFEB9C
    🔴 阻塞            红底 FFC7CE
    ⚪ 未开始          灰底 F2F2F2

为什么不用 xlsx_kit build 直接上色：build 的 spec 只支持
name/header/rows/freeze/widths/chart，逐格填充不支持。
因此本脚本用 office-layer 的 venv 与 openpyxl 做一层薄后处理，
仍属 office-layer 通道，不手搓 OOXML。
"""

from __future__ import annotations

import argparse
import json
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit("缺少 openpyxl：请用 vendor/office-layer/.venv/bin/python 运行本脚本")

# 状态 → (底色, 字色)
STATUS_MAP = {
    "正常": ("C6EFCE", "006100"),
    "已完成": ("C6EFCE", "006100"),
    "滞后": ("FFEB9C", "9C6500"),
    "阻塞": ("FFC7CE", "9C0006"),
    "未开始": ("F2F2F2", "808080"),
}
CN = "微软雅黑"


def normalize(v) -> str:
    return str(v or "").strip().replace("🟢", "").replace("🟡", "").replace("🔴", "").replace("⚪", "").strip()


def main() -> int:
    ap = argparse.ArgumentParser(prog="status_colorize", description="进度表状态色标")
    ap.add_argument("--in", dest="src", required=True, help="xlsx_kit build 产出的进度表")
    ap.add_argument("--out", help="输出路径，缺省原地覆盖")
    ap.add_argument("--sheet", help="工作表名，缺省第一张")
    ap.add_argument("--status-col", type=int, required=True, help="状态列序号（从 1 起）")
    ap.add_argument("--legend-sheet", default="色标说明", help="写图例的工作表名，传 '' 则不写")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.src)
    if args.sheet and args.sheet not in wb.sheetnames:
        sys.exit(f"找不到工作表 {args.sheet!r}，现有：{wb.sheetnames}")
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]

    tally: dict[str, int] = {}
    unknown: list[str] = []
    applied = 0
    col = args.status_col

    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=col)
        key = normalize(cell.value)
        if not key:
            continue
        if key not in STATUS_MAP:
            unknown.append(f"{cell.coordinate}={key!r}")
            continue
        bg, fg = STATUS_MAP[key]
        fill = PatternFill("solid", fgColor=bg)
        font = Font(name=CN, size=10.5, bold=True, color=fg)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        tally[key] = tally.get(key, 0) + 1
        applied += 1

    if args.legend_sheet:
        if args.legend_sheet in wb.sheetnames:
            del wb[args.legend_sheet]
        lg = wb.create_sheet(args.legend_sheet)
        lg.append(["状态", "含义", "导师动作"])
        for c in range(1, 4):
            h = lg.cell(row=1, column=c)
            h.fill = PatternFill("solid", fgColor="1F2D3D")
            h.font = Font(name=CN, size=11, bold=True, color="FFFFFF")
            h.alignment = Alignment(horizontal="center", vertical="center")
        legend = [
            ("🟢 正常", "按计划推进", "常规过会"),
            ("🟡 滞后", "落后于阶段目标", "会上要学生给出追赶时间点"),
            ("🔴 阻塞", "卡在外部条件或方法瓶颈", "导师介入，明确破局责任人"),
            ("⚪ 未开始", "尚未启动该阶段", "确认启动条件是否具备"),
        ]
        for name, meaning, action in legend:
            lg.append([name, meaning, action])
            row = lg.max_row
            key = normalize(name)
            if key in STATUS_MAP:
                bg, fg = STATUS_MAP[key]
                lg.cell(row=row, column=1).fill = PatternFill("solid", fgColor=bg)
                lg.cell(row=row, column=1).font = Font(name=CN, bold=True, color=fg)
        lg.column_dimensions["A"].width = 12
        lg.column_dimensions["B"].width = 30
        lg.column_dimensions["C"].width = 34

    out = args.out or args.src
    wb.save(out)

    print(json.dumps({
        "out": out,
        "sheet": ws.title,
        "上色单元格": applied,
        "状态分布": tally,
        "未识别状态": unknown,
        "需关注": tally.get("滞后", 0) + tally.get("阻塞", 0),
        "提醒": "色标只是提示位，是否算滞后/阻塞属判断性内容，请教师确认",
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
