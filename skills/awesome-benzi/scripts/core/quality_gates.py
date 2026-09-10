#!/usr/bin/env python3
"""Run the single combined content and DOCX gate for awesome-benzi."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import statistics
import sys
import zipfile
from pathlib import Path
from typing import Any, Iterable

from .errors import ensure_utf8_stdout
from .ledgers import validate_reasoning_protocol
from .style_profile import check_style

REQUIRED_REFERENCES = {
    "INPUT_MODEL.md",
    "WRITING_MODEL.md",
    "EVIDENCE_MODEL.md",
    "LANGUAGE_MODEL.md",
    "MARKDOWN_DOCX_CONTRACT.md",
    "QUALITY_GATES.md",
}
REQUIRED_WORKFLOWS = {
    "social-science-grant",
    "natural-science-grant",
    "innovation-competition",
    "student-innovation",
    "challenge-cup",
}
FORBIDDEN_RUNTIME_NARRATIVE = re.compile(
    r"github\.com|awesome-ai-research-writing|基于.{0,20}(?:项目|Prompt)|改编自|"
    r"Gemini\s*\d|Claude\s*\d|DeepSeek\s*R1|许可.{0,10}开源协议",
    re.I,
)

HARD_PATTERNS = {
    "unsupported-inflation": re.compile(r"大力推动|全面赋能|意义重大|颠覆性|范式转移|不可磨灭|令人惊叹|毋庸置疑|开启新篇章"),
    "chatbot-residue": re.compile(r"希望这对你有帮助|如果你愿意|如需我继续|下面我将|以下是为你|作为(?:一个)?AI"),
}
OVERCLAIM_RE = re.compile(r"填补(?:了)?(?:国内|国际)?(?:领域)?空白|形成(?:了)?闭环|首次(?:提出|发现|实现)|首创|开创性")


# Colons and dashes are normal Chinese academic punctuation ("研究对象：…",
# "2019—2023年"); banning every occurrence made ordinary prose undeliverable.
# They are checked per paragraph as over-use instead.
COLON_RE = re.compile(r"[:：]")
DASH_RE = re.compile(r"[—–]")
# Sequential connectives ("首先…其次…最后…") are legitimate enumeration in
# Chinese proposals; only heavy repetition is flagged, and never at blocked level.
MECHANICAL_SEQUENCE_RE = re.compile(r"(?:^|[。；;\n]\s*|，\s*)(第一|第二|首先|其次|再次|最后)(?:[，,、。:：]|步|点|个|方面)")

COLLOQUIAL = re.compile(r"我们觉得|大家都知道|说白了|搞清楚|弄明白|效果变好|很厉害|非常牛|一大堆")
EMPTY_META = re.compile(r"值得注意的是|需要指出的是|必须看到的是|综上所述|总而言之|展望未来|迈向新的高度")
VAGUE_ENDING = re.compile(r"具有(?:重要|重大|深远)?(?:理论|现实|实践|学术|社会)?意义|具有重要价值|提供有力支撑|奠定坚实基础")
PLACEHOLDER = re.compile(r"【待补：[^】]{1,100}】")
NUMERIC_CITATION = re.compile(r"\[(?:\d{1,3})(?:[-–,，]\d{1,3})*\]")
AUTHOR_YEAR_CITATION = re.compile(r"[（(][^（）()\n]{1,40}[，,]\s*(?:19|20)\d{2}[a-z]?[）)]", re.I)
CITATION_HINT_RE = re.compile(
    r"(?:已有研究|现有研究|相关研究|前人研究|国外研究|国内研究)(?:表明|发现|指出|显示|认为)|"
    r"研究(?:表明|发现|指出|显示|证实)|文献(?:显示|表明|指出|认为|报道)|"
    r"学者(?:指出|认为|发现|提出)|政策(?:规定|要求|显示|指出)|"
    r"统计(?:数据|结果)?(?:显示|表明)|数据(?:显示|表明)|调查(?:显示|发现|表明)|"
    r"报告(?:显示|指出)|据(?:相关|公开|最新)?(?:统计|调查|报告|研究)"
)
EVIDENCE_MARKER = re.compile(r"\d|文献|调查|访谈|案例|数据|政策|统计|研究发现|已有研究|样本|实验|测量|观察|比较|机制|边界|范围|对象|方法|材料|理论")
METHOD_NAME = re.compile(r"(?:采用|运用|使用).{0,16}(?:法|模型|分析|编码|检验|调查|访谈|实验)")
METHOD_OPERATION = re.compile(r"数据|材料|样本|访谈|文献|案例|问卷|实验|分析|编码|检验|步骤|收集|来源|抽样|测量")
SYMMETRIC_PATTERN = re.compile(r"一是.{2,80}二是|既要.{2,80}又要|不仅.{2,80}而且|从.{1,30}到.{1,30}从.{1,30}到")
GENERIC_INTERNAL_HEADING = re.compile(r"^(?:[（(][一二三四五六七八九十百]+[）)]|\d+[.、])?\s*(?:问题分析|主要内容|相关情况|基本情况|具体分析|总体分析|综合分析|其他内容|若干问题)$")
META_INSTRUCTION_RESIDUE = re.compile(r"(?:按|根据)(?:用户|申请人)(?:要求|需求)|用户(?:要求|希望|提出)|我(?:想|希望|要求)|帮我|不要(?:使用|采用|写|增加|出现)|只(?:使用|采用|做|保留)|需要(?:把|将).{0,12}(?:改为|写成|加入|删除)")
PROMPT_DIRECTIVE_CUE = re.compile(r"不要|不能|避免|需要|应当|必须|只(?:用|做|保留|写)|采用|使用|增加|补充|删除|修改|强调|突出|围绕|聚焦|保持")
ALLOWED_PARAGRAPH_STYLES = {"Normal", "Instruction", "Table Text", "List Bullet", "List Number"}
FORBIDDEN_OFFICE_HEADING_STYLES = {"Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5", "Heading 6", "Heading 7", "Heading 8", "Heading 9"}
DEFAULT_STYLE_PROFILE = {
    "title_font": "黑体", "title_size": 22.0,
    "heading1_font": "黑体", "heading1_size": 16.0,
    "heading2_font": "仿宋_GB2312", "heading2_size": 12.0,
    "heading3_font": "仿宋_GB2312", "heading3_size": 12.0,
    "body_font": "仿宋_GB2312", "body_size": 12.0,
    "table_font": "宋体", "table_size": 10.5,
    "line_spacing": 1.5, "first_line_chars": 2.0,
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def verify_social_template(path: Path | None = None) -> dict[str, Any]:
    skill_root = Path(__file__).resolve().parents[2]
    references = skill_root / "references"
    workflows = skill_root / "workflows"
    missing_references = sorted(name for name in REQUIRED_REFERENCES if not (references / name).is_file())
    missing_workflows = sorted(name for name in REQUIRED_WORKFLOWS if not (workflows / name / "WORKFLOW.md").is_file())
    runtime_files = [
        candidate
        for candidate in skill_root.rglob("*")
        if candidate.is_file() and candidate.suffix.lower() in {".md", ".yaml", ".py"}
    ]
    forbidden_hits: list[dict[str, str]] = []
    rule_owners: dict[str, str] = {}
    duplicate_rules: list[str] = []
    rule_pattern = re.compile(r"`((?:IN|EV|WR|MD|QG)-[A-Z0-9-]+)`")
    for candidate in runtime_files:
        text = candidate.read_text(encoding="utf-8", errors="ignore")
        forbidden_match = None if candidate.resolve() == Path(__file__).resolve() else FORBIDDEN_RUNTIME_NARRATIVE.search(text)
        if forbidden_match:
            forbidden_hits.append({"path": str(candidate.relative_to(skill_root)), "match": forbidden_match.group(0)})
        discovered_rules = rule_pattern.findall(text) if candidate.parent == references else []
        for rule_id in discovered_rules:
            owner = str(candidate.relative_to(skill_root))
            if rule_id in rule_owners and rule_owners[rule_id] != owner:
                duplicate_rules.append(rule_id)
            else:
                rule_owners[rule_id] = owner
    skill_text = (skill_root / "SKILL.md").read_text(encoding="utf-8")
    required_links = [f"references/{name}" for name in REQUIRED_REFERENCES] + [
        f"workflows/{name}/WORKFLOW.md" for name in REQUIRED_WORKFLOWS
    ]
    missing_links = [link for link in required_links if link not in skill_text]
    passed = not any((missing_references, missing_workflows, forbidden_hits, duplicate_rules, missing_links))
    return {
        "type": "writing-system-verification",
        "schema_version": 2,
        "path": str(skill_root),
        "compatibility_template_argument": str(path) if path is not None else None,
        "missing_references": missing_references,
        "missing_workflows": missing_workflows,
        "missing_skill_links": missing_links,
        "duplicate_rule_ids": sorted(set(duplicate_rules)),
        "forbidden_runtime_narrative": forbidden_hits,
        "rules": len(rule_owners),
        "passed": passed,
    }


def _paragraphs(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n\s*\n", text.replace("\r\n", "\n")) if part.strip()]


def _sentences(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"[。！？!?]+", text) if part.strip()]


def _compact_length(text: str) -> int:
    return len(re.sub(r"\s+", "", text))


def _issue(kind: str, level: str, message: str, match: re.Match[str] | None = None, paragraph: int | None = None, **extra: Any) -> dict[str, Any]:
    item: dict[str, Any] = {"kind": kind, "level": level, "message": message}
    if match is not None:
        item.update({"offset_start": match.start(), "offset_end": match.end(), "excerpt": match.group(0)})
    if paragraph is not None:
        item["paragraph"] = paragraph
    item.update(extra)
    return item


def count_proposal_characters(text: str) -> int:
    return _compact_length(text)

_CJK_WORD_CHAR = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")
_LATIN_WORD_SEGMENT = re.compile(r"[A-Za-z0-9]+(?:[._%+/-][A-Za-z0-9]+)*")


def count_proposal_words(text: str) -> int:
    """Count proposal length the way Chinese official forms do.

    Each CJK ideograph counts as one word; each contiguous Latin/digit token
    (including DOI-style punctuation) counts as one word. Whitespace and
    Chinese/ASCII punctuation do not count.
    """
    return len(_CJK_WORD_CHAR.findall(text)) + len(_LATIN_WORD_SEGMENT.findall(text))


def _check_length(count: int, limit: int, mode: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    if limit <= 0:
        raise ValueError("word limit must be positive")
    if mode not in {"first-draft", "strict"}:
        raise ValueError("length mode must be first-draft or strict")
    # Unified policy: the proposal must reach the stated word target and may
    # exceed it by at most 10 percent. The unit is proposal "words" (each CJK
    # character or Latin/digit token counts as one), not raw characters.
    ratio = count / limit
    issues: list[dict[str, Any]] = []
    target_min, target_max, permitted_max = limit, round(limit * 1.10), round(limit * 1.10)
    if count < limit:
        issues.append({
            "kind": "under-target-length",
            "level": "blocked",
            "message": "the proposal is below the stated word target; expand concepts, evidence, mechanisms, methods, controls, innovation derivation, and result-use paths instead of padding",
        })
    elif count > permitted_max:
        issues.append({
            "kind": "over-permitted-length",
            "level": "blocked",
            "message": "the proposal exceeds the stated word target by more than 10 percent; remove genuine repetition before delivery",
        })
    return {
        "mode": mode,
        "unit": "words",
        "limit": limit,
        "actual": count,
        "ratio": round(ratio, 3),
        "target_min": target_min,
        "target_max": target_max,
        "permitted_max": permitted_max,
    }, issues


def _paragraph_density_issues(paragraphs: list[str]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    lengths = [count_proposal_words(paragraph) for paragraph in paragraphs]
    issues: list[dict[str, Any]] = []
    if len(lengths) >= 3:
        midpoint = float(statistics.median(lengths))
        if midpoint < 140:
            issues.append(_issue("fragmented-draft", "blocked", "the median paragraph is below 140 words; merge fragments that serve the same proposition", median=midpoint))
        elif midpoint < 220:
            issues.append(_issue("thin-paragraph-density", "needs-review", "the median explanatory paragraph is below the 220-word depth target", median=midpoint))
        elif midpoint > 520:
            issues.append(_issue("overloaded-paragraph-density", "needs-review", "the median paragraph is unusually long; split only where the proposition changes", median=midpoint))
    consecutive = 0
    longest = 0
    for length in lengths:
        if 20 <= length < 140:
            consecutive += 1
            longest = max(longest, consecutive)
        else:
            consecutive = 0
    if longest >= 3:
        issues.append(_issue("consecutive-fragment-paragraphs", "blocked", "three or more consecutive short paragraphs fragment the argument", longest_run=longest))
    elif longest == 2:
        issues.append(_issue("adjacent-short-paragraphs", "needs-review", "merge adjacent short paragraphs when they perform the same argumentative function", longest_run=longest))
    return {
        "lengths": lengths,
        "median": float(statistics.median(lengths)) if lengths else 0.0,
        "target_range": [220, 420],
        "longest_short_run": longest,
    }, issues


def _repetition_issues(text: str) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    sentences = _sentences(text)
    prefixes: dict[str, list[int]] = {}
    for index, sentence in enumerate(sentences, start=1):
        prefix = re.sub(r"^[“\"'（(\s]+", "", sentence)[:5]
        if len(prefix) >= 4:
            prefixes.setdefault(prefix, []).append(index)
    repeated = {prefix: positions for prefix, positions in prefixes.items() if len(positions) >= 3}
    if repeated:
        issues.append(_issue("repeated-sentence-openers", "blocked", "three or more sentences reuse the same opening and create mechanical rhythm", repeated=repeated))
    symmetric_hits = list(SYMMETRIC_PATTERN.finditer(text))
    if len(symmetric_hits) >= 2:
        issues.append(_issue("forced-symmetry", "blocked", "repeated symmetrical constructions replace analysis with patterned prose", count=len(symmetric_hits)))
    for index, paragraph in enumerate(_paragraphs(text), start=1):
        if paragraph.count("；") + paragraph.count(";") >= 4:
            issues.append(_issue("semicolon-parallelism", "needs-review", "replace a long semicolon series with a developed argument or a template-required list", paragraph=index))
    return issues


def _evidence_issues(text: str, evidence_cards: list[dict[str, Any]] | None) -> tuple[list[str], list[dict[str, Any]]]:
    markers = list(dict.fromkeys(
        match.group(0)
        for pattern in (NUMERIC_CITATION, AUTHOR_YEAR_CITATION)
        for match in pattern.finditer(text)
    ))
    issues: list[dict[str, Any]] = []
    # Literature/data/policy claims must be explicitly cited in the final DOCX.
    # A sentence that says "research shows..." without an inline [n] marker is
    # indistinguishable from the applicant's own assertion and is not delivered.
    for index, paragraph in enumerate(_paragraphs(text), start=1):
        if CITATION_HINT_RE.search(paragraph) and not NUMERIC_CITATION.search(paragraph) and not AUTHOR_YEAR_CITATION.search(paragraph):
            issues.append(_issue(
                "missing-inline-citation",
                "blocked",
                "literature, data or policy claims must carry an inline numeric citation such as [1] in the same paragraph",
                paragraph=index,
            ))
    if not markers and not evidence_cards:
        return markers, issues
    if evidence_cards is None:
        return markers, issues + [_issue("missing-evidence-cards", "blocked", "new citations must be backed by verified in-memory evidence cards", markers=markers)]
    if not isinstance(evidence_cards, list) or not all(isinstance(item, dict) for item in evidence_cards):
        return markers, issues + [_issue("invalid-evidence-cards", "blocked", "evidence cards must be an array of objects")]
    indexed = {str(card.get("citation_key", "")).strip(): card for card in evidence_cards if str(card.get("citation_key", "")).strip()}
    for marker in markers:
        card = indexed.get(marker)
        if card is None:
            issues.append(_issue("untraced-citation", "blocked", "a citation in new prose has no matching evidence card", marker=marker))
            continue
        source = str(card.get("source", ""))
        locator = str(card.get("locator", ""))
        has_page = bool(re.search(r"https?://|doi\.org/|\b10\.\d{4,9}/", source, re.I))
        if card.get("verified") is not True or not has_page or not locator.strip():
            issues.append(_issue("unverified-citation", "blocked", "a cited evidence card lacks page-level verification, a source page or DOI, or a locator", marker=marker))
    for card in evidence_cards:
        if not card.get("used_in_text"):
            continue
        key = str(card.get("citation_key", "")).strip()
        if key and key not in markers:
            issues.append(_issue("declared-citation-not-found", "blocked", "an evidence card marked as used is not present in new prose as an inline citation", marker=key))
        elif key and not NUMERIC_CITATION.fullmatch(key):
            issues.append(_issue("non-numeric-citation-key", "blocked", "evidence-card citation keys must use numeric [n] form; renumber the cards and reference list accordingly", marker=key))
    return markers, issues


def _fact_consistency_issues(text: str, fact_ledger: dict[str, Any] | None) -> list[dict[str, Any]]:
    if fact_ledger is None:
        return []
    if not isinstance(fact_ledger, dict):
        return [_issue("invalid-fact-ledger", "blocked", "the fact ledger must be an object")]
    issues: list[dict[str, Any]] = []
    for field, raw in fact_ledger.items():
        record = raw if isinstance(raw, dict) else {"expected": raw}
        expected = str(record.get("expected", "")).strip()
        forbidden = [str(value).strip() for value in record.get("forbidden", []) if str(value).strip()]
        conflicts = [value for value in forbidden if value != expected and value in text]
        if conflicts:
            issues.append(_issue("fact-inconsistency", "blocked", "new prose contradicts a locked project fact", field=str(field), expected=expected, conflicts=conflicts))
    return issues


def _normalize_for_prompt_match(text: str) -> str:
    normalized = re.sub(r"[\s，。；、！？：:,.!?;（）()【】\[\]“”‘’\-—]+", "", text)
    return re.sub(r"本项目|本课题|本研究", "", normalized)


def _instruction_contamination_issues(text: str, task_prompt: str | None) -> list[dict[str, Any]]:
    issues = [
        _issue("user-instruction-residue", "blocked", "translate user instructions into evidence-bounded academic decisions; never copy task-management language into the proposal", match)
        for match in META_INSTRUCTION_RESIDUE.finditer(text)
    ]
    if not task_prompt:
        return issues
    normalized_text = _normalize_for_prompt_match(text)
    copied: list[str] = []
    for fragment in re.split(r"[\n。；;！？!?]+", task_prompt):
        fragment = fragment.strip()
        normalized = _normalize_for_prompt_match(fragment)
        normalized = re.sub(r"^(?:请你|请帮我|麻烦|请)", "", normalized)
        if not PROMPT_DIRECTIVE_CUE.search(fragment) or len(normalized) < 8:
            continue
        if normalized in normalized_text:
            copied.append(fragment[:120])
    if copied:
        issues.append(_issue("copied-prompt-directive", "blocked", "raw prompt directives must be internalized as design decisions and cannot appear verbatim in proposal prose", fragments=copied[:10]))
    return issues


def check_new_prose(
    text: str,
    review_scores: dict[str, float] | None = None,
    word_limit: int | None = None,
    length_mode: str = "first-draft",
    evidence_cards: list[dict[str, Any]] | None = None,
    fact_ledger: dict[str, Any] | None = None,
    task_prompt: str | None = None,
    style_profile: dict[str, Any] | None = None,
) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    for kind, pattern in HARD_PATTERNS.items():
        for match in pattern.finditer(text):
            issues.append(_issue(kind, "blocked", "rewrite the newly written prose before insertion", match))
    for match in OVERCLAIM_RE.finditer(text):
        before = text[:match.start()]
        paragraph_context = before.rsplit(chr(10)+chr(10), 1)[-1]
        level = "needs-review" if EVIDENCE_MARKER.search(paragraph_context) else "blocked"
        issues.append(_issue(
            "unsupported-inflation", level,
            "restrict the claim to the verified object, time, sample and boundary; strong novelty words are acceptable only with cited material evidence",
            match,
        ))
    for match in COLLOQUIAL.finditer(text):
        issues.append(_issue("colloquial", "blocked", "replace colloquial explanation with precise academic prose", match))
    for match in EMPTY_META.finditer(text):
        issues.append(_issue("empty-meta", "blocked", "remove meta-commentary and state the substantive judgment", match))

    paragraphs = _paragraphs(text)
    for index, paragraph in enumerate(paragraphs, start=1):
        colon_count = len(COLON_RE.findall(paragraph))
        dash_count = len(DASH_RE.findall(paragraph))
        if colon_count >= 4:
            issues.append(_issue("colon-overuse", "blocked", "too many colons in one paragraph; convert the list-like formatting into natural prose", paragraph=index, count=colon_count))
        elif colon_count == 3:
            issues.append(_issue("colon-heavy-paragraph", "needs-review", "three colons in one paragraph suggest list-like formatting inside prose", paragraph=index, count=colon_count))
        if dash_count >= 3:
            issues.append(_issue("dash-overuse", "blocked", "too many dashes in one paragraph", paragraph=index, count=dash_count))
        elif dash_count == 2:
            issues.append(_issue("dash-heavy-paragraph", "needs-review", "two dashes in one paragraph suggest forced contrast phrasing", paragraph=index, count=dash_count))
    sequence_hits = list(MECHANICAL_SEQUENCE_RE.finditer(text))
    if len(sequence_hits) >= 3:
        issues.append(_issue("mechanical-sequence", "needs-review", "repeated sequential connectives create mechanical rhythm; keep them only for genuine enumeration", count=len(sequence_hits)))
    density, density_issues = _paragraph_density_issues(paragraphs)
    issues.extend(density_issues)
    issues.extend(_repetition_issues(text))
    citation_markers, evidence_issues = _evidence_issues(text, evidence_cards)
    issues.extend(evidence_issues)
    issues.extend(_fact_consistency_issues(text, fact_ledger))
    issues.extend(_instruction_contamination_issues(text, task_prompt))
    for index, paragraph in enumerate(paragraphs, start=1):
        if VAGUE_ENDING.search(paragraph) and not EVIDENCE_MARKER.search(paragraph):
            issues.append(_issue("empty-significance", "blocked", "name the object, mechanism, user, or boundary instead of an empty value claim", paragraph=index))
        if len(paragraph) >= 100 and not EVIDENCE_MARKER.search(paragraph):
            issues.append(_issue("low-information-paragraph", "needs-review", "the paragraph lacks an object, mechanism, evidence, comparison, method, or boundary", paragraph=index))
        if METHOD_NAME.search(paragraph) and not METHOD_OPERATION.search(paragraph):
            issues.append(_issue("method-without-operation", "needs-review", "state why the method is used, what material it uses, and how it is executed", paragraph=index))
        if "创新" in paragraph and not re.search(r"不足|缺口|局限|研究内容|方法|成果|解释|机制|比较", paragraph):
            issues.append(_issue("generic-innovation", "needs-review", "bind the innovation to a documented gap, project component, and expected result", paragraph=index))

    placeholders = [{"text": match.group(0), "offset_start": match.start(), "offset_end": match.end()} for match in PLACEHOLDER.finditer(text)]

    score_result = None
    if review_scores is not None:
        required = {"value", "design", "method", "innovation", "foundation", "feasibility"}
        missing = sorted(required - set(review_scores))
        below = {key: value for key, value in review_scores.items() if key in required and float(value) < 8.0}
        score_result = {"missing": missing, "below_eight": below, "passed": not missing and not below}
        if missing:
            issues.append({"kind": "review-score-missing", "level": "blocked", "message": f"missing internal review dimensions: {missing}"})
        if below:
            issues.append({"kind": "review-score-low", "level": "needs-review", "message": "make one evidence-bounded targeted revision", "scores": below})

    length_result = None
    character_count = count_proposal_characters(text)
    word_count = count_proposal_words(text)
    if word_limit is not None:
        length_result, length_issues = _check_length(word_count, word_limit, length_mode)
        issues.extend(length_issues)

    style_result = check_style(text, style_profile)
    issues.extend(style_result.get("issues", []))

    hard = [item for item in issues if item["level"] == "blocked"]
    soft = [item for item in issues if item["level"] == "needs-review"]
    gate = "blocked" if hard else "needs-review" if soft else "passed"
    return {
        "type": "new-prose-quality.v2",
        "gate": gate,
        "metrics": {"characters": len(text), "proposal_characters": character_count, "words": word_count, "proposal_words": word_count, "paragraphs": len(paragraphs), "hard_issues": len(hard), "review_issues": len(soft)},
        "paragraph_density": density,
        "length_target": length_result,
        "placeholders": placeholders,
        "citation_markers": citation_markers,
        "review_scores": score_result,
        "style": style_result,
        "issues": issues,
    }


def _semantic_table_signature(node: dict[str, Any]) -> list[list[tuple[int, int, str | None]]]:
    return [
        [
            (
                int(cell.get("start_column", 0)),
                int(cell.get("grid_span", 1)),
                cell.get("vertical_merge"),
            )
            for cell in row.get("cells", [])
        ]
        for row in node.get("rows", [])
    ]


def _docx_table_signature(table: Any) -> list[list[tuple[int, int, str | None]]]:
    from docx.oxml.ns import qn

    signature: list[list[tuple[int, int, str | None]]] = []
    for tr in table._tbl.tr_lst:
        row_signature: list[tuple[int, int, str | None]] = []
        column = 0
        for tc in tr.tc_lst:
            tc_pr = tc.tcPr
            grid_span_element = tc_pr.find(qn("w:gridSpan")) if tc_pr is not None else None
            span_raw = grid_span_element.get(qn("w:val")) if grid_span_element is not None else None
            grid_span = int(span_raw) if span_raw and span_raw.isdigit() else 1
            vmerge_element = tc_pr.find(qn("w:vMerge")) if tc_pr is not None else None
            vertical_merge = None if vmerge_element is None else vmerge_element.get(qn("w:val")) or "continue"
            row_signature.append((column, grid_span, vertical_merge))
            column += grid_span
        signature.append(row_signature)
    return signature


def _style_font_name(style: Any) -> str | None:
    from docx.oxml.ns import qn

    rpr = style.element.rPr
    if rpr is not None and rpr.rFonts is not None:
        return rpr.rFonts.get(qn("w:eastAsia")) or style.font.name
    return style.font.name


def _style_contract_issues(document: Any, profile: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    expected = {
        "Normal": (profile["body_font"], profile["body_size"]),
        "Table Text": (profile["table_font"], profile["table_size"]),
    }
    for style_name, (font_name, font_size) in expected.items():
        try:
            style = document.styles[style_name]
        except KeyError:
            issues.append(_issue("missing-normalized-style", "blocked", "a required normalized paragraph style is absent", style=style_name))
            continue
        actual_name = _style_font_name(style)
        actual_size = style.font.size.pt if style.font.size else None
        if actual_name != font_name or actual_size is None or abs(actual_size - float(font_size)) > 0.2:
            issues.append(_issue(
                "normalized-font-mismatch", "blocked",
                "the rebuilt document font profile differs from the selected explicit or default profile",
                style=style_name, expected_font=font_name, actual_font=actual_name,
                expected_size=font_size, actual_size=actual_size,
            ))
        expected_bold = False
        if bool(style.font.bold) != expected_bold:
            issues.append(_issue("normalized-bold-mismatch", "blocked", "normalized body and table styles must remain regular; semantic headings use bold text in Normal paragraphs", style=style_name, expected=expected_bold, actual=style.font.bold))
    normal = document.styles["Normal"].paragraph_format
    expected_indent = float(profile["body_size"]) * float(profile["first_line_chars"])
    actual_indent = normal.first_line_indent.pt if normal.first_line_indent else 0.0
    if abs(actual_indent - expected_indent) > 0.5:
        issues.append(_issue("body-indent-mismatch", "blocked", "body first-line indentation differs from the normalized profile", expected=expected_indent, actual=actual_indent))
    actual_spacing = normal.line_spacing
    if profile.get("line_spacing_rule") == "exactly":
        actual_spacing_value = actual_spacing.pt if hasattr(actual_spacing, "pt") else actual_spacing
    else:
        actual_spacing_value = actual_spacing
    if not isinstance(actual_spacing_value, (int, float)) or abs(float(actual_spacing_value) - float(profile["line_spacing"])) > 0.05:
        issues.append(_issue("body-line-spacing-mismatch", "blocked", "body line spacing differs from the normalized profile", expected=profile["line_spacing"], actual=str(actual_spacing_value)))
    return issues


def check_docx_structure(path: Path, expected_structure: dict[str, Any] | None = None) -> dict[str, Any]:
    from docx import Document

    path = path.expanduser().resolve()
    issues: list[dict[str, Any]] = []
    if not path.exists() or path.suffix.lower() != ".docx" or not zipfile.is_zipfile(path):
        return {"type": "docx-structure.v1", "gate": "blocked", "issues": [{"kind": "invalid-docx", "level": "blocked", "message": "the final artifact is not a valid DOCX package"}]}
    try:
        document = Document(str(path))
    except Exception as exc:
        return {"type": "docx-structure.v1", "gate": "blocked", "issues": [{"kind": "unreadable-docx", "level": "blocked", "message": str(exc)}]}

    used_styles = {paragraph.style.name for paragraph in document.paragraphs if paragraph.style}
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                used_styles.update(paragraph.style.name for paragraph in cell.paragraphs if paragraph.style)
    unexpected = sorted(style for style in used_styles if style not in ALLOWED_PARAGRAPH_STYLES)
    if unexpected:
        issues.append(_issue("unexpected-paragraph-style", "blocked", "the rebuilt document uses styles outside the normalized style set", styles=unexpected))
    office_headings = sorted(style for style in used_styles if style in FORBIDDEN_OFFICE_HEADING_STYLES or style.startswith("Heading "))
    if office_headings:
        issues.append(_issue("office-heading-style-used", "blocked", "titles and headings must be bold Normal paragraphs, never Office Title or Heading styles", styles=office_headings))

    expected_semantic = [
        (node.get("kind"), node.get("text", "").strip(), int(node.get("internal_depth") or node.get("level") or 1))
        for node in (expected_structure or {}).get("nodes", [])
        if node.get("kind") in {"document-title", "heading", "internal-heading"} and node.get("text", "").strip()
    ]
    expected_internal = [(text, depth) for kind, text, depth in expected_semantic if kind == "internal-heading"]
    expected_semantic_texts = {text for _, text, _ in expected_semantic}
    direct_format_runs: list[dict[str, Any]] = []
    all_paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)
    for paragraph_index, paragraph in enumerate(all_paragraphs, start=1):
        for run_index, run in enumerate(paragraph.runs, start=1):
            if not run.text.strip():
                continue
            allowed_semantic_bold = (
                paragraph.style is not None
                and paragraph.style.name == "Normal"
                and paragraph.text.strip() in expected_semantic_texts
                and run.bold is True
                and all(value is None for value in (run.font.name, run.font.size, run.italic, run.underline, run.font.color.rgb))
            )
            if not allowed_semantic_bold and any(value is not None for value in (run.font.name, run.font.size, run.bold, run.italic, run.underline, run.font.color.rgb)):
                direct_format_runs.append({"paragraph": paragraph_index, "run": run_index})
    if direct_format_runs:
        issues.append(_issue("unexpected-direct-run-format", "blocked", "source direct formatting leaked into normalized text runs", runs=direct_format_runs[:20], count=len(direct_format_runs)))

    placeholder_runs: list[dict[str, Any]] = []
    illegal_red_runs: list[dict[str, Any]] = []
    markdown_runs: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(all_paragraphs, start=1):
        if re.search(r"(?:^|\n)\s*#{2,6}\s+", paragraph.text):
            markdown_runs.append({"paragraph": paragraph_index, "text": paragraph.text[:120]})
        for run_index, run in enumerate(paragraph.runs, start=1):
            text = run.text or ""
            style_name = run.style.name if run.style else None
            contains_placeholder = bool(PLACEHOLDER.search(text))
            if contains_placeholder:
                color = run.style.font.color.rgb if run.style and run.style.font.color else None
                if style_name != "Missing Fact" or str(color).upper() != "FF0000":
                    placeholder_runs.append({"paragraph": paragraph_index, "run": run_index, "style": style_name, "color": str(color) if color else None})
            elif style_name == "Missing Fact":
                illegal_red_runs.append({"paragraph": paragraph_index, "run": run_index, "text": text[:80]})
    if placeholder_runs:
        issues.append(_issue("missing-fact-not-red", "blocked", "every unresolved fact placeholder must use the red Missing Fact character style", runs=placeholder_runs[:20]))
    if illegal_red_runs:
        issues.append(_issue("non-placeholder-red", "blocked", "only unresolved fact placeholders may use red text", runs=illegal_red_runs[:20]))
    if markdown_runs:
        issues.append(_issue("markdown-heading-leak", "blocked", "Markdown markers must be converted into semantic internal headings before DOCX construction", paragraphs=markdown_runs[:20]))

    decorated_paragraphs: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(all_paragraphs, start=1):
        paragraph_properties = paragraph._p.pPr
        if paragraph_properties is None:
            continue
        decorations = [name for name in ("pBdr", "shd") if paragraph_properties.find(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{name}") is not None]
        if decorations:
            decorated_paragraphs.append({"paragraph": paragraph_index, "decorations": decorations, "text": paragraph.text[:100]})
    if decorated_paragraphs:
        issues.append(_issue("decorative-border-or-shading", "blocked", "semantic proposal text must not carry title borders, horizontal rules, or inherited paragraph shading", paragraphs=decorated_paragraphs[:20]))
    decorated_styles: list[dict[str, Any]] = []
    for style_name in sorted(used_styles):
        try:
            paragraph_properties = document.styles[style_name].element.pPr
        except (KeyError, AttributeError):
            continue
        if paragraph_properties is None:
            continue
        decorations = [name for name in ("pBdr", "shd") if paragraph_properties.find(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{name}") is not None]
        if decorations:
            decorated_styles.append({"style": style_name, "decorations": decorations})
    if decorated_styles:
        issues.append(_issue("decorated-paragraph-style", "blocked", "used paragraph styles must not add title borders, horizontal rules, or shading", styles=decorated_styles))

    semantic_layout: list[dict[str, Any]] = []
    semantic_actual: list[tuple[str, str, int]] = []
    paragraph_cursor = 0
    body_size = float(((expected_structure or {}).get("style_profile") or DEFAULT_STYLE_PROFILE).get("body_size", 12.0))
    for kind, text, depth in expected_semantic:
        match_index = next((index for index in range(paragraph_cursor, len(document.paragraphs)) if document.paragraphs[index].text.strip() == text), None)
        if match_index is None:
            semantic_layout.append({"heading": text, "reason": "semantic heading is missing or out of order"})
            continue
        paragraph = document.paragraphs[match_index]
        paragraph_cursor = match_index + 1
        semantic_actual.append((kind, text, depth))
        if not paragraph.style or paragraph.style.name != "Normal":
            semantic_layout.append({"heading": text, "reason": "all document and section headings must use Normal paragraph style", "actual_style": paragraph.style.name if paragraph.style else None})
        nonempty_runs = [run for run in paragraph.runs if run.text.strip()]
        if not nonempty_runs or any(run.bold is not True for run in nonempty_runs):
            semantic_layout.append({"heading": text, "reason": "semantic heading text must be bold"})
        first_indent = paragraph.paragraph_format.first_line_indent.pt if paragraph.paragraph_format.first_line_indent else 0.0
        left_indent = paragraph.paragraph_format.left_indent.pt if paragraph.paragraph_format.left_indent else 0.0
        expected_left = body_size * 2.0 if kind == "internal-heading" and depth == 2 else 0.0
        if abs(first_indent) > 0.5 or abs(left_indent - expected_left) > 0.5:
            semantic_layout.append({"heading": text, "reason": "semantic heading indentation is invalid", "expected_left": expected_left, "actual_left": left_indent, "actual_first": first_indent})
        if kind == "document-title" and paragraph.alignment != 1:
            semantic_layout.append({"heading": text, "reason": "document title must be centered without using the Title style"})
    if semantic_actual != expected_semantic:
        issues.append(_issue("heading-order-or-level", "blocked", "semantic title or heading order changed during DOCX construction", expected=expected_semantic, actual=semantic_actual))
    if semantic_layout:
        issues.append(_issue("semantic-heading-layout", "blocked", "all semantic headings must be ordinary bold body-style paragraphs without Office decorations", headings=semantic_layout[:20]))

    if expected_internal:
        actual_internal: list[tuple[str, int]] = []
        internal_layout: list[dict[str, Any]] = []
        blocks = list(document.iter_inner_content())
        text_to_depth = {text: depth for text, depth in expected_internal}
        body_size = float(((expected_structure or {}).get("style_profile") or DEFAULT_STYLE_PROFILE).get("body_size", 12.0))
        for block_index, block in enumerate(blocks):
            if not hasattr(block, "text") or block.text.strip() not in text_to_depth:
                continue
            text = block.text.strip()
            depth = text_to_depth[text]
            actual_internal.append((text, depth))
            if not block.style or block.style.name != "Normal":
                internal_layout.append({"heading": text, "reason": "internal heading must use Normal paragraph style"})
            nonempty_runs = [run for run in block.runs if run.text.strip()]
            if not nonempty_runs or any(run.bold is not True for run in nonempty_runs):
                internal_layout.append({"heading": text, "reason": "internal heading text must be bold"})
            first_indent = block.paragraph_format.first_line_indent.pt if block.paragraph_format.first_line_indent else 0.0
            left_indent = block.paragraph_format.left_indent.pt if block.paragraph_format.left_indent else 0.0
            expected_left = body_size * 2.0 if depth == 2 else 0.0
            if abs(first_indent) > 0.5 or abs(left_indent - expected_left) > 0.5:
                internal_layout.append({"heading": text, "reason": "internal heading indentation does not match its depth", "expected_left": expected_left, "actual_left": left_indent, "actual_first": first_indent})
            next_block = blocks[block_index + 1] if block_index + 1 < len(blocks) else None
            if next_block is None:
                internal_layout.append({"heading": text, "reason": "missing following body"})
                continue
            if hasattr(next_block, "text") and next_block.text.strip() in text_to_depth:
                if text_to_depth[next_block.text.strip()] != 2 or depth != 1:
                    internal_layout.append({"heading": text, "reason": "only a level-two child may immediately follow a level-one internal heading"})
                continue
            if hasattr(next_block, "rows"):
                continue
            if not getattr(next_block, "style", None) or next_block.style.name != "Normal" or not next_block.text.strip():
                internal_layout.append({"heading": text, "reason": "a leaf internal heading must be followed immediately by body text"})
        if actual_internal != expected_internal:
            issues.append(_issue("internal-heading-order", "blocked", "internal heading order or depth changed during DOCX construction", expected=expected_internal, actual=actual_internal))
        if internal_layout:
            issues.append(_issue("internal-heading-layout", "blocked", "internal headings must be ordinary bold body paragraphs with valid nesting and immediate body text", headings=internal_layout[:20]))

    paragraphs = document.paragraphs
    consecutive_empty = 0
    max_empty = 0
    heading_without_body: list[int] = []
    for index, paragraph in enumerate(paragraphs):
        if not paragraph.text.strip() and not paragraph._p.xpath(".//a:blip"):
            consecutive_empty += 1
            max_empty = max(max_empty, consecutive_empty)
        else:
            consecutive_empty = 0
    blocks = list(document.iter_inner_content())
    heading_number = 0
    for block_index, block in enumerate(blocks):
        if not hasattr(block, "style") or not block.style or not block.style.name.startswith("Heading"):
            continue
        heading_number += 1
        current_level_match = re.search(r"(\d+)$", block.style.name)
        current_level = int(current_level_match.group(1)) if current_level_match else 1
        has_body = False
        for candidate in blocks[block_index + 1:]:
            if hasattr(candidate, "rows"):
                has_body = True
                break
            if not candidate.text.strip() and not candidate._p.xpath(".//a:blip"):
                continue
            if candidate.style and candidate.style.name.startswith("Heading"):
                next_level_match = re.search(r"(\d+)$", candidate.style.name)
                next_level = int(next_level_match.group(1)) if next_level_match else 1
                if next_level > current_level:
                    has_body = True
                break
            has_body = True
            break
        if not has_body:
            heading_without_body.append(heading_number)
    if max_empty >= 2:
        issues.append(_issue("blank-paragraph-spacing", "blocked", "use paragraph spacing instead of consecutive empty paragraphs", longest_run=max_empty))
    if heading_without_body:
        issues.append(_issue("orphan-heading", "needs-review", "one or more headings have no body before the next heading", paragraphs=heading_without_body))

    subheading_layout: list[dict[str, Any]] = []
    for block_index, block in enumerate(blocks):
        if not hasattr(block, "style") or not block.style or block.style.name not in {"Heading 2", "Heading 3"}:
            continue
        next_block = blocks[block_index + 1] if block_index + 1 < len(blocks) else None
        if next_block is None:
            subheading_layout.append({"heading": block.text, "reason": "missing following body"})
            continue
        if hasattr(next_block, "rows"):
            continue
        if hasattr(next_block, "style") and next_block.style and next_block.style.name == "Normal" and next_block.text.strip():
            continue
        if hasattr(next_block, "style") and next_block.style and block.style.name == "Heading 2" and next_block.style.name == "Heading 3":
            continue
        subheading_layout.append({"heading": block.text, "reason": "body must immediately follow in a separate Normal paragraph"})
    if subheading_layout:
        issues.append(_issue("subheading-body-layout", "blocked", "small headings must occupy their own paragraph and be followed immediately by indented body text", headings=subheading_layout[:20]))

    profile = dict(DEFAULT_STYLE_PROFILE)
    if expected_structure:
        profile.update(expected_structure.get("style_profile", {}) or {})
    section = document.sections[0]
    page_width_cm = section.page_width.cm if section.page_width else 0
    page_height_cm = section.page_height.cm if section.page_height else 0
    expected_width = float(profile.get("page_width_cm", 21.0))
    expected_height = float(profile.get("page_height_cm", 29.7))
    if abs(page_width_cm - expected_width) > 0.2 or abs(page_height_cm - expected_height) > 0.2:
        issues.append(_issue("page-size-mismatch", "blocked", "the rebuilt page size differs from the selected explicit or default profile", expected=[expected_width, expected_height], actual=[round(page_width_cm, 2), round(page_height_cm, 2)]))
    margin_pairs = (
        ("top_margin_cm", "margin_top_cm", section.top_margin),
        ("bottom_margin_cm", "margin_bottom_cm", section.bottom_margin),
        ("left_margin_cm", "margin_left_cm", section.left_margin),
        ("right_margin_cm", "margin_right_cm", section.right_margin),
    )
    for profile_key, _label, margin in margin_pairs:
        actual_cm = margin.cm if margin is not None else 0.0
        expected_cm = float(profile.get(profile_key, 2.54))
        if abs(actual_cm - expected_cm) > 0.1:
            issues.append(_issue("page-margin-mismatch", "blocked", "the rebuilt page margins differ from the selected explicit or default profile", field=profile_key, expected=expected_cm, actual=round(actual_cm, 2)))

    with zipfile.ZipFile(path) as package:
        names = package.namelist()
        missing_parts = sorted({
            name for name in ("[Content_Types].xml", "word/document.xml", "word/_rels/document.xml.rels")
            if name not in names
        })
        if missing_parts:
            issues.append(_issue("missing-docx-relationship-part", "blocked", "the DOCX package is missing a required relationship or content part", parts=missing_parts))
        oversized_parts = sorted(name for name in names if package.getinfo(name).file_size > 64 * 1024 * 1024)
        if oversized_parts:
            issues.append(_issue("oversized-docx-part", "blocked", "a DOCX package part exceeds the 64 MB safety limit", parts=oversized_parts[:10]))
        image_parts = [name for name in names if name.startswith("word/media/")]
    metrics = {
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "images": len(image_parts),
        "used_paragraph_styles": sorted(used_styles),
        "page_width_cm": round(page_width_cm, 2),
        "page_height_cm": round(page_height_cm, 2),
    }
    if expected_structure:
        metadata = expected_structure.get("metadata", {})
        if metrics["tables"] != int(metadata.get("table_count", metrics["tables"])):
            issues.append(_issue("table-topology-count", "blocked", "table count changed during semantic rebuild", expected=metadata.get("table_count"), actual=metrics["tables"]))
        expected_images = int(metadata.get("image_count", 0))
        if expected_images and metrics["images"] < expected_images:
            issues.append(_issue("missing-official-images", "blocked", "one or more source images were not preserved", expected=expected_images, actual=metrics["images"]))
        expected_tables = [node for node in expected_structure.get("nodes", []) if node.get("kind") == "table"]
        for index, (expected_table, actual_table) in enumerate(zip(expected_tables, document.tables), start=1):
            expected_signature = _semantic_table_signature(expected_table)
            actual_signature = _docx_table_signature(actual_table)
            if actual_signature != expected_signature:
                issues.append(_issue("table-merge-topology", "blocked", "table rows, columns, or merged cells changed during semantic rebuild", table=index, expected=expected_signature, actual=actual_signature))
            expected_columns = int(expected_table.get("columns", 0))
            actual_columns = len(actual_table.columns)
            if expected_columns and actual_columns != expected_columns:
                issues.append(_issue("table-column-count", "blocked", "table column count changed during semantic rebuild", table=index, expected=expected_columns, actual=actual_columns))
        issues.extend(_style_contract_issues(document, profile))

    hard = [item for item in issues if item["level"] == "blocked"]
    soft = [item for item in issues if item["level"] == "needs-review"]
    return {"type": "docx-structure.v1", "gate": "blocked" if hard else "needs-review" if soft else "passed", "metrics": metrics, "issues": issues}


def check_writing_ledger(ledger: dict[str, Any] | None) -> dict[str, Any]:
    if ledger is None:
        return {"type": "writing-ledger-check.v1", "gate": "passed", "enforced": False, "issues": [], "coverage": None}
    issues: list[dict[str, Any]] = []
    if ledger.get("type") != "writing-ledger.v1":
        issues.append(_issue("invalid-writing-ledger", "blocked", "writing ledger type must be writing-ledger.v1"))
    for reasoning_issue in validate_reasoning_protocol(ledger):
        details = {
            key: value
            for key, value in reasoning_issue.items()
            if key not in {"kind", "message"}
        }
        issues.append(
            _issue(
                reasoning_issue.get("kind", "invalid-reasoning-protocol"),
                "blocked",
                reasoning_issue.get("message", "structured reasoning protocol is invalid"),
                **details,
            )
        )
    instruction_translation = ledger.get("instruction_translation")
    if isinstance(instruction_translation, dict) and instruction_translation.get("required"):
        if instruction_translation.get("status") != "complete":
            issues.append(_issue("instruction-translation-incomplete", "blocked", "content directives must be converted into academic design decisions before drafting"))
        decisions = instruction_translation.get("decisions", [])
        decision_errors: list[dict[str, Any]] = []
        required_decision_fields = {"instruction_summary", "academic_decision", "basis", "target_sections"}
        if not isinstance(decisions, list) or not decisions:
            decision_errors.append({"reason": "no translated design decisions"})
        else:
            for index, decision in enumerate(decisions, start=1):
                if not isinstance(decision, dict):
                    decision_errors.append({"decision": index, "reason": "decision is not an object"})
                    continue
                missing = sorted(field for field in required_decision_fields if not decision.get(field))
                if missing:
                    decision_errors.append({"decision": index, "reason": "missing decision fields", "fields": missing})
        if decision_errors:
            issues.append(_issue("invalid-instruction-translation", "blocked", "each user directive must map to an evidence basis, an academic decision, and target sections without copying the directive into prose", decisions=decision_errors[:20]))
    method_modules = ledger.get("method_modules")
    if method_modules is None:
        modules = {int(item.get("module")): item.get("status") for item in ledger.get("modules", []) if str(item.get("module", "")).isdigit()}
        required_modules = {int(item.get("module")) for item in ledger.get("modules", []) if item.get("required", True) and str(item.get("module", "")).isdigit()}
        incomplete_modules = [index for index in sorted(required_modules) if modules.get(index) != "complete"]
        if incomplete_modules:
            issues.append(_issue("writing-modules-incomplete", "blocked", "all ten writing modules must be completed before delivery", modules=incomplete_modules))
    if method_modules is not None:
        expected_ids = [f"M{index}" for index in range(1, 11)]
        actual_ids = [item.get("module_id") for item in method_modules if isinstance(item, dict)]
        if actual_ids != expected_ids:
            issues.append(_issue("method-module-contract", "blocked", "method_modules must contain ordered M1 through M10"))
        incomplete_method_modules = [
            item.get("module_id")
            for item in method_modules
            if item.get("required", True) and item.get("status") != "complete"
        ]
        if incomplete_method_modules:
            issues.append(
                _issue(
                    "method-modules-incomplete",
                    "blocked",
                    "all required integrated Prompt modules must be completed before delivery",
                    modules=incomplete_method_modules,
                )
            )
        revision_count = int(ledger.get("m10_targeted_revision_count", 0) or 0)
        if revision_count > 1:
            issues.append(
                _issue(
                    "m10-revision-limit",
                    "blocked",
                    "M10 may trigger at most one targeted Markdown revision",
                    revision_count=revision_count,
                )
            )
        if revision_count == 1:
            affected = set(ledger.get("m10_affected_modules", [])) | {"M8", "M9", "M10"}
            incomplete_reruns = [
                item.get("module_id")
                for item in method_modules
                if item.get("module_id") in affected
                and (
                    item.get("status") != "complete"
                    or item.get("completed_revision") != item.get("input_revision")
                )
            ]
            if incomplete_reruns:
                issues.append(
                    _issue(
                        "m10-rerun-incomplete",
                        "blocked",
                        "affected modules and M8-M10 must be rerun after the single targeted revision",
                        modules=incomplete_reruns,
                    )
                )
    sections = ledger.get("sections", [])
    accepted = {"complete", "placeholder-complete", "source-present"}
    incomplete_sections = [item.get("section_id") for item in sections if item.get("status") not in accepted]
    if incomplete_sections:
        issues.append(_issue("writing-sections-incomplete", "blocked", "every required section must be complete or explicitly completed with a missing-fact placeholder", sections=incomplete_sections))
    if ledger.get("review", {}).get("status") != "complete":
        issues.append(_issue("writing-review-incomplete", "blocked", "independent review must be completed before delivery"))
    depth_metrics: list[dict[str, Any]] = []
    required_tree_fields = {"central_thesis", "subproblems", "facts", "evidence", "mechanisms", "methods", "boundary", "conclusion", "units"}
    required_unit_fields = {"unit_id", "title", "depth", "subproblem", "facts", "evidence", "mechanism", "method", "boundary", "conclusion", "body_character_count", "paragraph_count", "status"}
    for section in sections:
        tree = section.get("argument_tree")
        section_id = section.get("section_id")
        if not isinstance(tree, dict):
            issues.append(_issue("missing-argument-tree", "blocked", "every proposal section must carry an internal argument tree", section=section_id))
            continue
        missing_tree_fields = sorted(required_tree_fields - set(tree))
        if missing_tree_fields:
            issues.append(_issue("incomplete-argument-tree", "blocked", "the section argument tree is missing required reasoning fields", section=section_id, fields=missing_tree_fields))
        units = tree.get("units", []) if isinstance(tree.get("units"), list) else []
        unit_errors: list[dict[str, Any]] = []
        titles: list[str] = []
        for index, unit in enumerate(units):
            if not isinstance(unit, dict):
                unit_errors.append({"unit": index + 1, "reason": "unit is not an object"})
                continue
            missing_unit_fields = sorted(required_unit_fields - set(unit))
            if missing_unit_fields:
                unit_errors.append({"unit": index + 1, "reason": "missing reasoning fields", "fields": missing_unit_fields})
            title = str(unit.get("title", "")).strip()
            titles.append(title)
            if not title or GENERIC_INTERNAL_HEADING.fullmatch(title):
                unit_errors.append({"unit": index + 1, "reason": "empty or generic internal heading", "title": title})
            depth = int(unit.get("depth") or 0)
            if depth not in {1, 2}:
                unit_errors.append({"unit": index + 1, "reason": "internal heading depth must be one or two", "depth": depth})
            has_child = bool(unit.get("has_child"))
            if not has_child and (int(unit.get("body_character_count") or 0) <= 0 or int(unit.get("paragraph_count") or 0) <= 0):
                unit_errors.append({"unit": index + 1, "reason": "leaf internal heading has no substantive body"})
            if int(unit.get("paragraph_count") or 0) > 3:
                issues.append(_issue("overloaded-argument-unit", "needs-review", "an internal argument unit normally should contain one to three substantive paragraphs", section=section_id, unit=index + 1, paragraphs=unit.get("paragraph_count")))
        duplicate_titles = sorted({title for title in titles if title and titles.count(title) > 1})
        if duplicate_titles:
            unit_errors.append({"reason": "duplicate internal headings", "titles": duplicate_titles})
        if unit_errors:
            issues.append(_issue("invalid-argument-units", "blocked", "argument units must form meaningful headings with complete leaf bodies", section=section_id, units=unit_errors[:20]))
        requires_structure = bool(section.get("requires_internal_structure"))
        top_level_count = sum(int(unit.get("depth") or 0) == 1 for unit in units if isinstance(unit, dict))
        if requires_structure:
            if not str(tree.get("central_thesis", "")).strip():
                issues.append(_issue("missing-central-thesis", "blocked", "a complex section must state its central thesis in the writing ledger", section=section_id))
            if top_level_count < 2:
                issues.append(_issue("understructured-section", "blocked", "a complex section must contain at least two first-level argument units", section=section_id, top_level_units=top_level_count))
        if section.get("modified") and int(tree.get("actual_character_count") or 0) >= 600 and not units:
            issues.append(_issue("monolithic-long-section", "blocked", "a long generated section cannot remain one undifferentiated block", section=section_id, characters=tree.get("actual_character_count")))
        depth_metrics.append({
            "section_id": section_id,
            "required": requires_structure,
            "characters": int(tree.get("actual_character_count") or 0),
            "units": len(units),
            "top_level_units": top_level_count,
            "max_depth": max((int(unit.get("depth") or 0) for unit in units if isinstance(unit, dict)), default=0),
        })
    completed = sum(item.get("status") in accepted for item in sections)
    coverage = round(completed / len(sections), 4) if sections else 1.0
    hard = [item for item in issues if item["level"] == "blocked"]
    soft = [item for item in issues if item["level"] == "needs-review"]
    reasoning_plan = ledger.get("reasoning_plan")
    reasoning_metrics = None
    if isinstance(reasoning_plan, dict):
        reasoning_metrics = {
            "mode": reasoning_plan.get("mode"),
            "sections": len(reasoning_plan.get("sections", [])),
            "verification_questions": sum(
                len(section.get("verification_questions", []))
                for section in reasoning_plan.get("sections", [])
                if isinstance(section, dict)
            ),
            "revision_rounds": int(ledger.get("verification_summary", {}).get("revision_rounds", 0) or 0),
        }
    return {
        "type": "writing-ledger-check.v1",
        "gate": "blocked" if hard else "needs-review" if soft else "passed",
        "enforced": True,
        "issues": issues,
        "coverage": coverage,
        "argument_depth": depth_metrics,
        "reasoning_protocol": reasoning_metrics,
    }


def check_document(
    docx_path: Path,
    new_text: str,
    expected_structure: dict[str, Any] | None = None,
    review_scores: dict[str, float] | None = None,
    word_limit: int | None = None,
    length_mode: str = "first-draft",
    evidence_cards: list[dict[str, Any]] | None = None,
    fact_ledger: dict[str, Any] | None = None,
    writing_ledger: dict[str, Any] | None = None,
    task_prompt: str | None = None,
    style_profile: dict[str, Any] | None = None,
) -> dict[str, Any]:
    format_result = check_docx_structure(docx_path, expected_structure)
    content_result = check_new_prose(new_text, review_scores, word_limit, length_mode, evidence_cards, fact_ledger, task_prompt, style_profile)
    writing_result = check_writing_ledger(writing_ledger)
    gates = {format_result["gate"], content_result["gate"], writing_result["gate"]}
    gate = "blocked" if "blocked" in gates else "needs-review" if "needs-review" in gates else "passed"
    return {
        "type": "combined-final-gate.v1",
        "passes": sum(1 for value in gates if value == "passed"),
        "gate": gate,
        "format": format_result,
        "content": content_result,
        "writing_ledger": writing_result,
        "rendering_invoked": False,
        "screenshot_invoked": False,
        "pdf_export_invoked": False,
    }


def _read_text(path: str | None) -> str:
    return Path(path).read_text(encoding="utf-8-sig") if path else sys.stdin.read()


def _read_optional_text(path: str | None) -> str:
    return Path(path).read_text(encoding="utf-8-sig") if path else ""


def _read_json(path: str | None) -> dict[str, Any] | None:
    if not path:
        return None
    value = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError("JSON input must be an object")
    return value


def _read_json_value(path: str | None) -> Any:
    return json.loads(Path(path).read_text(encoding="utf-8")) if path else None


def main(argv: Iterable[str] | None = None) -> int:
    ensure_utf8_stdout()
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    check_parser = sub.add_parser("check")
    check_parser.add_argument("--input", help="UTF-8 file outside the delivery directory; stdin when omitted")
    check_parser.add_argument("--review-scores")
    check_parser.add_argument("--word-limit", type=int)
    check_parser.add_argument("--length-mode", choices=["first-draft", "strict"], default="first-draft")
    check_parser.add_argument("--evidence-cards")
    check_parser.add_argument("--fact-ledger")
    check_parser.add_argument("--style-profile", help="style-profile.v1 JSON for the naturalness and seriousness gate")
    check_parser.add_argument("--task-prompt", help="UTF-8 task prompt file used only for instruction-contamination checks")
    document_parser = sub.add_parser("check-document")
    document_parser.add_argument("--docx", required=True)
    document_parser.add_argument("--new-text", help="UTF-8 file containing only newly written prose; stdin when omitted")
    document_parser.add_argument("--structure", help="semantic-document JSON from standard output or a temporary path")
    document_parser.add_argument("--review-scores")
    document_parser.add_argument("--word-limit", type=int)
    document_parser.add_argument("--length-mode", choices=["first-draft", "strict"], default="first-draft")
    document_parser.add_argument("--evidence-cards")
    document_parser.add_argument("--fact-ledger")
    document_parser.add_argument("--writing-ledger", help="writing-ledger.v1 JSON with method_modules and reasoning protocol")
    document_parser.add_argument("--style-profile", help="style-profile.v1 JSON for the naturalness and seriousness gate")
    document_parser.add_argument("--task-prompt", help="UTF-8 task prompt file used only for instruction-contamination checks")
    verify_parser = sub.add_parser("verify-writing-system")
    verify_parser.add_argument("--template")
    args = parser.parse_args(argv)
    try:
        if args.command == "check":
            result = check_new_prose(_read_text(args.input), _read_json(args.review_scores), args.word_limit, args.length_mode, _read_json_value(args.evidence_cards), _read_json(args.fact_ledger), _read_optional_text(args.task_prompt), _read_json(args.style_profile))
        elif args.command == "check-document":
            result = check_document(Path(args.docx), _read_text(args.new_text), _read_json(args.structure), _read_json(args.review_scores), args.word_limit, args.length_mode, _read_json_value(args.evidence_cards), _read_json(args.fact_ledger), _read_json(args.writing_ledger), _read_optional_text(args.task_prompt), _read_json(args.style_profile))
        else:
            result = verify_social_template(Path(args.template) if args.template else None)
        code = 0 if result.get("gate", "passed") != "blocked" and result.get("passed", True) else 2
    except Exception as exc:
        result, code = {"type": "quality-error", "gate": "blocked", "error": str(exc)}, 2
    json.dump(result, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    return code


if __name__ == "__main__":
    raise SystemExit(main())
