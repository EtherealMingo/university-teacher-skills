#!/usr/bin/env python3
"""Build a clean proposal DOCX from a semantic source structure.

The builder never edits the source. It creates the document in an operating-
system temporary directory and commits only the final DOCX to the delivery path.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import time
from copy import deepcopy
from contextlib import ExitStack
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import render as render_core
from .contracts import SessionState
from .template_schema import (
    InternalRecoveryRequired,
    SESSION_PREFIX,
    VisionOcrRequired,
    _assert_user_material_path,
    _extract_docx,
    _legacy_doc_to_docx,
    _protected_skill_owner,
    _split_table_row,
    _text_table_node,
    extract_structure,
)
from .quality_gates import check_document
from .draft_store import create_draft, freeze_draft, structure_to_markdown, verify_frozen
from .errors import AwesomeBenziError, ensure_utf8_stdout
from .ledgers import complete_modules, integrate_method_modules
from .parity import compare_frozen_to_docx
from .paths import assert_user_path, atomic_commit, safe_output_path, sha256 as core_sha256
from .render import render_docx, validate_render_review
from .style_profile import check_style, extract_style_profile, render_style_brief
from .session import (
    cleanup_session,
    consume_execution_budget,
    normalize_execution_budget,
    read_manifest,
    session_lock,
    transition,
)

SIZE_MAP = {"二号": 22.0, "三号": 16.0, "小三": 15.0, "四号": 14.0, "小四": 12.0, "五号": 10.5, "小五": 9.0}
DEFAULT_PROFILE = {
    "page_width_cm": 21.0,
    "page_height_cm": 29.7,
    "margin_top_cm": 2.54,
    "margin_bottom_cm": 2.54,
    "margin_left_cm": 2.54,
    "margin_right_cm": 2.54,
    "title_font": "黑体",
    "title_size": 22.0,
    "heading1_font": "黑体",
    "heading1_size": 16.0,
    "heading2_font": "仿宋_GB2312",
    "heading2_size": 12.0,
    "heading3_font": "仿宋_GB2312",
    "heading3_size": 12.0,
    "body_font": "仿宋_GB2312",
    "body_size": 12.0,
    "table_font": "宋体",
    "table_size": 10.5,
    "line_spacing": 1.5,
    "line_spacing_rule": "multiple",
    "first_line_chars": 2.0,
}
MISSING_FACT_RE = re.compile(r"【待补：[^】]{1,100}】")
INTERNAL_HEADING_PREFIX = re.compile(r"^(?:[（(][一二三四五六七八九十百]+[）)]|\d+[.、])\s*")
CHINESE_NUMERALS = ("一", "二", "三", "四", "五", "六", "七", "八", "九", "十")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _profile_from_directives(directives: list[str]) -> dict[str, Any]:
    profile = dict(DEFAULT_PROFILE)
    text = chr(10).join(directives)
    page_sizes = {
        "A3": (29.7, 42.0), "A4": (21.0, 29.7), "A5": (14.8, 21.0),
        "B4": (25.0, 35.3), "B5": (17.6, 25.0),
    }
    page = re.search(r"(?:^|[^A-Z])(A3|A4|A5|B4|B5)(?:$|[^0-9A-Z])", text, re.I)
    if page:
        width, height = page_sizes[page.group(1).upper()]
        profile["page_width_cm"] = width
        profile["page_height_cm"] = height
    margin_match = re.search(
        r"(?:上下|纵向)?页边距[^0-9]{0,8}([0-9]+(?:[.][0-9]+)?)[ 	]*(?:厘米|cm)[^。；;]{0,12}(?:左右|横向)?页边距[^0-9]{0,8}([0-9]+(?:[.][0-9]+)?)[ 	]*(?:厘米|cm)",
        text, re.I,
    )
    if margin_match:
        vertical = float(margin_match.group(1))
        horizontal = float(margin_match.group(2))
        if 1.0 <= vertical <= 5.0:
            profile["margin_top_cm"] = profile["margin_bottom_cm"] = vertical
        if 1.0 <= horizontal <= 5.0:
            profile["margin_left_cm"] = profile["margin_right_cm"] = horizontal
    else:
        margin = re.search(r"(?:四边|上下左右)?页边距[^0-9]{0,8}([0-9]+(?:[.][0-9]+)?)[ 	]*(?:厘米|cm)", text, re.I)
        if margin:
            value = float(margin.group(1))
            if 1.0 <= value <= 5.0:
                for key in ("margin_top_cm", "margin_bottom_cm", "margin_left_cm", "margin_right_cm"):
                    profile[key] = value
    body = re.search(r"正文[^。；;]{0,35}(宋体|仿宋(?:_GB2312)?|楷体|微软雅黑)[^。；;]{0,20}(小四|四号|五号|小五)", text)
    if not body:
        body = re.search(r"正文[^。；;]{0,35}(小四|四号|五号|小五)[^。；;]{0,20}(宋体|仿宋(?:_GB2312)?|楷体|微软雅黑)", text)
        if body:
            profile["body_size"] = SIZE_MAP[body.group(1)]
            profile["body_font"] = body.group(2)
    else:
        profile["body_font"] = body.group(1)
        profile["body_size"] = SIZE_MAP[body.group(2)]
    heading = re.search(r"(?:标题|题目)[^。；;]{0,20}(宋体|黑体|仿宋(?:_GB2312)?|楷体|微软雅黑)[^。；;]{0,12}(二号|三号|小三|四号|小四)", text)
    if heading:
        profile["title_font"] = heading.group(1)
        profile["title_size"] = SIZE_MAP[heading.group(2)]
    subheading = re.search(r"小标题[^。；;]{0,35}(宋体|仿宋(?:_GB2312)?|楷体|微软雅黑|黑体)[^。；;]{0,20}(小四|四号|五号|小五)", text)
    if not subheading:
        subheading = re.search(r"小标题[^。；;]{0,35}(小四|四号|五号|小五)[^。；;]{0,20}(宋体|仿宋(?:_GB2312)?|楷体|微软雅黑|黑体)", text)
        if subheading:
            subheading_font, subheading_size = subheading.group(2), SIZE_MAP[subheading.group(1)]
        else:
            subheading_font = subheading_size = None
    else:
        subheading_font, subheading_size = subheading.group(1), SIZE_MAP[subheading.group(2)]
    if subheading_font and subheading_size:
        for level in (2, 3):
            profile[f"heading{level}_font"] = subheading_font
            profile[f"heading{level}_size"] = subheading_size
    else:
        for level in (2, 3):
            profile[f"heading{level}_font"] = profile["body_font"]
            profile[f"heading{level}_size"] = profile["body_size"]
    fixed_spacing = re.search(r"固定值?[^。；;]{0,8}([0-9]+(?:[.][0-9]+)?)[ 	]*磅", text, re.I)
    if fixed_spacing:
        profile["line_spacing"] = float(fixed_spacing.group(1))
        profile["line_spacing_rule"] = "exactly"
    elif re.search(r"单倍[ 	]*行距|1[.]0[ 	]*倍?行距", text):
        profile["line_spacing"] = 1.0
        profile["line_spacing_rule"] = "multiple"
    elif re.search(r"(1[.]5|一点五)[ 	]*倍?行距", text):
        profile["line_spacing"] = 1.5
        profile["line_spacing_rule"] = "multiple"
    elif re.search(r"2(?:[ 	]*[.][ 	]*0)?[ 	]*倍?行距|两倍[ 	]*行距", text):
        profile["line_spacing"] = 2.0
        profile["line_spacing_rule"] = "multiple"
    return profile
def _set_style_font(style: Any, font_name: str, size: float, bold: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), font_name)


def _remove_paragraph_decorations(paragraph_or_style: Any) -> None:
    """Remove decorations that Word/WPS can inherit from built-in title styles."""
    if hasattr(paragraph_or_style, "_p"):
        paragraph_properties = paragraph_or_style._p.get_or_add_pPr()
    else:
        paragraph_properties = paragraph_or_style.element.get_or_add_pPr()
    for tag in ("w:pBdr", "w:shd", "w:outlineLvl"):
        for child in list(paragraph_properties.findall(qn(tag))):
            paragraph_properties.remove(child)


def _configure_styles(document: Document, profile: dict[str, Any]) -> None:
    section = document.sections[0]
    section.page_width = Cm(profile["page_width_cm"])
    section.page_height = Cm(profile["page_height_cm"])
    section.top_margin = Cm(profile["margin_top_cm"])
    section.bottom_margin = Cm(profile["margin_bottom_cm"])
    section.left_margin = Cm(profile["margin_left_cm"])
    section.right_margin = Cm(profile["margin_right_cm"])

    styles = document.styles
    normal = styles["Normal"]
    _set_style_font(normal, profile["body_font"], profile["body_size"])
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if profile.get("line_spacing_rule") == "exactly":
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        normal.paragraph_format.line_spacing = Pt(float(profile["line_spacing"]))
    else:
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        normal.paragraph_format.line_spacing = float(profile["line_spacing"])
    normal.paragraph_format.first_line_indent = Pt(profile["body_size"] * profile["first_line_chars"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    _remove_paragraph_decorations(normal)

    if "Table Text" not in [style.name for style in styles]:
        table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_text = styles["Table Text"]
    _set_style_font(table_text, profile["table_font"], profile["table_size"])
    table_text.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    table_text.paragraph_format.line_spacing = 1.0
    table_text.paragraph_format.first_line_indent = Pt(0)
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)

    if "Instruction" not in [style.name for style in styles]:
        instruction = styles.add_style("Instruction", WD_STYLE_TYPE.PARAGRAPH)
    else:
        instruction = styles["Instruction"]
    _set_style_font(instruction, profile["body_font"], max(profile["body_size"] - 1, 9))
    instruction.font.italic = True
    instruction.paragraph_format.first_line_indent = Pt(0)
    instruction.paragraph_format.line_spacing = 1.0

    if "Missing Fact" not in [style.name for style in styles]:
        missing_fact = styles.add_style("Missing Fact", WD_STYLE_TYPE.CHARACTER)
    else:
        missing_fact = styles["Missing Fact"]
    missing_fact.font.color.rgb = RGBColor(255, 0, 0)


def _split_body(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    parts = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    return parts or ([text.strip()] if text.strip() else [])


def _append_styled_text(paragraph: Any, text: str, bold: bool = False) -> None:
    cursor = 0
    for match in MISSING_FACT_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            if bold:
                run.bold = True
        run = paragraph.add_run(match.group(0))
        run.style = "Missing Fact"
        if bold:
            run.bold = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if bold:
            run.bold = True


def _content_table_node(rows: list[list[str]], anchor: str, index: int) -> dict[str, Any]:
    table_index = max(index, 1)
    node = _text_table_node(rows, table_index)
    node["source_anchor"] = f"{anchor}-t{table_index:03d}"
    node["confidence"] = 1.0
    node["origin"] = "generated"
    return node


def _normalize_table_rows(raw_rows: Any) -> list[list[str]]:
    if not isinstance(raw_rows, list):
        raise ValueError("table rows must be an array")
    rows: list[list[str]] = []
    for raw_row in raw_rows:
        if isinstance(raw_row, str):
            cells = [raw_row]
        elif isinstance(raw_row, (list, tuple)):
            cells = [str(value) for value in raw_row]
        elif isinstance(raw_row, dict):
            cells = [
                str(cell.get("text", ""))
                for cell in raw_row.get("cells", [])
                if isinstance(cell, dict)
            ]
            if not cells:
                cells = [str(raw_row.get("text", ""))]
        else:
            cells = [str(raw_row)]
        rows.append([cell.strip() for cell in cells])
    width = max((len(row) for row in rows), default=0)
    return [row + [""] * (width - len(row)) for row in rows] if width else []


def _number_internal_headings(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    top = 0
    child = 0
    for block in blocks:
        if block.get("kind") != "internal-heading":
            continue
        depth = int(block.get("level") or 1)
        if depth == 2 and top == 0:
            depth = 1
            block["level"] = 1
        text = block.get("text", "").strip()
        if depth == 1:
            top += 1
            child = 0
            prefix = f"（{CHINESE_NUMERALS[top - 1]}）" if top <= len(CHINESE_NUMERALS) else f"（{top}）"
        else:
            child += 1
            prefix = f"{child}."
        if text and not INTERNAL_HEADING_PREFIX.match(text):
            block["text"] = f"{prefix}{text}"
        block["internal_depth"] = depth
    return blocks


def _normalize_block(block: Any, anchor: str, index: int, default_heading_level: int = 2) -> dict[str, Any]:
    if isinstance(block, (int, float, bool)):
        block = str(block)
    if isinstance(block, str):
        return {"source_anchor": f"{anchor}-b{index:03d}", "kind": "paragraph", "text": block.strip(), "confidence": 1.0, "level": None, "images": [], "origin": "generated"}
    if not isinstance(block, dict):
        raise ValueError("structured section blocks must be strings or objects")
    raw_kind = str(block.get("kind", "paragraph"))
    if raw_kind in {"subheading", "sub-subheading"}:
        if "depth" in block:
            depth = int(block["depth"])
        elif raw_kind == "sub-subheading":
            depth = 2
        elif "level" in block:
            depth = 2 if int(block["level"]) >= 3 else 1
        else:
            depth = 1
        kind, level = "internal-heading", max(1, min(depth, 2))
    elif raw_kind == "missing-fact":
        label = str(block.get("label") or block.get("text") or "需补充事实").strip()
        text = label if MISSING_FACT_RE.fullmatch(label) else f"【待补：{label}】"
        return {"source_anchor": f"{anchor}-b{index:03d}", "kind": "paragraph", "text": text, "confidence": 1.0, "level": None, "images": [], "origin": "missing-fact", "missing_fact": True}
    elif raw_kind in {"paragraph", "table-content"}:
        kind, level = "paragraph", None
    elif raw_kind == "list-item":
        list_type = str(block.get("list_type") or block.get("ordered") or "unordered")
        if list_type in {"1", "true", "ordered", "ordered-list"}:
            list_type = "ordered"
        else:
            list_type = "unordered"
        return {
            "source_anchor": f"{anchor}-b{index:03d}", "kind": "list-item",
            "text": str(block.get("text", "")).strip(), "confidence": 1.0,
            "level": None, "images": [], "origin": str(block.get("origin", "generated")),
            "list_type": list_type,
        }
    elif raw_kind == "table":
        rows = _normalize_table_rows(block.get("rows", []))
        if not rows:
            raise ValueError("table blocks must contain at least one row")
        return _content_table_node(rows, anchor, index)
    else:
        raise ValueError(f"unsupported structured section block kind: {raw_kind}")
    return {
        "source_anchor": f"{anchor}-b{index:03d}", "kind": kind,
        "text": str(block.get("text", "")).strip(), "confidence": 1.0,
        "level": level, "images": [], "origin": str(block.get("origin", "generated")),
    }


def _expand_content_blocks(value: Any) -> list[Any]:
    if not isinstance(value, list):
        return value
    expanded: list[Any] = []
    for item in value:
        if isinstance(item, dict) and str(item.get("kind")) in {"ordered-list", "unordered-list"}:
            list_type = "ordered" if str(item.get("kind")) == "ordered-list" else "unordered"
            items = item.get("items")
            if not isinstance(items, list):
                raise ValueError(f"{item.get('kind')} blocks require an items array")
            for list_item in items:
                expanded.append({
                    "kind": "list-item",
                    "text": list_item.get("text") if isinstance(list_item, dict) else list_item,
                    "list_type": list_type,
                    "origin": item.get("origin", "generated"),
                })
        else:
            expanded.append(item)
    return expanded


def _section_blocks(value: Any, anchor: str, default_heading_level: int = 2) -> list[dict[str, Any]]:
    if isinstance(value, list):
        expanded = _expand_content_blocks(value)
        return _number_internal_headings([
            _normalize_block(block, anchor, index, default_heading_level)
            for index, block in enumerate(expanded, start=1)
            if not isinstance(block, str) or block.strip()
        ])
    if not isinstance(value, str):
        raise ValueError("content.sections values must be strings or arrays of structured blocks")
    blocks: list[dict[str, Any]] = []
    buffer: list[str] = []
    heading_re = re.compile(r"^[ 	]*(#{2,6})[ 	]+(.+?)[ 	]*$")
    list_item_re = re.compile(r"^[ 	]*(?:[-+*]|[0-9]+[.)])[ 	]+(.+)$")
    ordered_item_re = re.compile(r"^[ 	]*[0-9]+[.)][ 	]+")

    def flush() -> None:
        if not buffer:
            return
        text = chr(10).join(buffer).strip()
        if text:
            blocks.append(_normalize_block(text, anchor, len(blocks) + 1, default_heading_level))
        buffer.clear()

    lines = value.replace(chr(13)+chr(10), chr(10)).replace(chr(13), chr(10)).split(chr(10))
    cursor = 0
    while cursor < len(lines):
        line = lines[cursor]
        if not line.strip():
            flush()
            cursor += 1
            continue
        stripped = line.strip()
        heading = heading_re.match(stripped)
        if heading:
            flush()
            depth = 1 if len(heading.group(1)) <= 3 else 2
            blocks.append(_normalize_block({"kind": "subheading", "depth": depth, "text": heading.group(2)}, anchor, len(blocks) + 1))
            cursor += 1
            continue
        list_match = list_item_re.match(stripped)
        if list_match:
            flush()
            list_type = "ordered" if ordered_item_re.match(stripped) else "unordered"
            blocks.append(_normalize_block({"kind": "list-item", "text": list_match.group(1).strip(), "list_type": list_type}, anchor, len(blocks) + 1))
            cursor += 1
            continue
        if stripped.startswith("|") and cursor + 1 < len(lines):
            table_lines = [stripped]
            look = cursor + 1
            while look < len(lines) and lines[look].strip().startswith("|"):
                table_lines.append(lines[look].strip())
                look += 1
            parsed_rows = [_split_table_row(row) for row in table_lines]
            if len(parsed_rows) >= 2 and all(
                re.fullmatch(r":?-{3,}:?", cell.replace(" ", ""))
                for cell in parsed_rows[1]
            ):
                flush()
                blocks.append(_content_table_node(parsed_rows, anchor, len(blocks) + 1))
                cursor = look
                continue
        buffer.append(stripped)
        cursor += 1
    flush()
    return _number_internal_headings(blocks)
def _walk_text_nodes(nodes: list[dict[str, Any]]) -> Iterable[dict[str, Any]]:
    for node in nodes:
        if node["kind"] == "table":
            for row in node.get("rows", []):
                for cell in row.get("cells", []):
                    yield from _walk_text_nodes(cell.get("paragraphs", []))
        else:
            yield node


SHORT_FIELD_SECTION = re.compile(r"项目名称|摘要|关键词|基本信息|申请人信息|负责人|经费|预算|签字|承诺|推荐意见")


def _generated_quality_text(structure: dict[str, Any]) -> str:
    """Join only substantive agent-generated prose for the new-prose quality gate.

    Short form fields (project name, abstract, keywords, budget, sign-off, etc.)
    are checked separately by `style-gate` and the section-aware content gate.
    They must not participate in paragraph-density statistics.
    """
    parts: list[str] = []
    current_heading = ""
    for node in structure.get("nodes", []):
        if node.get("kind") == "heading":
            current_heading = str(node.get("text", "")).strip()
            continue
        if node.get("kind") == "table":
            if SHORT_FIELD_SECTION.search(current_heading):
                continue
            for row in node.get("rows", []):
                for cell in row.get("cells", []):
                    for paragraph in cell.get("paragraphs", []):
                        if (
                            paragraph.get("origin") == "generated"
                            and not paragraph.get("missing_fact")
                            and str(paragraph.get("text", "")).strip()
                        ):
                            parts.append(str(paragraph["text"]).strip())
            continue
        if SHORT_FIELD_SECTION.search(current_heading):
            continue
        if (
            node.get("kind") in {"paragraph", "list-item"}
            and node.get("origin") == "generated"
            and not node.get("missing_fact")
            and str(node.get("text", "")).strip()
        ):
            parts.append(str(node["text"]).strip())
    return chr(10)+chr(10).join(parts)
def _generated_short_field_text(structure: dict[str, Any]) -> str:
    parts: list[str] = []
    current_heading = ""
    for node in structure.get("nodes", []):
        if node.get("kind") == "heading":
            current_heading = str(node.get("text", "")).strip()
            continue
        if not SHORT_FIELD_SECTION.search(current_heading):
            continue
        if node.get("kind") == "table":
            for row in node.get("rows", []):
                for cell in row.get("cells", []):
                    for paragraph in cell.get("paragraphs", []):
                        if paragraph.get("origin") == "generated" and str(paragraph.get("text", "")).strip():
                            parts.append(str(paragraph["text"]).strip())
            continue
        if (
            node.get("kind") in {"paragraph", "list-item"}
            and node.get("origin") == "generated"
            and not node.get("missing_fact")
            and str(node.get("text", "")).strip()
        ):
            parts.append(str(node["text"]).strip())
    return chr(10)+chr(10).join(parts)


def _apply_content(structure: dict[str, Any], content: dict[str, Any]) -> dict[str, int]:
    nodes = structure["nodes"]
    text_nodes = list(_walk_text_nodes(nodes))
    replacements = content.get("replacements", {}) or {}
    if not isinstance(replacements, dict):
        raise ValueError("content.replacements must be an object")
    hits = {str(key): 0 for key in replacements}
    for node in text_nodes:
        anchor = node["source_anchor"]
        if anchor in replacements:
            replacement = _normalize_block(replacements[anchor], anchor, 1)
            for key in ("text", "level", "missing_fact", "images"):
                if key not in replacement:
                    node.pop(key, None)
            node.update({key: value for key, value in replacement.items() if key != "source_anchor"})
            if replacement["kind"] == "paragraph" and node["kind"] in {"heading", "document-title", "internal-heading"}:
                node["kind"] = node["kind"]
            else:
                node["kind"] = replacement["kind"]
            hits[anchor] += 1
        for old, new in replacements.items():
            if old == anchor or not old:
                continue
            if old in node.get("text", ""):
                replacement = _normalize_block(new, anchor, 1)
                if replacement["kind"] not in {"paragraph", "missing-fact"}:
                    raise ValueError("text replacements must be prose or missing-fact placeholders")
                node["text"] = node["text"].replace(str(old), replacement["text"])
                if replacement.get("missing_fact"):
                    node["missing_fact"] = True
                if replacement["kind"] == "paragraph" and node["kind"] in {"heading", "document-title", "internal-heading"}:
                    pass
                else:
                    node["kind"] = replacement["kind"]
                hits[str(old)] += 1

    title = content.get("title")
    if title:
        target = next((node for node in text_nodes if node["kind"] == "document-title"), None)
        if target is None:
            nodes.insert(0, {"source_anchor": "generated-title", "kind": "document-title", "text": str(title), "confidence": 1.0, "level": None, "images": []})
        else:
            target["text"] = str(title)

    sections = content.get("sections", {}) or {}
    if not isinstance(sections, dict):
        raise ValueError("content.sections must be an object")
    for heading, body in sections.items():
        heading_index = next((index for index, node in enumerate(nodes) if node["kind"] in {"heading", "document-title"} and node.get("text") == heading), None)
        if heading_index is None:
            raise ValueError(f"section heading not found: {heading}")
        replacement_index = None
        for index in range(heading_index + 1, min(heading_index + 5, len(nodes))):
            if nodes[index]["kind"] in {"heading", "document-title", "table"}:
                break
            if nodes[index]["kind"] in {"placeholder", "empty"}:
                replacement_index = index
                break
        parent_level = int(nodes[heading_index].get("level") or 1)
        replacements_for_section = _section_blocks(body, f"section-{heading_index:04d}", min(parent_level + 1, 3))
        if not replacements_for_section:
            replacements_for_section = [_normalize_block({"kind": "missing-fact", "label": heading}, f"section-{heading_index:04d}", 1)]
        if replacement_index is None:
            nodes[heading_index + 1:heading_index + 1] = replacements_for_section
        else:
            replacements_for_section[0]["source_anchor"] = nodes[replacement_index]["source_anchor"]
            nodes[replacement_index:replacement_index + 1] = replacements_for_section

    for node in _walk_text_nodes(nodes):
        if node.get("kind") != "placeholder":
            continue
        label = re.sub(r"[【】_\s]", "", str(node.get("text", "")))
        label = re.sub(r"^(?:请|在此)?(?:填写|输入|待补)[:：]?", "", label) or "正文"
        node.update(
            {
                "kind": "paragraph",
                "text": f"【待补：{label[:100]}】",
                "origin": "missing-fact",
                "missing_fact": True,
            }
        )

    missing = [key for key, count in hits.items() if count == 0]
    if missing:
        raise ValueError(f"replacement targets not found: {missing}")
    return hits


def _docx_image_blobs(path: Path) -> dict[str, bytes]:
    document = Document(str(path))
    blobs: dict[str, bytes] = {}
    for rel_id, relationship in document.part.rels.items():
        target = getattr(relationship, "target_part", None)
        if target is not None and str(getattr(target, "content_type", "")).startswith("image/"):
            blobs[rel_id] = target.blob
    return blobs


def _pdf_image_blobs(path: Path) -> dict[str, bytes]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return {}
    blobs: dict[str, bytes] = {}
    reader = PdfReader(str(path))
    for page_index, page in enumerate(reader.pages):
        try:
            images = page.images
        except Exception:
            continue
        for image_index, image in enumerate(images):
            try:
                blobs[f"pdf:{page_index}:{image_index}"] = image.data
            except Exception:
                continue
    return blobs


def _add_images(paragraph: Any, refs: list[dict[str, Any]], image_blobs: dict[str, bytes], max_width_cm: float) -> int:
    inserted = 0
    for item in refs:
        blob = image_blobs.get(item.get("image_ref"))
        if not blob:
            continue
        run = paragraph.add_run()
        try:
            run.add_picture(BytesIO(blob), width=Cm(max_width_cm))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inserted += 1
        except Exception:
            continue
    return inserted


def _add_text_paragraph(document: Document, node: dict[str, Any], image_blobs: dict[str, bytes], max_width_cm: float, profile: dict[str, Any]) -> list[Any]:
    kind = node["kind"]
    text_parts = _split_body(node.get("text", ""))
    paragraphs: list[Any] = []
    if kind == "document-title":
        paragraph = document.add_paragraph(style="Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        _remove_paragraph_decorations(paragraph)
        _append_styled_text(paragraph, text_parts[0] if text_parts else "", bold=True)
        paragraphs.append(paragraph)
    elif kind == "heading":
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        _remove_paragraph_decorations(paragraph)
        _append_styled_text(paragraph, text_parts[0] if text_parts else "", bold=True)
        paragraphs.append(paragraph)
    elif kind == "internal-heading":
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(profile["body_size"] * 2.0) if int(node.get("internal_depth") or node.get("level") or 1) == 2 else Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        _remove_paragraph_decorations(paragraph)
        _append_styled_text(paragraph, text_parts[0] if text_parts else "", bold=True)
        paragraphs.append(paragraph)
    else:
        style = "Instruction" if kind == "instruction" else "Normal"
        if kind == "list-item":
            style = "List Number" if node.get("list_type") == "ordered" else "List Bullet"
        for part in text_parts:
            paragraph = document.add_paragraph(style=style)
            _append_styled_text(paragraph, part)
            if kind == "list-item":
                paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraphs.append(paragraph)
    if not paragraphs and node.get("images"):
        paragraphs.append(document.add_paragraph(style="Normal"))
    if paragraphs:
        _add_images(paragraphs[-1], node.get("images", []), image_blobs, max_width_cm)
    return paragraphs


def _clear_cell(cell: Any) -> None:
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    first = cell.paragraphs[0]
    for child in list(first._p):
        if not child.tag.endswith("}pPr"):
            first._p.remove(child)
    first.style = "Table Text"


def _fill_cell(cell: Any, paragraph_nodes: list[dict[str, Any]], image_blobs: dict[str, bytes], max_width_cm: float) -> None:
    _clear_cell(cell)
    target = cell.paragraphs[0]
    wrote = False
    for node in paragraph_nodes:
        parts = _split_body(node.get("text", ""))
        for part in parts:
            paragraph = target if not wrote else cell.add_paragraph(style="Table Text")
            paragraph.style = "Table Text"
            _append_styled_text(paragraph, part)
            wrote = True
        if node.get("images"):
            paragraph = target if not wrote else cell.add_paragraph(style="Table Text")
            _add_images(paragraph, node["images"], image_blobs, max_width_cm)
            wrote = True


def _set_table_borders(table: Any) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")


def _add_table(document: Document, node: dict[str, Any], image_blobs: dict[str, bytes], max_width_cm: float) -> Any:
    rows_data = node.get("rows", [])
    columns = max(int(node.get("columns") or 1), 1)
    table = document.add_table(rows=max(len(rows_data), 1), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    _set_table_borders(table)
    vertical_starts: dict[tuple[int, int], Any] = {}
    for row_index, row_data in enumerate(rows_data):
        for cell_data in row_data.get("cells", []):
            start = min(int(cell_data.get("start_column", 0)), columns - 1)
            span = max(int(cell_data.get("grid_span", 1)), 1)
            end = min(start + span - 1, columns - 1)
            cell = table.cell(row_index, start)
            if end > start:
                cell = cell.merge(table.cell(row_index, end))
            merge_state = cell_data.get("vertical_merge")
            key = (start, span)
            if merge_state == "restart":
                vertical_starts[key] = cell
            elif merge_state == "continue" and key in vertical_starts:
                cell = vertical_starts[key].merge(cell)
                continue
            _fill_cell(cell, cell_data.get("paragraphs", []), image_blobs, max_width_cm)
    return table


def _safe_output(output_dir: Path, stem: str, requested_name: str | None) -> Path:
    skill_root = Path(__file__).resolve().parents[2]
    try:
        output_dir = assert_user_path(output_dir, skill_root)
    except AwesomeBenziError:
        raise
    owner = _protected_skill_owner(output_dir)
    if owner is not None:
        raise AwesomeBenziError("E_PATH_PROTECTED", "delivery", f"awesome-benzi runtime package is read-only and cannot be a delivery directory: {output_dir}", path=str(output_dir), recoverability="user-input")
    if not output_dir.is_dir():
        raise AwesomeBenziError("E_PATH_NOT_FOUND", "delivery", "delivery directory must already exist", path=str(output_dir), recoverability="user-input")
    name = requested_name or f"{stem}-完成稿.docx"
    if Path(name).suffix.lower() != ".docx" or Path(name).name != name:
        raise ValueError("output name must be a DOCX filename without a directory")
    candidate = output_dir / name
    if not candidate.exists():
        return candidate
    for index in range(2, 1000):
        candidate = output_dir / f"{Path(name).stem}-{index:02d}.docx"
        if not candidate.exists():
            return candidate
    raise RuntimeError("could not allocate a non-conflicting output name")


def _read_content(path: str | None) -> dict[str, Any]:
    if not path:
        return {}
    value = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    if not isinstance(value, dict):
        raise ValueError("content mapping must be a JSON object")
    return value


def _load_prepared_session(path: Path) -> tuple[dict[str, Any], Path]:
    requested = path.expanduser().resolve()
    manifest_path = requested / "manifest.json" if requested.is_dir() else requested
    session_dir = manifest_path.parent.resolve()
    temp_root = Path(tempfile.gettempdir()).resolve()
    if session_dir.parent != temp_root or not session_dir.name.startswith(SESSION_PREFIX):
        raise ValueError("prepared session must be an awesome-benzi directory under the operating-system temporary root")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("type") != "prepared-session.v1":
        raise ValueError("prepared session manifest type is invalid")
    for key in ("semantic_structure_path", "writing_ledger_path"):
        candidate = Path(manifest[key]).resolve()
        try:
            candidate.relative_to(session_dir)
        except ValueError:
            raise ValueError(f"prepared session {key} escapes its session directory")
    return manifest, session_dir


def _heading_body_status(nodes: list[dict[str, Any]], heading_text: str) -> str:
    heading_index = next((index for index, node in enumerate(nodes) if node.get("kind") == "heading" and node.get("text", "").strip() == heading_text), None)
    if heading_index is None:
        return "pending"
    level = int(nodes[heading_index].get("level") or 1)
    body_texts: list[str] = []
    for node in nodes[heading_index + 1:]:
        if node.get("kind") == "heading":
            if int(node.get("level") or 1) <= level:
                break
            body_texts.append(node.get("text", ""))
            continue
        if node.get("kind") == "table":
            body_texts.extend(part.get("text", "") for row in node.get("rows", []) for cell in row.get("cells", []) for part in cell.get("paragraphs", []))
        elif node.get("kind") not in {"empty", "instruction"}:
            body_texts.append(node.get("text", ""))
    text = "\n".join(body_texts).strip()
    if MISSING_FACT_RE.search(text):
        return "placeholder-complete"
    return "complete" if text else "pending"


def _compact_characters(text: str) -> int:
    return len(re.sub(r"\s+", "", text))


def _section_argument_tree(nodes: list[dict[str, Any]], section: dict[str, Any]) -> dict[str, Any]:
    heading_text = section.get("heading", "").strip()
    heading_index = next((index for index, node in enumerate(nodes) if node.get("kind") == "heading" and node.get("text", "").strip() == heading_text), None)
    existing = deepcopy(section.get("argument_tree") or {})
    for key, default in (
        ("central_thesis", ""), ("subproblems", []), ("facts", []), ("evidence", []),
        ("mechanisms", []), ("methods", []), ("boundary", ""), ("conclusion", ""), ("units", []),
    ):
        existing.setdefault(key, default)
    if heading_index is None:
        existing.update({"actual_character_count": 0, "paragraph_count": 0, "status": "pending"})
        return existing
    section_level = int(nodes[heading_index].get("level") or 1)
    section_nodes: list[dict[str, Any]] = []
    for node in nodes[heading_index + 1:]:
        if node.get("kind") == "heading" and int(node.get("level") or 1) <= section_level:
            break
        section_nodes.append(node)

    prior_units = existing.get("units", []) if isinstance(existing.get("units"), list) else []
    units: list[dict[str, Any]] = []
    preface: list[str] = []
    current: dict[str, Any] | None = None
    all_body_parts: list[str] = []
    paragraph_count = 0

    def table_text(node: dict[str, Any]) -> list[str]:
        return [
            part.get("text", "").strip()
            for row in node.get("rows", [])
            for cell in row.get("cells", [])
            for part in cell.get("paragraphs", [])
            if part.get("text", "").strip()
        ]

    for node in section_nodes:
        if node.get("kind") == "internal-heading":
            index = len(units)
            prior = deepcopy(prior_units[index]) if index < len(prior_units) and isinstance(prior_units[index], dict) else {}
            current = {
                "unit_id": prior.get("unit_id") or node.get("source_anchor") or f"unit-{index + 1}",
                "title": node.get("text", "").strip(),
                "depth": int(node.get("internal_depth") or node.get("level") or 1),
                "subproblem": prior.get("subproblem", ""),
                "facts": prior.get("facts", []), "evidence": prior.get("evidence", []),
                "mechanism": prior.get("mechanism", ""), "method": prior.get("method", ""),
                "boundary": prior.get("boundary", ""), "conclusion": prior.get("conclusion", ""),
                "body_character_count": 0, "paragraph_count": 0, "status": "pending",
            }
            units.append(current)
            continue
        parts = table_text(node) if node.get("kind") == "table" else ([node.get("text", "").strip()] if node.get("kind") not in {"empty", "instruction", "placeholder"} and node.get("text", "").strip() else [])
        for part in parts:
            all_body_parts.append(part)
            paragraph_count += 1
            if current is None:
                preface.append(part)
            else:
                current["body_character_count"] += _compact_characters(part)
                current["paragraph_count"] += 1

    for index, unit in enumerate(units):
        has_child = False
        for candidate in units[index + 1:]:
            if candidate["depth"] <= unit["depth"]:
                break
            has_child = True
            break
        unit["has_child"] = has_child
        unit["status"] = "complete" if unit["body_character_count"] > 0 or has_child else "pending"
    actual_count = sum(_compact_characters(part) for part in all_body_parts)
    if not existing.get("central_thesis") and all_body_parts:
        existing["central_thesis"] = all_body_parts[0][:180]
    if not existing.get("conclusion") and all_body_parts:
        existing["conclusion"] = all_body_parts[-1][-180:]
    existing.update({
        "units": units,
        "preface_character_count": sum(_compact_characters(part) for part in preface),
        "actual_character_count": actual_count,
        "paragraph_count": paragraph_count,
        "status": "complete" if actual_count > 0 and all(unit["status"] == "complete" for unit in units) else "pending",
    })
    return existing


def _finalize_writing_ledger(
    ledger: dict[str, Any] | None,
    structure: dict[str, Any],
    content: dict[str, Any],
    review_scores: dict[str, float] | None,
) -> dict[str, Any] | None:
    if ledger is None:
        return None
    ledger = deepcopy(content.get("writing_ledger") or ledger)
    completed_modules = {int(value) for value in content.get("completed_modules", []) if str(value).isdigit()}
    for item in ledger.get("modules", []):
        if int(item.get("module", 0)) in completed_modules:
            item["status"] = "complete"
    for item in ledger.get("sections", []):
        item["status"] = _heading_body_status(structure.get("nodes", []), item.get("heading", ""))
        tree = _section_argument_tree(structure.get("nodes", []), item)
        item["argument_tree"] = tree
        section_heading = item.get("heading", "").strip()
        section_index = next((index for index, node in enumerate(structure.get("nodes", [])) if node.get("kind") == "heading" and node.get("text", "").strip() == section_heading), None)
        modified = False
        if section_index is not None:
            level = int(structure["nodes"][section_index].get("level") or 1)
            for node in structure["nodes"][section_index + 1:]:
                if node.get("kind") == "heading" and int(node.get("level") or 1) <= level:
                    break
                modified = modified or node.get("origin") == "generated"
        item["modified"] = modified
        budget = int(item.get("word_budget") or 0)
        tasks = item.get("argument_tasks", []) if isinstance(item.get("argument_tasks"), list) else []
        item["requires_internal_structure"] = bool(
            modified
            and not item.get("exempt_from_internal_structure", False)
            and (budget >= 600 or len(tasks) >= 2 or int(tree.get("actual_character_count", 0)) >= 600)
        )
    required_scores = {"value", "design", "method", "innovation", "foundation", "feasibility"}
    if review_scores and required_scores.issubset(review_scores):
        ledger.setdefault("review", {})["status"] = "complete"
        ledger["review"]["dimensions"] = {key: float(review_scores[key]) for key in sorted(required_scores)}
    if "m10_targeted_revision_count" in content:
        ledger["m10_targeted_revision_count"] = int(content["m10_targeted_revision_count"])
    if "m10_affected_modules" in content:
        ledger["m10_affected_modules"] = [
            f"M{value}" if str(value).isdigit() else str(value)
            for value in content.get("m10_affected_modules", [])
        ]
    if isinstance(content.get("reasoning_plan"), dict):
        ledger["reasoning_plan"] = deepcopy(content["reasoning_plan"])
    if isinstance(content.get("verification_summary"), dict):
        ledger["verification_summary"] = deepcopy(content["verification_summary"])
    ledger["unresolved_placeholders"] = len(MISSING_FACT_RE.findall("\n".join(node.get("text", "") for node in _walk_text_nodes(structure.get("nodes", [])))))
    ledger = integrate_method_modules(ledger)
    reasoning_sections = {
        str(item.get("section_id")): item
        for item in ledger.get("reasoning_plan", {}).get("sections", [])
        if isinstance(item, dict)
    }
    for section in ledger.get("sections", []):
        reasoning_section = reasoning_sections.get(str(section.get("section_id")))
        if reasoning_section is None:
            continue
        tree = section.get("argument_tree", {})
        reasoning_section["central_proposition"] = str(tree.get("central_thesis", "")).strip()
        reasoning_section["subquestions"] = list(tree.get("subproblems", []))
        reasoning_section["evidence_ids"] = list(tree.get("evidence", []))
        reasoning_section["unknowns"] = list(section.get("missing_items", []))
        reasoning_section["method"] = "；".join(str(value) for value in tree.get("methods", []) if value)
        reasoning_section["boundary"] = str(tree.get("boundary", "")).strip()
        reasoning_section["expected_conclusion"] = str(tree.get("conclusion", "")).strip()
        reasoning_section["word_budget"] = section.get("word_budget")
        reasoning_section["status"] = (
            "complete"
            if section.get("status") in {"complete", "placeholder-complete", "source-present"}
            else "pending"
        )
    verification_questions = [
        question
        for section in reasoning_sections.values()
        for question in section.get("verification_questions", [])
        if isinstance(question, dict)
    ]
    summary = ledger.get("verification_summary", {})
    statuses = [str(question.get("status", "pending")) for question in verification_questions]
    terminal_statuses = {"supported", "limited", "unsupported", "unknown", "verified"}
    summary.update(
        {
            "total_questions": len(verification_questions),
            "verified": statuses.count("verified") + statuses.count("supported"),
            "supported": statuses.count("supported"),
            "limited": statuses.count("limited"),
            "unsupported": statuses.count("unsupported"),
            "unknown": statuses.count("unknown"),
            "revision_rounds": int(ledger.get("m10_targeted_revision_count", 0) or 0),
        }
    )
    if verification_questions and all(status in terminal_statuses for status in statuses):
        summary["status"] = "complete"
    elif not verification_questions and ledger.get("review", {}).get("status") == "complete":
        summary["status"] = "complete"
        summary["early_stop_reason"] = "no-structured-high-risk-claims"
    ledger["verification_summary"] = summary
    ledger = complete_modules(
        ledger,
        sorted(completed_modules),
        int(content.get("draft_revision", 1)),
        module_results=content.get("module_results"),
    )
    return ledger


def build_docx(
    template: Path | None,
    output_dir: Path,
    content: dict[str, Any] | None = None,
    output_name: str | None = None,
    new_text: str | None = None,
    review_scores: dict[str, float] | None = None,
    word_limit: int | None = None,
    length_mode: str = "first-draft",
    evidence_cards: list[dict[str, Any]] | None = None,
    fact_ledger: dict[str, Any] | None = None,
    prepared_session: Path | None = None,
    preserve_prepared_session: bool = False,
) -> dict[str, Any]:
    prepared_manifest: dict[str, Any] | None = None
    cleanup_session: Path | None = None
    initial_writing_ledger: dict[str, Any] | None = None
    if prepared_session is not None:
        prepared_manifest, cleanup_session = _load_prepared_session(prepared_session)
        template = Path(prepared_manifest["source_template_path"])
        initial_writing_ledger = json.loads(Path(prepared_manifest["writing_ledger_path"]).read_text(encoding="utf-8"))
    if template is None:
        raise ValueError("template or prepared session is required")
    template = _assert_user_material_path(template)
    if not template.is_file():
        raise AwesomeBenziError("E_PATH_NOT_FOUND", "build", f"template is not a file: {template}", path=str(template), recoverability="user-input")
    if prepared_manifest is not None:
        if sha256(template) != prepared_manifest.get("source_hash"):
            raise AwesomeBenziError("E_SOURCE_CHANGED", "build", "source template changed after semantic preparation", path=str(template), recoverability="fatal", user_action_required=True, suggested_action="Prepare a new session from the current template.")
    source_hash = prepared_manifest.get("source_hash") if prepared_manifest else sha256(template)
    destination = _safe_output(output_dir, template.stem, output_name)
    delivered = False
    content = content or {}
    recovered_source_text = content.get("recovered_source_text")
    if recovered_source_text is not None and not isinstance(recovered_source_text, str):
        raise ValueError("content.recovered_source_text must be a string")
    with ExitStack() as stack:
        if cleanup_session is not None and not preserve_prepared_session:
            stack.callback(shutil.rmtree, cleanup_session, True)
        raw = stack.enter_context(tempfile.TemporaryDirectory(prefix="awesome-benzi-build-"))
        temp_root = Path(raw)
        semantic_source = Path(prepared_manifest["semantic_source_path"]) if prepared_manifest else template
        if prepared_manifest:
            structure = json.loads(Path(prepared_manifest["semantic_structure_path"]).read_text(encoding="utf-8"))
        elif template.suffix.lower() == ".doc":
            semantic_source = temp_root / f"{template.stem}.docx"
            recovery_backend = _legacy_doc_to_docx(template, semantic_source)
            structure = _extract_docx(semantic_source)
            structure["metadata"].update({"source_path": str(template), "source_format": ".doc", "input_recovery": {"status": "recovered", "backend": recovery_backend}})
        else:
            structure = extract_structure(template, recovered_source_text)
        if structure["semantic_structure_confidence"] < 0.75:
            raise ValueError("template semantic structure confidence is below 0.75; request a clearer template once")
        structure = deepcopy(structure)
        replacement_hits = _apply_content(structure, content)
        structure["metadata"]["table_count"] = sum(1 for node in structure["nodes"] if node.get("kind") == "table")
        structure["metadata"]["heading_count"] = sum(1 for node in structure["nodes"] if node.get("kind") == "heading")
        structure["metadata"]["title_count"] = sum(1 for node in structure["nodes"] if node.get("kind") == "document-title")
        writing_ledger = _finalize_writing_ledger(initial_writing_ledger or content.get("writing_ledger"), structure, content, review_scores)
        supplied_directives = content.get("style_directives", []) or []
        if not isinstance(supplied_directives, list) or not all(isinstance(item, str) for item in supplied_directives):
            raise ValueError("content.style_directives must be an array of strings")
        structure["style_directives"] = list(dict.fromkeys(supplied_directives + structure.get("style_directives", [])))
        profile = _profile_from_directives(structure.get("style_directives", []))
        image_blobs = _docx_image_blobs(semantic_source) if semantic_source.suffix.lower() == ".docx" else _pdf_image_blobs(template) if template.suffix.lower() == ".pdf" else {}

        document = Document()
        _configure_styles(document, profile)
        usable_width = profile["page_width_cm"] - profile["margin_left_cm"] - profile["margin_right_cm"]
        inserted_images = 0
        for node in structure["nodes"]:
            if node["kind"] == "table":
                _add_table(document, node, image_blobs, max(usable_width - 0.5, 5.0))
            elif node["kind"] == "image":
                paragraph = document.add_paragraph(style="Normal")
                inserted_images += _add_images(paragraph, node.get("images", []), image_blobs, max(usable_width - 0.5, 5.0))
            elif node["kind"] != "empty" or node.get("images"):
                paragraphs = _add_text_paragraph(document, node, image_blobs, max(usable_width - 0.5, 5.0), profile)
                inserted_images += sum(1 for paragraph in paragraphs if paragraph._p.xpath(".//a:blip"))

        staged = temp_root / destination.name
        document.save(str(staged))
        expected_structure = deepcopy(structure)
        expected_structure["style_profile"] = profile
        generated_text = _generated_quality_text(structure)
        quality_text = new_text if new_text is not None else generated_text
        quality_result = check_document(
            staged,
            quality_text,
            expected_structure,
            review_scores,
            word_limit,
            length_mode,
            evidence_cards,
            fact_ledger,
            writing_ledger,
            prepared_manifest.get("task_prompt", "") if prepared_manifest else str(content.get("task_prompt", "")),
            content.get("style_profile"),
        )
        short_field_text = _generated_short_field_text(structure)
        if short_field_text:
            short_style = check_style(short_field_text, content.get("style_profile"))
            quality_result["content"].setdefault("style", {}).setdefault("short_fields", short_style)
            for item in short_style.get("issues", []):
                if item.get("level") == "blocked":
                    quality_result["content"].setdefault("issues", []).append(item)
            content_issues = quality_result["content"].get("issues", [])
            content_hard = [item for item in content_issues if item.get("level") == "blocked"]
            content_soft = [item for item in content_issues if item.get("level") == "needs-review"]
            if content_hard:
                quality_result["content"]["gate"] = "blocked"
            elif content_soft and quality_result["content"]["gate"] == "passed":
                quality_result["content"]["gate"] = "needs-review"
            if quality_result["content"]["gate"] == "blocked":
                quality_result["gate"] = "blocked"
            elif quality_result["content"]["gate"] == "needs-review" and quality_result["gate"] == "passed":
                quality_result["gate"] = "needs-review"
        if quality_result["gate"] == "blocked":
            issue_kinds = sorted({
                item["kind"]
                for section in (quality_result["format"], quality_result["content"], quality_result["writing_ledger"])
                for item in section.get("issues", [])
                if item.get("level") == "blocked"
            })
            raise ValueError(f"combined final gate blocked delivery: {issue_kinds}")
        try:
            atomic_commit(staged, destination)
            delivered = True
        finally:
            if not delivered and destination.exists():
                destination.unlink()

    return {
        "type": "delivery.v3",
        "artifact_path": str(destination),
        "source_template_path": str(template),
        "source_template_hash": source_hash,
        "output_hash": sha256(destination),
        "output_strategy": "semantic-docx-rebuild",
        "source_format": template.suffix.lower(),
        "output_format": "docx",
        "semantic_structure_confidence": structure["semantic_structure_confidence"],
        "style_directives": structure.get("style_directives", []),
        "applied_style_profile": profile,
        "replacement_hits": replacement_hits,
        "preserved_elements": structure.get("preserved_elements", []),
        "inserted_images": inserted_images,
        "input_recovery": structure.get("metadata", {}).get("input_recovery"),
        "prepared_session": prepared_manifest.get("type") if prepared_manifest else None,
        "extraction_stats": prepared_manifest.get("extraction_stats") if prepared_manifest else None,
        "writing_coverage": quality_result["writing_ledger"].get("coverage"),
        "unresolved_placeholders": len(MISSING_FACT_RE.findall("\n".join(node.get("text", "") for node in _walk_text_nodes(structure.get("nodes", []))))),
        "final_gate": quality_result,
        "process_files_in_delivery_directory": [],
    }


def _prepared_paths(prepared_session: Path) -> tuple[dict[str, Any], Path, dict[str, Any], dict[str, Any]]:
    manifest, session = _load_prepared_session(prepared_session)
    structure = json.loads(Path(manifest["semantic_structure_path"]).read_text(encoding="utf-8"))
    ledger = json.loads(Path(manifest["writing_ledger_path"]).read_text(encoding="utf-8"))
    for name in ("draft", "build", "qa", "qa/pages"):
        (session / name).mkdir(parents=True, exist_ok=True)
    return manifest, session, structure, ledger


def _stage_delivery(
    template: Path | None,
    output_dir: Path,
    content: dict[str, Any] | None = None,
    output_name: str | None = None,
    new_text: str | None = None,
    review_scores: dict[str, float] | None = None,
    word_limit: int | None = None,
    length_mode: str = "first-draft",
    evidence_cards: list[dict[str, Any]] | None = None,
    fact_ledger: dict[str, Any] | None = None,
    prepared_session: Path | None = None,
    budget_seconds: int | float | None = None,
) -> dict[str, Any]:
    stage_started = time.perf_counter()
    if prepared_session is None:
        if template is None:
            raise AwesomeBenziError("E_TEMPLATE_MISSING", "prepare", "template or prepared session is required")
        from .template_schema import prepare_session

        prepared_session = Path(
            prepare_session(
                template,
                budget_seconds=budget_seconds if budget_seconds is not None else 300,
            )["manifest_path"]
        )
    prepared, session, structure, initial_ledger = _prepared_paths(prepared_session)
    raw_budget = prepared.get("execution_budget")
    if budget_seconds is not None and isinstance(raw_budget, dict):
        raw_budget = {**raw_budget, "limit_ms": round(float(budget_seconds) * 1000, 2)}
    execution_budget = normalize_execution_budget(
        raw_budget,
        budget_seconds if budget_seconds is not None else 300,
    )
    backend_attempts = list(prepared.get("backend_attempts", []))
    render_capability_report = render_core.render_capabilities()
    capabilities = {
        **prepared.get("capabilities", {}),
        **render_capability_report,
    }
    render_gate_caps = render_capability_report["render_gate"]
    render_available = bool(render_gate_caps["available"])
    current_state = SessionState(read_manifest(session)["state"])
    if current_state in {SessionState.PREPARED, SessionState.WAITING_FOR_USER}:
        transition(session, SessionState.PLANNED)
        transition(session, SessionState.DRAFTING)
    elif current_state == SessionState.PLANNED:
        transition(session, SessionState.DRAFTING)
    elif current_state == SessionState.FROZEN:
        transition(session, SessionState.DRAFTING)
    elif current_state != SessionState.DRAFTING:
        raise AwesomeBenziError(
            "E_INVALID_STATE_TRANSITION",
            "session",
            f"first-phase build cannot start from {current_state.value}",
            recoverability="internal-invariant",
        )
    template = Path(prepared["source_template_path"])
    if core_sha256(template) != prepared.get("source_hash"):
        raise AwesomeBenziError(
            "E_SOURCE_CHANGED",
            "prepare",
            "source template changed after semantic preparation",
            recoverability="fatal",
            user_action_required=True,
            suggested_action="Prepare a new session from the current template.",
        )
    content = deepcopy(content or {})
    if not content.get("sections") and not content.get("replacements"):
        legacy_body = content.get("blocks") or new_text
        first_heading = next(
            (
                str(node.get("text", "")).strip()
                for node in structure.get("nodes", [])
                if node.get("kind") == "heading" and str(node.get("text", "")).strip()
            ),
            "",
        )
        if legacy_body and first_heading:
            content["sections"] = {first_heading: legacy_body}
    completed = content.get("completed_modules", [])
    preview_structure = deepcopy(structure)
    _apply_content(preview_structure, content)
    writing_ledger = _finalize_writing_ledger(
        content.get("writing_ledger") or initial_ledger,
        preview_structure,
        content,
        review_scores,
    )
    if writing_ledger is not None:
        writing_ledger = complete_modules(
            writing_ledger,
            completed,
            1,
            module_results=content.get("module_results"),
        )
        content["writing_ledger"] = writing_ledger
    markdown = structure_to_markdown(preview_structure)
    draft_manifest = create_draft(session, markdown)
    content["draft_revision"] = draft_manifest["revision"]

    from .quality_gates import check_new_prose, check_writing_ledger

    style_profile = content.get("style_profile") or prepared.get("style_profile")
    if content.get("style_samples"):
        sample_texts = [
            item.get("text") if isinstance(item, dict) else item
            for item in content["style_samples"]
            if isinstance(item, str) or (isinstance(item, dict) and item.get("text"))
        ]
        style_profile = extract_style_profile(sample_texts, prepared.get("task_prompt", ""), [str(item.get("heading", "")) for item in structure.get("nodes", []) if item.get("kind") == "heading"])
    if style_profile is not None:
        content["style_profile"] = style_profile
        content["style_brief"] = render_style_brief(style_profile)
    quality_text = _generated_quality_text(preview_structure)
    prose_result = check_new_prose(
        quality_text,
        review_scores,
        word_limit,
        length_mode,
        evidence_cards,
        fact_ledger,
        prepared.get("task_prompt", ""),
        style_profile,
    )
    short_field_text = _generated_short_field_text(preview_structure)
    if short_field_text:
        short_style = check_style(short_field_text, style_profile)
        prose_result.setdefault("style", {})["short_fields"] = short_style
        for item in short_style.get("issues", []):
            if item.get("level") == "blocked":
                prose_result.setdefault("issues", []).append(item)
        hard = [item for item in prose_result.get("issues", []) if item.get("level") == "blocked"]
        soft = [item for item in prose_result.get("issues", []) if item.get("level") == "needs-review"]
        prose_result["gate"] = "blocked" if hard else "needs-review" if soft else "passed"
    ledger_result = check_writing_ledger(writing_ledger)
    gates = {prose_result["gate"], ledger_result["gate"]}
    content_report = {
        "schema_type": "content-quality",
        "schema_version": 3,
        "gate": "blocked" if "blocked" in gates else "needs-review" if "needs-review" in gates else "passed",
        "revision": draft_manifest["revision"],
        "issues": prose_result.get("issues", []) + ledger_result.get("issues", []),
        "prose": prose_result,
        "writing_ledger": ledger_result,
        "style": prose_result.get("style"),
    }
    if content_report["gate"] == "blocked":
        freeze_draft(session, content_report)
    transition(session, SessionState.CONTENT_PASSED)
    frozen = freeze_draft(session, content_report)
    verify_frozen(frozen)
    transition(session, SessionState.FROZEN)
    content_finished = time.perf_counter()
    consume_execution_budget(execution_budget, content_finished - stage_started)

    legacy_delivery = build_docx(
        template,
        session / "build",
        content,
        output_name="staged.docx",
        new_text=new_text,
        review_scores=review_scores,
        word_limit=word_limit,
        length_mode=length_mode,
        evidence_cards=evidence_cards,
        fact_ledger=fact_ledger,
        prepared_session=prepared_session,
        preserve_prepared_session=True,
    )
    staged = Path(legacy_delivery["artifact_path"])
    transition(session, SessionState.DOCX_STAGED)
    docx_finished = time.perf_counter()
    consume_execution_budget(execution_budget, docx_finished - content_finished)
    locked_text = {
        str(node.get("text", "")).strip()
        for node in structure.get("nodes", [])
        if node.get("kind") == "instruction" and str(node.get("text", "")).strip()
    }
    parity = compare_frozen_to_docx(Path(frozen["draft_path"]), staged, locked_text=locked_text)
    if parity["gate"] == "blocked":
        staged.unlink(missing_ok=True)
        transition(session, SessionState.DRAFTING)
        raise AwesomeBenziError(
            "E_DRAFT_DOCX_MISMATCH",
            "parity",
            "staged DOCX does not match frozen Markdown",
            recoverability="internal-retry",
            details=parity,
        )
    transition(session, SessionState.PARITY_PASSED)
    parity_finished = time.perf_counter()
    consume_execution_budget(execution_budget, parity_finished - docx_finished)
    structure_gate = legacy_delivery["final_gate"]["format"]
    if structure_gate["gate"] == "blocked":
        staged.unlink(missing_ok=True)
        transition(session, SessionState.FROZEN)
        raise AwesomeBenziError(
            "E_DOCX_STRUCTURE_BLOCKED",
            "structure",
            "staged DOCX failed the structure gate",
            recoverability="internal-retry",
            details=structure_gate,
        )
    transition(session, SessionState.STRUCTURE_PASSED)
    structure_finished = time.perf_counter()
    consume_execution_budget(execution_budget, structure_finished - parity_finished)
    if render_available:
        try:
            render_report = render_docx(
                staged,
                session / "qa",
                timeout_seconds=60,
                budget=execution_budget,
                attempts=backend_attempts,
            )
        except Exception:
            transition(
                session,
                SessionState.FROZEN,
                execution_budget=execution_budget,
                backend_attempts=backend_attempts,
            )
            raise
        transition(
            session,
            SessionState.RENDER_REVIEW_PENDING,
            execution_budget=execution_budget,
            backend_attempts=backend_attempts,
        )
        render_finished = time.perf_counter()
    else:
        # Rendering is optional: warn and deliver without page-level review.
        render_report = {
            "schema_type": "render-report",
            "schema_version": 1,
            "gate": "skipped",
            "status": "renderer-unavailable",
            "missing": render_gate_caps["missing"],
            "message": (
                "LibreOffice 或 Poppler 不可用：已跳过机械渲染与逐页视觉复核，"
                "一致性与结构检查通过后直接交付 DOCX。"
            ),
        }
        render_finished = structure_finished
    stage_timings_ms = {
        "prepare": float(prepared.get("performance", {}).get("prepare_ms", 0.0)),
        "plan_content_freeze": round((content_finished - stage_started) * 1000, 2),
        "docx_assembly": round((docx_finished - content_finished) * 1000, 2),
        "parity": round((parity_finished - docx_finished) * 1000, 2),
        "structure": round((structure_finished - parity_finished) * 1000, 2),
        "render": round((render_finished - structure_finished) * 1000, 2),
        "first_phase_total": round((render_finished - stage_started) * 1000, 2),
    }
    target = safe_output_path(output_dir, template.stem, output_name, Path(__file__).resolve().parents[2])
    if not render_available:
        # Direct delivery: consistency and structure gates passed, render review
        # was skipped, so there is no page review to wait for. The session is
        # cleaned by the caller after the exclusive lock is released (Windows
        # keeps the lock file open while the lock is held).
        output_hash = atomic_commit(staged, target)
        transition(session, SessionState.DELIVERED)
        return {
            **legacy_delivery,
            "schema_type": "delivery",
            "schema_version": 3,
            "type": "delivery.v3",
            "status": "delivered",
            "artifact_path": str(target),
            "resume_session": None,
            "draft_sha256": frozen["draft_sha256"],
            "output_hash": output_hash,
            "parity_gate": parity,
            "structure_gate": structure_gate,
            "render_gate": render_report,
            "warnings": [render_report["message"]],
            "stage_timings_ms": stage_timings_ms,
            "capabilities": capabilities,
            "execution_budget": execution_budget,
            "backend_attempts": backend_attempts,
            "recovery_checkpoints": prepared.get("recovery_checkpoints", []),
            "reasoning_plan": writing_ledger.get("reasoning_plan") if writing_ledger else None,
            "verification_summary": writing_ledger.get("verification_summary") if writing_ledger else None,
            "style_gate": prose_result.get("style"),
            "style_profile": style_profile,
            "style_brief": render_style_brief(style_profile),
            "process_files_in_delivery_directory": [],
        }
    delivery_session = {
        "schema_type": "delivery-session",
        "schema_version": 1,
        "status": "awaiting-render-review",
        "session_id": session.name.removeprefix(SESSION_PREFIX),
        "session_root": str(session),
        "template_path": str(template),
        "source_template_hash": prepared.get("source_hash") or sha256(template),
        "frozen_draft_path": str(session / "draft" / "frozen-draft.json"),
        "draft_sha256": frozen["draft_sha256"],
        "staged_docx_path": str(staged),
        "staged_docx_sha256": core_sha256(staged),
        "render_report_path": str(session / "qa" / "render.json"),
        "render_report_sha256": core_sha256(session / "qa" / "render.json"),
        "output_dir": str(output_dir.resolve()),
        "output_name": target.name,
        "parity_gate": parity,
        "structure_gate": structure_gate,
        "stage_timings_ms": stage_timings_ms,
        "capabilities": capabilities,
        "execution_budget": execution_budget,
        "backend_attempts": backend_attempts,
        "recovery_checkpoints": prepared.get("recovery_checkpoints", []),
        "reasoning_plan": writing_ledger.get("reasoning_plan") if writing_ledger else None,
        "verification_summary": writing_ledger.get("verification_summary") if writing_ledger else None,
        "style_gate": prose_result.get("style"),
        "style_profile": style_profile,
        "style_brief": render_style_brief(style_profile),
    }
    delivery_session_path = session / "delivery-session.json"
    delivery_session_path.write_text(json.dumps(delivery_session, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        **legacy_delivery,
        "schema_type": "delivery",
        "schema_version": 3,
        "type": "delivery.v3",
        "status": "awaiting-render-review",
        "artifact_path": None,
        "resume_session": str(delivery_session_path),
        "draft_sha256": frozen["draft_sha256"],
        "parity_gate": parity,
        "structure_gate": structure_gate,
        "render_gate": render_report,
        "stage_timings_ms": stage_timings_ms,
        "capabilities": capabilities,
        "execution_budget": execution_budget,
        "backend_attempts": backend_attempts,
        "recovery_checkpoints": prepared.get("recovery_checkpoints", []),
        "reasoning_plan": writing_ledger.get("reasoning_plan") if writing_ledger else None,
        "verification_summary": writing_ledger.get("verification_summary") if writing_ledger else None,
        "style_gate": prose_result.get("style"),
        "style_profile": style_profile,
        "style_brief": render_style_brief(style_profile),
        "process_files_in_delivery_directory": [],
    }


def stage_delivery(
    template: Path | None,
    output_dir: Path,
    content: dict[str, Any] | None = None,
    output_name: str | None = None,
    new_text: str | None = None,
    review_scores: dict[str, float] | None = None,
    word_limit: int | None = None,
    length_mode: str = "first-draft",
    evidence_cards: list[dict[str, Any]] | None = None,
    fact_ledger: dict[str, Any] | None = None,
    prepared_session: Path | None = None,
    budget_seconds: int | float | None = None,
) -> dict[str, Any]:
    if prepared_session is None:
        if template is None:
            raise AwesomeBenziError("E_TEMPLATE_MISSING", "prepare", "template or prepared session is required")
        from .template_schema import prepare_session

        prepared_session = Path(
            prepare_session(
                template,
                budget_seconds=budget_seconds if budget_seconds is not None else 300,
            )["manifest_path"]
        )
    prepared_data, loaded_session = _load_prepared_session(prepared_session)
    prepared_session = Path(prepared_data.get("manifest_path") or (loaded_session / "manifest.json"))
    session = loaded_session
    try:
        with session_lock(session):
            result = _stage_delivery(
                template,
                output_dir,
                content,
                output_name,
                new_text,
                review_scores,
                word_limit,
                length_mode,
                evidence_cards,
                fact_ledger,
                prepared_session,
                budget_seconds,
            )
    except AwesomeBenziError as exc:
        if exc.recoverability in {"environment", "fatal", "internal-invariant"} and session.exists():
            cleanup_session(session)
        raise
    if result.get("status") == "delivered" and session.exists():
        # Render review was skipped, so there is no second phase. Clean the
        # session only after the exclusive lock is released: Windows keeps the
        # lock file open while the lock is held.
        cleanup_session(session)
    return result


def _resume_delivery(resume_session: Path, render_review: dict[str, Any]) -> dict[str, Any]:
    resume_started = time.perf_counter()
    resume_session = resume_session.expanduser().resolve()
    data = json.loads(resume_session.read_text(encoding="utf-8"))
    session = Path(data["session_root"]).resolve()
    if resume_session.parent != session or not session.name.startswith(SESSION_PREFIX):
        raise AwesomeBenziError("E_PATH_ESCAPE", "delivery", "invalid delivery session path", recoverability="internal-invariant")
    current_state = SessionState(read_manifest(session)["state"])
    if current_state != SessionState.RENDER_REVIEW_PENDING:
        raise AwesomeBenziError(
            "E_INVALID_STATE_TRANSITION",
            "session",
            f"render review cannot resume from {current_state.value}",
            recoverability="internal-invariant",
        )
    staged = Path(data["staged_docx_path"]).resolve()
    frozen_path = Path(data["frozen_draft_path"]).resolve()
    render_path = Path(data["render_report_path"]).resolve()
    for candidate in (staged, frozen_path, render_path):
        try:
            candidate.relative_to(session)
        except ValueError as exc:
            raise AwesomeBenziError("E_PATH_ESCAPE", "delivery", "delivery session contains an escaped path", recoverability="internal-invariant") from exc
    if core_sha256(staged) != data["staged_docx_sha256"]:
        raise AwesomeBenziError("E_DRAFT_CHANGED_AFTER_FREEZE", "delivery", "staged DOCX changed before review")
    frozen = verify_frozen(frozen_path)
    if frozen["draft_sha256"] != data["draft_sha256"]:
        raise AwesomeBenziError("E_DRAFT_CHANGED_AFTER_FREEZE", "delivery", "delivery session draft hash does not match the frozen draft")
    template_path = Path(data["template_path"]).resolve()
    if not template_path.is_file() or core_sha256(template_path) != data["source_template_hash"]:
        raise AwesomeBenziError("E_SOURCE_CHANGED", "delivery", "source template changed before delivery")
    if core_sha256(render_path) != data["render_report_sha256"]:
        raise AwesomeBenziError("E_RENDER_REVIEW_BLOCKED", "render", "render report changed before review")
    render_report = json.loads(render_path.read_text(encoding="utf-8"))
    pages = render_report.get("pages", [])
    expected_page_numbers = list(range(1, int(render_report.get("page_count", 0)) + 1))
    if [item.get("page") for item in pages] != expected_page_numbers:
        raise AwesomeBenziError("E_RENDER_REVIEW_BLOCKED", "render", "render page manifest is incomplete or out of order")
    for page in pages:
        page_path = Path(page["path"]).resolve()
        try:
            page_path.relative_to(session)
        except ValueError as exc:
            raise AwesomeBenziError("E_PATH_ESCAPE", "render", "render page escapes the session", recoverability="internal-invariant") from exc
        if not page_path.is_file() or core_sha256(page_path) != page.get("sha256"):
            raise AwesomeBenziError("E_RENDER_REVIEW_BLOCKED", "render", "a rendered page changed before review")
    review_validation = validate_render_review(
        render_review,
        render_report,
        data["session_id"],
        expected_draft_sha256=frozen["draft_sha256"],
    )
    if review_validation["gate"] == "blocked":
        raise AwesomeBenziError(
            "E_RENDER_REVIEW_BLOCKED",
            "render",
            "visual review did not cover and pass every page",
            recoverability="internal-retry",
            details=review_validation,
        )
    transition(session, SessionState.RENDER_PASSED)
    skill_root = Path(__file__).resolve().parents[2]
    destination = safe_output_path(Path(data["output_dir"]), Path(data["template_path"]).stem, data["output_name"], skill_root)
    output_hash = atomic_commit(staged, destination)
    transition(session, SessionState.DELIVERED)
    result = {
        "schema_type": "delivery",
        "schema_version": 3,
        "type": "delivery.v3",
        "status": "delivered",
        "artifact_path": str(destination),
        "source_template_path": data["template_path"],
        "source_template_hash": data["source_template_hash"],
        "draft_sha256": frozen["draft_sha256"],
        "output_hash": output_hash,
        "output_strategy": "semantic-docx-rebuild",
        "output_format": "docx",
        "parity_gate": data["parity_gate"],
        "structure_gate": data["structure_gate"],
        "render_gate": {**render_report, "gate": "passed", "visual_review": review_validation},
        "stage_timings_ms": {
            **data.get("stage_timings_ms", {}),
            "resume_and_commit": round((time.perf_counter() - resume_started) * 1000, 2),
        },
        "capabilities": data.get("capabilities", {}),
        "execution_budget": data.get("execution_budget", {}),
        "backend_attempts": data.get("backend_attempts", []),
        "recovery_checkpoints": data.get("recovery_checkpoints", []),
        "reasoning_plan": data.get("reasoning_plan"),
        "verification_summary": data.get("verification_summary"),
        "process_files_in_delivery_directory": [],
    }
    return result


def resume_delivery(resume_session: Path, render_review: dict[str, Any]) -> dict[str, Any]:
    resume_session = resume_session.expanduser().resolve()
    data = json.loads(resume_session.read_text(encoding="utf-8"))
    session = Path(data["session_root"]).resolve()
    try:
        with session_lock(session):
            result = _resume_delivery(resume_session, render_review)
    except AwesomeBenziError as exc:
        if exc.recoverability in {"environment", "fatal", "internal-invariant"} and session.exists():
            cleanup_session(session)
        raise
    cleanup_session(session)
    return result


def main(argv: Iterable[str] | None = None) -> int:
    ensure_utf8_stdout()
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    build_parser = sub.add_parser("build")
    build_parser.add_argument("--template")
    build_parser.add_argument("--prepared-session")
    build_parser.add_argument("--content", help="JSON mapping stored outside the delivery directory")
    build_parser.add_argument("--output-dir")
    build_parser.add_argument("--output-name")
    build_parser.add_argument("--new-text", help="UTF-8 file containing only newly written prose")
    build_parser.add_argument("--review-scores", help="JSON object with six internal review scores")
    build_parser.add_argument("--word-limit", type=int)
    build_parser.add_argument("--length-mode", choices=["first-draft", "strict"], default="first-draft")
    build_parser.add_argument("--evidence-cards", help="JSON array of verified in-memory literature evidence cards")
    build_parser.add_argument("--fact-ledger", help="JSON object containing locked project facts and forbidden variants")
    build_parser.add_argument("--budget-seconds", type=float, default=None)
    build_parser.add_argument("--resume-session", help="delivery-session.json returned by the first build phase")
    build_parser.add_argument("--render-review", help="render-review.v1 JSON used to resume atomic delivery")
    args = parser.parse_args(argv)
    try:
        if args.resume_session:
            if not args.render_review:
                raise AwesomeBenziError("E_RENDER_REVIEW_BLOCKED", "render", "--render-review is required with --resume-session")
            result = resume_delivery(
                Path(args.resume_session),
                json.loads(Path(args.render_review).read_text(encoding="utf-8")),
            )
        else:
            if not args.output_dir:
                raise AwesomeBenziError("E_PATH_NOT_FOUND", "delivery", "--output-dir is required for the first build phase")
            result = stage_delivery(
                Path(args.template) if args.template else None,
                Path(args.output_dir),
                _read_content(args.content),
                args.output_name,
                Path(args.new_text).read_text(encoding="utf-8-sig") if args.new_text else None,
                _read_content(args.review_scores) if args.review_scores else None,
                args.word_limit,
                args.length_mode,
                json.loads(Path(args.evidence_cards).read_text(encoding="utf-8")) if args.evidence_cards else None,
                _read_content(args.fact_ledger) if args.fact_ledger else None,
                Path(args.prepared_session) if args.prepared_session else None,
                args.budget_seconds,
            )
        code = 0
    except VisionOcrRequired as exc:
        result = {
            "type": "internal-input-recovery.v1", "status": "requires-internal-vision-ocr",
            "artifact_path": None, "error": str(exc), "action": exc.action,
            "user_action_required": False, "question_batch": [],
            "backend_attempts": exc.backend_attempts,
            "execution_budget": exc.execution_budget,
        }
        code = 0
    except InternalRecoveryRequired as exc:
        result = {
            "type": "internal-input-recovery.v1", "status": "requires-internal-document-recovery",
            "artifact_path": None, "error": str(exc), "action": getattr(exc, "action", None),
            "user_action_required": False, "question_batch": [],
            "backend_attempts": getattr(exc, "backend_attempts", []),
            "execution_budget": getattr(exc, "execution_budget", None),
        }
        code = 0
    except AwesomeBenziError as exc:
        result = {**exc.to_dict(), "status": "blocked", "artifact_path": None}
        code = exc.exit_code
    except Exception as exc:
        result = {"type": "delivery-error", "status": "blocked", "artifact_path": None, "error": str(exc)}
        code = 2
    json.dump(result, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    return code


if __name__ == "__main__":
    raise SystemExit(main())
