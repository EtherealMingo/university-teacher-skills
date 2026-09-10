#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
xlsx_kit — 中文表格与成绩分析工具（office-layer 自建）

两个用途：
  build     JSON 规格 → XLSX（表头、列宽、冻结窗格、原生图表）
  analyze   读成绩表 → 逐题得分率/区分度/分布 → 输出分析工作簿

`analyze` 按中国教育测量学通行口径计算：
  难度 P  = 该题平均得分 / 满分（即得分率，越大越容易）
  区分度 D = 高分组得分率 − 低分组得分率（高低分组各取总分前/后 27%）
  D >= 0.4 优良；0.3~0.39 良好；0.2~0.29 尚可需修改；< 0.2 应淘汰
  另给出偏度/峰度作正态性快评

依赖：openpyxl
"""
from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("缺少 openpyxl：请运行 skills/office-layer/bootstrap.sh")

CN = "微软雅黑"
HDR_FILL = PatternFill("solid", fgColor="1F2D3D")
HDR_FONT = Font(name=CN, size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name=CN, size=10.5)
THIN = Side(style="thin", color="BFC7D1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _style_header(ws, row=1, ncol=None):
    ncol = ncol or ws.max_column
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def _autosize(ws, min_w=8, max_w=42):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min_w
        for cell in col:
            if cell.value is not None:
                # 中文按 2 个字符宽度估算
                s = str(cell.value)
                w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in s) + 2
                width = max(width, w)
        ws.column_dimensions[letter].width = min(width, max_w)


# ---------------------------------------------------------------- build
def build_from_spec(spec, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in spec.get("sheets", []):
        ws = wb.create_sheet(sh.get("name", "Sheet")[:31])
        headers = sh.get("header", [])
        if headers:
            ws.append(headers)
            _style_header(ws, 1, len(headers))
        for row in sh.get("rows", []):
            ws.append(row)
        for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=max(1, len(headers))):
            for cell in r:
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        if sh.get("freeze", True) and headers:
            ws.freeze_panes = "A2"
        if sh.get("widths"):
            for i, w in enumerate(sh["widths"], 1):
                ws.column_dimensions[get_column_letter(i)].width = w
        else:
            _autosize(ws)
        ch = sh.get("chart")
        if ch and headers and ws.max_row > 1:
            chart = BarChart() if ch.get("type", "bar") == "bar" else LineChart()
            chart.title = ch.get("title", "")
            chart.height, chart.width = ch.get("height", 9), ch.get("width", 18)
            data = Reference(ws, min_col=ch.get("value_col", 2),
                             min_row=1 if ch.get("header_row", True) else 2,
                             max_row=ws.max_row)
            cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
            chart.add_data(data, titles_from_data=ch.get("header_row", True))
            chart.set_categories(cats)
            anchor = ch.get("anchor", f"{get_column_letter(max(2, len(headers)) + 2)}2")
            ws.add_chart(chart, anchor)
    wb.save(str(out_path))
    return out_path


# ---------------------------------------------------------------- analyze
def _stats(vals):
    n = len(vals)
    if n == 0:
        return {}
    mean = sum(vals) / n
    var = sum((v - mean) ** 2 for v in vals) / n if n > 1 else 0.0
    sd = math.sqrt(var)
    skew = kurt = 0.0
    if sd > 0 and n > 2:
        skew = sum(((v - mean) / sd) ** 3 for v in vals) / n
        kurt = sum(((v - mean) / sd) ** 4 for v in vals) / n - 3
    return {"n": n, "mean": mean, "sd": sd, "min": min(vals), "max": max(vals),
            "skew": skew, "kurt": kurt}


def analyze_scores(src, out_path, sheet=None, full_marks=None,
                   id_col=1, total_col=None, first_q_col=None, n_questions=None,
                   high_pct=0.27):
    """
    约定表结构：第 1 行为表头；前若干列为学号/姓名等标识；其后为各题得分列。
    未显式指定时：标识取第 id_col 列；题目列 = 从 first_q_col 到最后一列；
    若 total_col 为空则用各题之和作总分。
    """
    wb = openpyxl.load_workbook(str(src), data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"error": "空工作表"}
    header = [str(h) if h is not None else f"列{i+1}" for i, h in enumerate(rows[0])]

    fq = (first_q_col - 1) if first_q_col else 1
    q_idx = list(range(fq, len(header))) if n_questions is None else \
        list(range(fq, fq + n_questions))
    if not q_idx:
        return {"error": "未识别到题目列"}

    def num(v):
        try:
            return float(v) if v is not None and v != "" else None
        except (TypeError, ValueError):
            return None

    students = []
    for r in rows[1:]:
        if r is None or all(v is None or v == "" for v in r):
            continue
        scores = [num(r[i]) if i < len(r) else None for i in q_idx]
        if all(s is None for s in scores):
            continue
        scores = [0.0 if s is None else s for s in scores]
        sid = r[id_col - 1] if id_col - 1 < len(r) else ""
        total = num(r[total_col - 1]) if total_col and total_col - 1 < len(r) else sum(scores)
        students.append({"id": sid, "scores": scores, "total": total if total is not None else sum(scores)})

    if not students:
        return {"error": "未解析到有效学生成绩行"}

    totals = [s["total"] for s in students]
    tstat = _stats(totals)
    ranked = sorted(students, key=lambda x: x["total"], reverse=True)
    k = max(1, int(round(len(ranked) * high_pct)))
    hi, lo = ranked[:k], ranked[-k:]
    marks = full_marks if isinstance(full_marks, list) else \
        [max((s["scores"][j] for s in students), default=1.0) or 1.0 for j in range(len(q_idx))]

    qstats = []
    for j, qi in enumerate(q_idx):
        full = marks[j] if j < len(marks) else 1.0
        full = full or 1.0
        all_s = [s["scores"][j] for s in students]
        mean = sum(all_s) / len(all_s)
        P = mean / full
        PH = (sum(s["scores"][j] for s in hi) / len(hi)) / full
        PL = (sum(s["scores"][j] for s in lo) / len(lo)) / full
        D = PH - PL
        if D >= 0.4:
            verdict = "优良"
        elif D >= 0.3:
            verdict = "良好"
        elif D >= 0.2:
            verdict = "尚可，建议修改"
        else:
            verdict = "区分度低，建议淘汰或重编"
        if P >= 0.9:
            verdict += "；难度过低"
        elif P <= 0.3:
            verdict += "；难度过高"
        qstats.append({"题号": header[qi], "满分": full, "平均分": round(mean, 2),
                       "得分率P": round(P, 3), "高分组得分率": round(PH, 3),
                       "低分组得分率": round(PL, 3), "区分度D": round(D, 3), "评价": verdict})

    # 输出工作簿
    out = openpyxl.Workbook()
    s1 = out.active
    s1.title = "逐题分析"
    s1.append(["题号", "满分", "平均分", "得分率P", "高分组得分率", "低分组得分率", "区分度D", "评价"])
    for q in qstats:
        s1.append([q["题号"], q["满分"], q["平均分"], q["得分率P"],
                   q["高分组得分率"], q["低分组得分率"], q["区分度D"], q["评价"]])
    _style_header(s1, 1, 8)
    for r in range(2, s1.max_row + 1):
        for c in range(1, 9):
            cell = s1.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
        d = s1.cell(row=r, column=7).value
        if isinstance(d, (int, float)) and d < 0.2:
            for c in range(1, 9):
                s1.cell(row=r, column=c).fill = WARN_FILL
    _autosize(s1)
    ch = BarChart()
    ch.title = "逐题得分率与区分度"
    ch.height, ch.width = 9, 20
    ch.add_data(Reference(s1, min_col=4, max_col=4, min_row=1, max_row=s1.max_row), titles_from_data=True)
    ch.add_data(Reference(s1, min_col=7, max_col=7, min_row=1, max_row=s1.max_row), titles_from_data=True)
    ch.set_categories(Reference(s1, min_col=1, min_row=2, max_row=s1.max_row))
    s1.add_chart(ch, "J2")

    s2 = out.create_sheet("总体概况")
    for row in [["项目", "数值"],
                ["考生数", tstat["n"]], ["平均分", round(tstat["mean"], 2)],
                ["标准差", round(tstat["sd"], 2)], ["最高分", tstat["max"]],
                ["最低分", tstat["min"]], ["及格率(>=60%)",
                 round(sum(1 for t in totals if t >= 0.6 * (sum(marks) or 1)) / len(totals), 3)],
                ["偏度", round(tstat["skew"], 3)], ["峰度", round(tstat["kurt"], 3)],
                ["高分组成员数", k], ["低分组成员数", k]]:
        s2.append(row)
    _style_header(s2, 1, 2)
    _autosize(s2)

    s3 = out.create_sheet("分数段分布")
    s3.append(["分数段", "人数", "占比"])
    mx = tstat["max"] or 1
    bins = [(0.9, 1.01, "90-100%"), (0.8, 0.9, "80-89%"), (0.7, 0.8, "70-79%"),
            (0.6, 0.7, "60-69%"), (0.0, 0.6, "60% 以下")]
    for lo_, hi_, label in bins:
        cnt = sum(1 for t in totals if lo_ * mx <= t < hi_ * mx)
        s3.append([label, cnt, round(cnt / len(totals), 3)])
    _style_header(s3, 1, 3)
    ch2 = BarChart(); ch2.title = "分数段分布"
    ch2.add_data(Reference(s3, min_col=2, min_row=1, max_row=s3.max_row), titles_from_data=True)
    ch2.set_categories(Reference(s3, min_col=1, min_row=2, max_row=s3.max_row))
    s3.add_chart(ch2, "E2")
    _autosize(s3)

    out.save(str(out_path))
    return {"students": len(students), "questions": len(qstats),
            "total_stats": {kk: round(vv, 3) for kk, vv in tstat.items()},
            "weak_questions": [q["题号"] for q in qstats if q["区分度D"] < 0.2],
            "hard_questions": [q["题号"] for q in qstats if q["得分率P"] <= 0.3],
            "easy_questions": [q["题号"] for q in qstats if q["得分率P"] >= 0.9],
            "out": str(out_path)}


def inspect(path):
    wb = openpyxl.load_workbook(str(path), data_only=True)
    return {name: {"rows": wb[name].max_row, "cols": wb[name].max_column}
            for name in wb.sheetnames}


def main():
    ap = argparse.ArgumentParser(prog="xlsx_kit", description="中文表格与成绩分析工具")
    sub = ap.add_subparsers(dest="cmd", required=True)

    a = sub.add_parser("build", help="JSON 规格 → XLSX")
    a.add_argument("--spec", required=True)
    a.add_argument("--out", required=True)

    b = sub.add_parser("analyze", help="成绩表 → 逐题分析工作簿")
    b.add_argument("--in", dest="src", required=True)
    b.add_argument("--out", required=True)
    b.add_argument("--sheet")
    b.add_argument("--id-col", type=int, default=1)
    b.add_argument("--total-col", type=int)
    b.add_argument("--first-q-col", type=int)
    b.add_argument("--n-questions", type=int)
    b.add_argument("--full-marks", help="JSON 数组，各题满分；缺省取该题最高分")

    c = sub.add_parser("inspect", help="读回工作簿结构")
    c.add_argument("--in", dest="src", required=True)

    args = ap.parse_args()
    if args.cmd == "build":
        build_from_spec(json.loads(Path(args.spec).read_text(encoding="utf-8")), args.out)
        print(json.dumps(inspect(args.out), ensure_ascii=False, indent=2))
    elif args.cmd == "analyze":
        fm = json.loads(args.full_marks) if args.full_marks else None
        res = analyze_scores(args.src, args.out, args.sheet, fm,
                             args.id_col, args.total_col, args.first_q_col, args.n_questions)
        print(json.dumps(res, ensure_ascii=False, indent=2))
    else:
        print(json.dumps(inspect(args.src), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
