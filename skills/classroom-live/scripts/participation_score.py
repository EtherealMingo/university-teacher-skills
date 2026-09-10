#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
participation_score — 课堂互动与平时分计算（自建，无上游可复用）

把「考勤 / 课堂互动 / 平时作业 / 随堂小测 记录 + 教师设定的平时分规则」
算成一份**可逐行举证**的《平时分依据表》规格，再交 office-layer 的 xlsx_kit 出表。

两个子命令：
  template  生成空白《课堂互动记录表》spec.json + 默认《平时分规则》JSON
  score     记录表 + 规则 → 平时分依据表.spec.json（+ 追溯明细 CSV + 待核清单 MD）

用法：
  PY=vendor/office-layer/.venv/bin/python
  $PY skills/classroom-live/scripts/participation_score.py template \\
      --out 记录表模板.spec.json --rules-out 平时分规则.json --course "数据结构"
  $PY skills/classroom-live/scripts/participation_score.py score \\
      --records 课堂互动记录.xlsx --rules 平时分规则.json --roster 名单.csv --outdir ./平时分
  $PY vendor/office-layer/scripts/xlsx_kit.py build --spec 平时分/平时分依据表.spec.json \\
      --out 平时分依据表.xlsx

设计原则（对应 CONVENTIONS.md，改动前先读）：
  1. 不虚构 —— 记录缺失只留空并进《待核与异常》，绝不补分、绝不推测、绝不生成学生数据
  2. 可追溯 —— 每个分值都带 来源文件 / 原始行号 / 适用规则 / 折算过程；明细各分项之和
               就是「平时分(初算)」，学生质询时可逐行举证
  3. 人工确认位 —— 汇总表固定保留「平时分(教师核定)」列，值为【待教师确认：____】
  4. 敏感数据 —— 只在本机读写；stdout 不打印学生姓名，只打印条数、文件路径与校验和

退出码：0 正常｜非 0 表示规则或数据不合法（信息里带工作表/行号/原因，脚本不猜测修正，
        例如权重之和不为 100、状态值规则里没定义、满分留空）
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

D = Decimal
Q4 = D("0.0001")
Q2 = D("0.01")

# ------------------------------------------------------------------ 表头别名
ALIASES: dict[str, list[str]] = {
    "sid": ["学号", "学生学号", "编号", "学籍号", "id", "studentid", "no"],
    "name": ["姓名", "学生姓名", "name"],
    "date": ["日期", "上课日期", "布置日期", "小测日期", "date"],
    "week": ["周次", "周", "week"],
    "status": ["状态", "考勤状态", "出勤情况", "考勤", "status"],
    "itype": ["互动类型", "类型", "互动形式", "形式", "type"],
    "level": ["表现等级", "等级", "表现", "level", "grade"],
    "score": ["得分", "分数", "成绩", "score"],
    "full": ["满分", "总分", "full", "fullmarks"],
    "item": ["加分项", "项目", "事由", "item"],
    "points": ["加分", "分值", "points"],
    "note": ["备注", "说明", "note", "remark"],
}

STATUS_ALIAS = {
    "出勤": "出勤", "到": "出勤", "正常": "出勤", "√": "出勤", "出席": "出勤", "上课": "出勤",
    "迟到": "迟到", "早退": "早退", "中途离场": "早退",
    "事假": "事假", "请假": "事假", "私假": "事假",
    "病假": "病假", "病": "病假",
    "公假": "公假", "公务": "公假",
    "旷课": "旷课", "缺勤": "旷课", "未到": "旷课", "缺席": "旷课", "×": "旷课", "x": "旷课", "X": "旷课",
}

LEVEL_ALIAS = {
    "a": "A", "b": "B", "c": "C", "d": "D",
    "优": "A", "优秀": "A", "很好": "A", "积极参与": "A", "主动": "A",
    "良": "B", "良好": "B", "较好": "B", "参与": "B",
    "中": "C", "中等": "C", "合格": "C", "及格": "C", "一般": "C", "被动": "C",
    "差": "D", "待改进": "D", "未参与": "D", "不参与": "D", "消极": "D",
}

DEFAULT_RULES: dict = {
    "course": "【待补：课程名】",
    "class_name": "【待补：班级】",
    "term": "【待补：学年学期】",
    "teacher": "【待教师确认：任课教师】",
    "weights": {"attendance": 20, "interaction": 30, "homework": 30, "quiz": 20},
    "attendance": {
        "base": 100,
        "floor": 0,
        "status": {"出勤": 0, "迟到": -10, "早退": -10, "事假": -5, "病假": 0, "公假": 0, "旷课": -25},
        "absent_alert": 3,
    },
    "interaction": {
        "level_points": {"A": 3, "B": 2, "C": 1, "D": 0},
        "full_credit_points": 20,
    },
    "homework": {"drop_lowest": 0},
    "quiz": {"drop_lowest": 0},
    "bonus": {"enabled": True, "cap": 5},
    "cap_total": 100,
    "round": 2,
    "sheets": {
        "attendance": "考勤",
        "interaction": "课堂互动",
        "homework": "平时作业",
        "quiz": "随堂小测",
        "bonus": "加分项",
        "votes": "投票结果",
    },
}


def sha256_of(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def norm_header(s) -> str:
    s = "" if s is None else str(s)
    for ch in (" ", "\u3000", "\t", "(", ")", "（", "）", "\n", "\r"):
        s = s.replace(ch, "")
    return s.strip().lower()


def norm_sid(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def norm_text(v) -> str:
    return "" if v is None else str(v).strip()


# ------------------------------------------------------------------ 读表
def read_table(path: Path, sheet: str | None = None) -> tuple[list[list], str]:
    """返回 (rows, 表名)。rows[0] 是表头；行号 = 该表内的 1-based 行序。"""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            sys.exit("缺少 openpyxl：请用 vendor/office-layer/.venv/bin/python 运行本脚本")
        wb = openpyxl.load_workbook(str(path), data_only=True)
        if sheet and sheet not in wb.sheetnames:
            return [], sheet
        name = sheet or wb.sheetnames[0]
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        # 去掉尾部整行为空的行
        while rows and all(c is None or str(c).strip() == "" for c in rows[-1]):
            rows.pop()
        return rows, name
    if path.suffix.lower() in (".csv", ".txt"):
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            rows = [r for r in csv.reader(f)]
        while rows and all(c.strip() == "" for c in rows[-1]):
            rows.pop()
        return rows, path.name
    sys.exit(f"不支持的记录文件类型：{path.name}（只支持 .xlsx/.xlsm/.csv）")


def col_index(header: list, key: str, required: bool = False) -> int | None:
    want = {norm_header(a) for a in ALIASES[key]}
    for i, h in enumerate(header):
        if norm_header(h) in want:
            return i
    if required:
        shown = "/".join(ALIASES[key][:3])
        sys.exit(f"记录表缺少「{shown}」列，实际表头为：{header}")
    return None


def cell(row: list, idx: int | None) -> str:
    if idx is None or idx >= len(row):
        return ""
    return norm_text(row[idx])


# ------------------------------------------------------------------ 校验
def validate_rules(rules: dict) -> None:
    w = rules.get("weights") or {}
    need = ["attendance", "interaction", "homework", "quiz"]
    missing = [k for k in need if k not in w]
    if missing:
        sys.exit(f"规则的 weights 缺少键：{missing}")
    total = sum(D(str(w[k])) for k in need)
    if total != 100:
        sys.exit(f"权重之和必须为 100，当前为 {total}（考勤 {w['attendance']} + 互动 {w['interaction']}"
                 f" + 作业 {w['homework']} + 小测 {w['quiz']}）。请教师确认后再算，脚本不做自动归一。")
    st = (rules.get("attendance") or {}).get("status")
    if not isinstance(st, dict) or not st:
        sys.exit("规则的 attendance.status 必须是非空对象，如 {\"出勤\":0,\"迟到\":-10,\"旷课\":-25}")
    for k, v in st.items():
        if D(str(v)) > 0:
            sys.exit(f"attendance.status.{k} = {v}，扣分项必须 ≤ 0（加分请走 bonus 加分项）")
    lp = (rules.get("interaction") or {}).get("level_points")
    if not isinstance(lp, dict) or not lp:
        sys.exit("规则的 interaction.level_points 必须是非空对象，如 {\"A\":3,\"B\":2,\"C\":1,\"D\":0}")


def resolve_mapping(raw: str, alias: dict, allowed: dict, where: str):
    if raw in allowed:
        return raw, D(str(allowed[raw]))
    canon = alias.get(raw) or alias.get(raw.lower())
    if canon and canon in allowed:
        return canon, D(str(allowed[canon]))
    opts = "、".join(allowed.keys())
    sys.exit(f"{where}：出现规则里没有定义的值「{raw}」。规则已定义：{opts}；"
             f"别名表可识别：{'、'.join(list(alias.keys())[:12])}…。"
             f"请不要让脚本猜，先在规则里补上该口径或改正记录。")


# ------------------------------------------------------------------ 计分
def fmt(v: Decimal, nd: int = 4) -> str:
    q = Q4 if nd == 4 else Q2
    return str(v.quantize(q, rounding=ROUND_HALF_UP))


def score_student(recs: dict, rules: dict, detail: list, alerts: list) -> dict:
    """算一个学生。detail 追加带原始行号的追溯行，返回汇总 dict。"""
    sh = {**DEFAULT_RULES["sheets"], **(rules.get("sheets") or {})}
    label = {"attendance": sh["attendance"], "interaction": sh["interaction"],
             "homework": sh["homework"], "quiz": sh["quiz"]}
    w = rules["weights"]
    wa, wi = D(str(w["attendance"])), D(str(w["interaction"]))
    wh, wq = D(str(w["homework"])), D(str(w["quiz"]))
    caps = D(str(rules.get("cap_total", 100)))
    nd = int(rules.get("round", 2))

    sid = recs["sid"]
    name = recs["name"]
    sub = {"attendance": D(0), "interaction": D(0), "homework": D(0), "quiz": D(0), "bonus": D(0)}
    rate = {"attendance": D(0), "interaction": D(0), "homework": D(0), "quiz": D(0)}
    bonus_total = D(0)
    notes: list[str] = []

    def add(comp, value, src, rowno, when, raw, rule, remark):
        value = value.quantize(Q4, rounding=ROUND_HALF_UP) if isinstance(value, Decimal) else D(str(value))
        sub[comp] += value
        detail.append([comp, src, rowno, when, raw, rule, fmt(value), remark])

    # ---- 考勤
    if wa > 0:
        att = rules["attendance"]
        base = D(str(att.get("base", 100)))
        floor = D(str(att.get("floor", 0)))
        if recs["attendance"]:
            add("attendance", base * wa / 100, "（规则基准）", "—", "—",
                f"基准 {base} 分", f"attendance.base={base}", "全勤基准分")
        for r in recs["attendance"]:
            canon, delta = resolve_mapping(r["status"], STATUS_ALIAS, att["status"],
                                           f"考勤第 {r['rowno']} 行状态")
            if delta != 0:
                add("attendance", delta * wa / 100, r["src"], r["rowno"], r["when"],
                    f"{r['status']}（规则取值 {delta}）",
                    f"attendance.status.{canon}={delta}", f"考勤记录 {r['week'] or ''}")
        raw_score = base + sum(
            (resolve_mapping(r["status"], STATUS_ALIAS, att["status"], "x")[1] for r in recs["attendance"]), D(0))
        raw_score = max(floor, raw_score)
        # 无任何考勤记录 → 该组件计 0（不送基准分），并已在下方进《待核与异常》
        rate["attendance"] = (raw_score / 100) if recs["attendance"] else D(0)
        absent = sum(1 for r in recs["attendance"]
                     if resolve_mapping(r["status"], STATUS_ALIAS, att["status"], "x")[0] == "旷课")
        if absent >= int(att.get("absent_alert", 3)):
            notes.append(f"旷课 {absent} 次")
            alerts.append(["考勤预警", sid, name,
                           f"旷课 {absent} 次，达到规则阈值 {att.get('absent_alert')}",
                           "按学校学籍/考勤规定处理，必要时报教务；本表不作自动处分"])
        if not recs["attendance"]:
            alerts.append(["记录缺失", sid, name, f"{label['attendance']}表无该生任何记录",
                           "确认是否漏录；未确认前该组件按 0 计，不推测"])
    else:
        notes.append("考勤权重 0")

    # ---- 课堂互动
    if wi > 0:
        ip = rules["interaction"]["level_points"]
        full = D(str(rules["interaction"].get("full_credit_points", 20)))
        if full <= 0:
            sys.exit("interaction.full_credit_points 必须 > 0")
        pts = D(0)
        for r in recs["interaction"]:
            canon, p = resolve_mapping(r["level"], LEVEL_ALIAS, ip, f"互动第 {r['rowno']} 行等级")
            pts += p
            add("interaction", p * wi / full, r["src"], r["rowno"], r["when"],
                f"{r['itype'] or '互动'} / {r['level']}（规则 {p} 点）",
                f"interaction.level_points.{canon}={p}", f"{full} 点折满分")
        capped = min(pts, full)
        if pts > full:
            add("interaction", (capped - pts) * wi / full, "（封顶调整）", "—", "—",
                f"累计 {pts} 点 > 满分基准 {full} 点",
                f"interaction.full_credit_points={full}", "超出部分封顶，不加分")
            notes.append("互动已封顶")
        rate["interaction"] = capped / full
        if not recs["interaction"]:
            alerts.append(["记录缺失", sid, name, f"{label['interaction']}表无该生任何记录",
                           "确认是否从未参与还是漏录；未确认前按 0 计，不推测"])

    # ---- 平时作业 / 随堂小测（同一套算法，各自权重）
    for comp, wgt in (("homework", wh), ("quiz", wq)):
        if wgt <= 0:
            continue
        items = recs[comp]
        if not items:
            alerts.append(["记录缺失", sid, name, f"{label[comp]}表无该生成绩记录",
                           "确认是否漏录；未确认前该组件按 0 计"])
            continue
        drop = int((rules.get(comp) or {}).get("drop_lowest", 0) or 0)
        parsed = []
        for r in items:
            try:
                s = D(str(r["score"]))
                f = D(str(r["full"])) if r["full"] else D(0)
            except Exception:
                sys.exit(f"{comp} 第 {r['rowno']} 行得分/满分不是数字：{r['score']} / {r['full']}")
            if f <= 0:
                sys.exit(f"{comp} 第 {r['rowno']} 行满分必须 > 0，当前为「{r['full']}」。"
                         f"请不要留空满分，从规则或记录里补齐后再算。")
            parsed.append((r, s, f, s / f))
        kept = parsed
        dropped = []
        if drop > 0 and len(parsed) > 1:
            ordered = sorted(parsed, key=lambda x: x[3])
            dropped = ordered[:min(drop, len(parsed) - 1)]
            kept = [p for p in parsed if p not in dropped]
        n = len(kept)
        for r, s, f, rt in kept:
            add(comp, rt * wgt / n, r["src"], r["rowno"], r["when"],
                f"{fmt(s,2)}/{fmt(f,2)}",
                f"weights.{comp}={wgt}", f"该组件 {n} 条有效记录取平均")
        for r, s, f, rt in dropped:
            detail.append([comp, r["src"], r["rowno"], r["when"], f"{fmt(s,2)}/{fmt(f,2)}",
                           f"{comp}.drop_lowest={drop}", fmt(D(0)), "按规则剔除最低一次，不计入"])
        rate[comp] = sum(rt for *_x, rt in kept) / n if n else D(0)

    # ---- 加分项
    b = rules.get("bonus") or {}
    if b.get("enabled") and recs["bonus"]:
        cap = D(str(b.get("cap", 0)))
        for r in recs["bonus"]:
            try:
                v = D(str(r["points"]))
            except Exception:
                sys.exit(f"加分项第 {r['rowno']} 行加分值不是数字：{r['points']}")
            if v < 0:
                sys.exit(f"加分项第 {r['rowno']} 行加分为负（{v}）。扣分请走考勤/等级规则，不用负加分。")
            bonus_total += v
            add("bonus", v, r["src"], r["rowno"], r["when"], r["item"] or "加分",
                f"bonus.cap={cap}", "需有原始依据（教师记录/课堂记录）")
        if bonus_total > cap:
            add("bonus", cap - bonus_total, "（封顶调整）", "—", "—",
                f"加分合计 {fmt(bonus_total,2)} > 上限 {fmt(cap,2)}",
                f"bonus.cap={cap}", "超出部分封顶")
            notes.append("加分已封顶")
            bonus_total = cap
    elif recs["bonus"] and not b.get("enabled"):
        alerts.append(["规则外记录", sid, name, "记录表有加分项但规则里 bonus.enabled=false",
                       "请教师确认是否启用加分；确认前不加分"])

    component_sum = sub["attendance"] + sub["interaction"] + sub["homework"] + sub["quiz"]
    draft = component_sum + bonus_total
    if draft > caps:
        add("bonus", caps - draft, "（总分封顶）", "—", "—",
            f"合计 {fmt(draft,2)} > 上限 {fmt(caps,2)}", f"cap_total={caps}", "按上限封顶")
        notes.append("总分封顶")
        draft = caps
    final = draft.quantize(Q2 if nd == 2 else Q4, rounding=ROUND_HALF_UP)

    detail.append(["合计", "（各分项之和）", "—", "—", "—", "本表各分项相加",
                   fmt(final), "四舍五入至 " + str(nd) + " 位，即《平时分(初算)》"])

    return {
        "sid": sid, "name": name,
        "rate": {k: rate[k].quantize(Q4, rounding=ROUND_HALF_UP) for k in rate},
        "sub": {k: sub[k].quantize(Q2, rounding=ROUND_HALF_UP) for k in sub},
        "bonus": bonus_total.quantize(Q2, rounding=ROUND_HALF_UP),
        "final": final, "notes": "；".join(notes),
    }


# ------------------------------------------------------------------ 主流程
def cmd_score(a) -> int:
    rec_path, rules_path = Path(a.records), Path(a.rules)
    if not rec_path.exists():
        sys.exit(f"找不到记录表：{rec_path}")
    if not rules_path.exists():
        sys.exit(f"找不到规则文件：{rules_path}")
    rules = json.loads(rules_path.read_text(encoding="utf-8"))
    validate_rules(rules)

    sh = {**DEFAULT_RULES["sheets"], **(rules.get("sheets") or {})}
    detail: list[list] = []
    alerts: list[list] = []
    students: dict[str, dict] = {}
    order: list[str] = []
    sources: list[list] = []

    def bucket(sid, name):
        if sid not in students:
            students[sid] = {"sid": sid, "name": name,
                             "attendance": [], "interaction": [], "homework": [], "quiz": [], "bonus": []}
            order.append(sid)
        if name and not students[sid]["name"]:
            students[sid]["name"] = name
        return students[sid]

    # 名单（可选）：让整学期零互动的学生也出现在表里，而不是"消失"
    roster_sids: set[str] = set()
    if a.roster:
        rp = Path(a.roster)
        if not rp.exists():
            sys.exit(f"找不到名单文件：{rp}")
        rows, sname = read_table(rp)
        if not rows:
            sys.exit(f"名单文件为空：{rp.name}")
        si, ni = col_index(rows[0], "sid", True), col_index(rows[0], "name")
        for r in rows[1:]:
            sid = norm_sid(r[si] if si < len(r) else "")
            if sid:
                roster_sids.add(sid)
                bucket(sid, cell(r, ni))
        sources.append(["名单文件", f"{rp.name}（{len(rows)-1} 人，sha256 {sha256_of(rp)[:16]}）"])

    # 四个记录组件
    comp_map = {
        "attendance": {"sheet": sh["attendance"], "need": ["status"], "has_level": False},
        "interaction": {"sheet": sh["interaction"], "need": ["level"], "has_level": True},
        "homework": {"sheet": sh["homework"], "need": ["score"], "has_level": False},
        "quiz": {"sheet": sh["quiz"], "need": ["score"], "has_level": False},
    }
    loaded: dict[str, list[list]] = {}
    for comp, cfg in comp_map.items():
        if D(str(rules["weights"][comp])) <= 0:
            continue
        rows, sname = read_table(rec_path, cfg["sheet"])
        if not rows:
            sys.exit(f"记录表里找不到工作表「{cfg['sheet']}」（或该表为空）。"
                     f"可用工作表请用 xlsx_kit inspect 查看；表名可在规则的 sheets 里改。")
        h = rows[0]
        req = {"sid": True, "score": comp in ("homework", "quiz")}
        if comp == "attendance":
            req["status"] = True
        if comp == "interaction":
            req["level"] = True
        idx = {k: col_index(h, k, req.get(k, False)) for k in
               ("sid", "name", "date", "week", "status", "itype", "level", "score", "full", "note")}
        loaded[comp] = rows
        cnt = 0
        for i, r in enumerate(rows[1:], start=2):
            sid = norm_sid(r[idx["sid"]] if idx["sid"] is not None and idx["sid"] < len(r) else "")
            if not sid:
                if any(norm_text(c) for c in r):
                    alerts.append(["行数据异常", "—", "—",
                                   f"{cfg['sheet']} 第 {i} 行有内容但缺学号",
                                   "补学号或删空行后重算；本行未计入"])
                continue
            st = bucket(sid, cell(r, idx["name"]))
            when = " ".join(x for x in (cell(r, idx["date"]), cell(r, idx["week"])) if x)
            entry = {"src": f"{rec_path.name}/{sname}", "rowno": i, "when": when,
                     "week": cell(r, idx["week"]), "note": cell(r, idx["note"])}
            if comp == "attendance":
                raw = cell(r, idx["status"])
                if not raw:
                    alerts.append(["行数据异常", sid, st["name"],
                                   f"{cfg['sheet']} 第 {i} 行状态为空", "补状态后重算；本行未计入"])
                    continue
                entry["status"] = raw
            elif comp == "interaction":
                if not cell(r, idx["level"]):
                    alerts.append(["行数据异常", sid, st["name"],
                                   f"{cfg['sheet']} 第 {i} 行等级为空", "补等级后重算；本行未计入"])
                    continue
                entry["level"] = cell(r, idx["level"])
                entry["itype"] = cell(r, idx["itype"])
            else:
                s = cell(r, idx["score"])
                if s == "":
                    alerts.append(["行数据异常", sid, st["name"],
                                   f"{cfg['sheet']} 第 {i} 行得分为空",
                                   "补得分或标记缺交（缺交应有 0 分记录，而不是空）"])
                    continue
                entry["score"], entry["full"] = s, cell(r, idx["full"])
            st[comp].append(entry)
            cnt += 1
        sources.append([f"{cfg['sheet']}记录", f"{rec_path.name}/{sname}，{cnt} 条有效记录"])

    # 加分项
    b = rules.get("bonus") or {}
    if b.get("enabled"):
        rows, sname = read_table(rec_path, sh["bonus"])
        if rows:
            h = rows[0]
            idx = {k: col_index(h, k, k in ("sid", "points")) for k in ("sid", "name", "date", "item", "points")}
            cnt = 0
            for i, r in enumerate(rows[1:], start=2):
                sid = norm_sid(r[idx["sid"]] if idx["sid"] is not None and idx["sid"] < len(r) else "")
                if not sid:
                    continue
                st = bucket(sid, cell(r, idx["name"]))
                st["bonus"].append({"src": f"{rec_path.name}/{sname}", "rowno": i,
                                    "when": cell(r, idx["date"]), "item": cell(r, idx["item"]),
                                    "points": cell(r, idx["points"])})
                cnt += 1
            if cnt:
                sources.append(["加分项记录", f"{rec_path.name}/{sname}，{cnt} 条"])

    if not order:
        sys.exit("没有读到任何有效记录，无法出表。请检查记录表内容与规则的 sheets 表名。")

    for sid in order:
        st = students[sid]
        res = score_student(st, rules, detail, alerts)
        st["res"] = res

    # 名单与记录比对：只提示，不自行增删学生
    if roster_sids:
        for sid in order:
            if sid not in roster_sids:
                alerts.append(["名单不符", sid, students[sid]["name"],
                               "该学号出现在记录表，但不在提供的名单里",
                               "核对是否插班/退课/学号录入错误；确认前保留该生行，不要删除"])

    # ---- 组装 spec
    wa, wi = rules["weights"]["attendance"], rules["weights"]["interaction"]
    wh, wq = rules["weights"]["homework"], rules["weights"]["quiz"]
    summary_rows = []
    for sid in order:
        r = students[sid]["res"]
        summary_rows.append([
            sid, r["name"],
            float(r["rate"]["attendance"]), float(r["rate"]["interaction"]),
            float(r["rate"]["homework"]), float(r["rate"]["quiz"]),
            float(r["sub"]["attendance"]), float(r["sub"]["interaction"]),
            float(r["sub"]["homework"]), float(r["sub"]["quiz"]),
            float(r["sub"]["attendance"] + r["sub"]["interaction"] + r["sub"]["homework"] + r["sub"]["quiz"]),
            float(r["bonus"]), float(r["final"]),
            "【待教师确认：____】", r["notes"] or "",
        ])
    summary_rows.sort(key=lambda x: (-x[12], x[0]))

    # detail 是按学生顺序连续追加的，按「合计」行切块，逐块贴回学号
    detail_rows = []
    idx = 0
    for sid in order:
        r = students[sid]["res"]
        grp = []
        # 该学生的明细行连续存放，遇到「合计」行结束
        while idx < len(detail):
            row = detail[idx]
            idx += 1
            if row[0] == "合计":
                grp.append(row)
                break
            grp.append(row)
        for g in grp:
            detail_rows.append([sid, r["name"], g[0], g[1], g[2], g[3], g[4], g[5], g[6], g[7]])

    # 分数段分布
    bands = [(0, 59.999, "0–59"), (60, 69.999, "60–69"), (70, 79.999, "70–79"),
             (80, 89.999, "80–89"), (90, 100.001, "90–100")]
    dist = []
    for lo, hi, label in bands:
        c = sum(1 for r in summary_rows if lo <= r[12] < hi)
        dist.append([label, c])

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    rule_rows = [
        ["课程", rules.get("course", "")], ["班级", rules.get("class_name", "")],
        ["学年学期", rules.get("term", "")], ["任课教师（规则填写）", rules.get("teacher", "")],
        ["权重：考勤", f"{wa}%"], ["权重：课堂互动", f"{wi}%"],
        ["权重：平时作业", f"{wh}%"], ["权重：随堂小测", f"{wq}%"],
        ["考勤口径", "基准 %s 分；%s；下限 %s 分" % (
            rules["attendance"].get("base", 100),
            "；".join(f"{k} {v:+g}" for k, v in rules["attendance"]["status"].items()),
            rules["attendance"].get("floor", 0))],
        ["互动口径", "等级折算 %s；%s 点折满分（超出封顶）" % (
            "，".join(f"{k}={v}" for k, v in rules["interaction"]["level_points"].items()),
            rules["interaction"].get("full_credit_points", 20))],
        ["作业口径", f"各次得分率取平均；剔除最低 {rules.get('homework',{}).get('drop_lowest',0)} 次"],
        ["小测口径", f"各次得分率取平均；剔除最低 {rules.get('quiz',{}).get('drop_lowest',0)} 次"],
        ["加分上限", f"{rules.get('bonus',{}).get('cap',0)} 分（需有原始依据）"],
        ["总分上限", f"{rules.get('cap_total',100)} 分"],
        ["计算时间", now],
        ["计算口径", "各分项折算得分之和 = 平时分（初算）；明细表逐行可加总核对"],
        ["舍入口径", f"分项保留 4 位小数，汇总四舍五入至 {rules.get('round',2)} 位"],
        ["边界声明", "本表是计算草稿与举证材料，不是最终成绩；最终分值由任课教师核定后生效"],
        ["个人信息提示", "本文件含学生姓名与学号，属个人信息，请仅本地保管、不外传，用毕删除"],
        ["人工确认提示", "汇总表「平时分(教师核定)」列为空位，请教师逐项核对后填写"],
    ]
    alert_rows = alerts + [["教师核定", "—", "—", "汇总表「平时分(教师核定)」列需教师逐人核对填写",
                             "本表不代教师定分"]]

    spec = {
        "sheets": [
            {"name": "平时分汇总",
             "header": ["学号", "姓名",
                        f"考勤得分率({wa}%)", f"互动得分率({wi}%)",
                        f"作业得分率({wh}%)", f"小测得分率({wq}%)",
                        f"考勤折算({wa})", f"互动折算({wi})", f"作业折算({wh})", f"小测折算({wq})",
                        "加权小计", "加分", "平时分(初算)", "平时分(教师核定)", "自动备注"],
             "rows": summary_rows,
             "widths": [12, 10, 12, 12, 12, 12, 11, 11, 11, 11, 10, 8, 12, 20, 22],
             "freeze": True},
            {"name": "追溯明细",
             "header": ["学号", "姓名", "组件", "来源文件", "原始行号", "日期/周次",
                        "原始记录", "适用规则", "折算得分", "说明"],
             "rows": detail_rows, "freeze": True},
            {"name": "分数段分布",
             "header": ["平时分(初算)区间", "人数"], "rows": dist,
             "chart": {"type": "bar", "title": "平时分(初算)分布", "value_col": 2,
                       "header_row": True, "anchor": "D2"}},
            {"name": "待核与异常",
             "header": ["类型", "学号", "姓名", "说明", "建议处理"],
             "rows": alert_rows, "freeze": True},
            {"name": "数据来源与规则快照",
             "header": ["项目", "内容"],
             "rows": sources + [["规则文件", f"{rules_path.name}（sha256 {sha256_of(rules_path)[:16]}）"],
                                ["记录文件", f"{rec_path.name}（sha256 {sha256_of(rec_path)[:16]}）"]] + rule_rows,
             "freeze": True},
        ]
    }

    outdir = Path(a.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    spec_path = outdir / "平时分依据表.spec.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False, indent=2), encoding="utf-8")

    csv_path = outdir / "计算追溯明细.csv"
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["学号", "姓名", "组件", "来源文件", "原始行号", "日期/周次", "原始记录",
                    "适用规则", "折算得分", "说明"])
        w.writerows(detail_rows)

    md = ["# 待核与异常清单", "",
          f"生成时间：{now}　课程：{rules.get('course','')}　班级：{rules.get('class_name','')}", "",
          "> 本文件含学生姓名与学号，属个人信息，请仅本地保管，核对完后删除。", "",
          "| 类型 | 学号 | 姓名 | 说明 | 建议处理 |", "|---|---|---|---|---|"]
    for r in alert_rows:
        md.append("| " + " | ".join(str(x) for x in r) + " |")
    md += ["", "## 需教师核定的口径", "",
           "1. 权重与扣分细则是否就是本学期向学生公布的那一版（本表用的是规则文件里的版本）",
           "2. 上表每条「记录缺失/行数据异常」是漏录还是事实，确认后补录并重算",
           "3. 「平时分(教师核定)」列由教师逐人填写，本表不代定分"]
    md_path = outdir / "待核与异常清单.md"
    md_path.write_text("\n".join(md) + "\n", encoding="utf-8")

    print(json.dumps({
        "students": len(order),
        "detail_rows": len(detail_rows),
        "alerts": len(alerts),
        "spec": str(spec_path),
        "detail_csv": str(csv_path),
        "alert_md": str(md_path),
        "weights": rules["weights"],
        "note": "stdout 不打印学生姓名；后续请用 xlsx_kit build 出表，并在本地核对",
        "next": f'vendor/office-layer/.venv/bin/python vendor/office-layer/scripts/xlsx_kit.py '
                f'build --spec {spec_path} --out 平时分依据表.xlsx',
    }, ensure_ascii=False, indent=2))
    return 0


def _template_spec(course: str) -> dict:
    return {"sheets": [
        {"name": "考勤", "header": ["日期", "周次", "学号", "姓名", "状态", "备注"],
         "rows": [["2025-09-08", "第1周", "【学号】", "【姓名】", "出勤", ""],
                  ["2025-09-08", "第1周", "", "", "迟到", "状态可填：出勤/迟到/早退/事假/病假/公假/旷课"]],
         "widths": [12, 8, 12, 10, 8, 46], "freeze": True},
        {"name": "课堂互动", "header": ["日期", "周次", "学号", "姓名", "互动类型", "表现等级", "备注"],
         "rows": [["2025-09-08", "第1周", "【学号】", "【姓名】", "点名回答", "B", ""],
                  ["", "", "", "", "小组讨论", "A", "互动类型：点名回答/小组讨论/投票/随堂小测/主动提问"],
                  ["", "", "", "", "", "", "等级：A 优 / B 良 / C 中 / D 未参与"]],
         "widths": [12, 8, 12, 10, 12, 10, 40], "freeze": True},
        {"name": "投票结果", "header": ["日期", "周次", "题号", "选项", "人数", "正确率", "主要误选项"],
         "rows": [["2025-09-08", "第1周", "Q1", "B", 0, "", "用于课中诊断，不直接进平时分，除教师另有规定"],
                  ["", "", "Q1", "C", 0, "", "误选集中项即本次课的认知盲点"]],
         "widths": [12, 8, 8, 8, 8, 10, 40], "freeze": True},
        {"name": "平时作业", "header": ["布置日期", "周次", "学号", "姓名", "得分", "满分", "备注"],
         "rows": [["2025-09-15", "第2周", "【学号】", "【姓名】", 0, 10, "缺交请填 0 分，不要留空"]],
         "widths": [12, 8, 12, 10, 8, 8, 40], "freeze": True},
        {"name": "随堂小测", "header": ["日期", "周次", "学号", "姓名", "得分", "满分", "备注"],
         "rows": [["2025-09-22", "第3周", "【学号】", "【姓名】", 0, 20, ""]],
         "widths": [12, 8, 12, 10, 8, 8, 40], "freeze": True},
        {"name": "加分项", "header": ["日期", "学号", "姓名", "加分项", "加分", "备注"],
         "rows": [["2025-10-13", "【学号】", "【姓名】", "课堂纠错/作业展示/学科竞赛", 0, "必须有原始依据才可加分"]],
         "widths": [12, 12, 10, 24, 8, 40], "freeze": True},
        {"name": "填写说明", "header": ["项目", "内容"], "rows": [
            ["课程", course],
            ["用途", "课堂互动与考勤的过程记录表；学期末由 participation_score.py 算平时分"],
            ["填写时机", "下课后 5 分钟内录入，趁印象还在；不要期末凭回忆集中补录"],
            ["记录原则", "只记事实，不记评价性判语；未发生的事不留记录，缺失就留空"],
            ["隐私提醒", "本表含学生姓名与学号，属个人信息，仅本地保管、不外传；用毕删除中间文件"],
            ["表名要求", "工作表名必须是 考勤/课堂互动/投票结果/平时作业/随堂小测/加分项（可在规则 JSON 的 sheets 里改名）"],
        ], "widths": [14, 90], "freeze": True},
    ]}


def cmd_template(a) -> int:
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(_template_spec(a.course), ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "spec": str(out),
        "next": f'vendor/office-layer/.venv/bin/python vendor/office-layer/scripts/xlsx_kit.py '
                f'build --spec {out} --out 课堂互动记录表.xlsx',
    }, ensure_ascii=False, indent=2))
    if a.rules_out:
        rp = Path(a.rules_out)
        r = json.loads(json.dumps(DEFAULT_RULES))
        r["course"] = a.course
        rp.write_text(json.dumps(r, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"已写默认规则：{rp}（请教师核对权重与扣分细则后再用）")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(prog="participation_score",
                                 description="课堂互动与平时分计算（可追溯）")
    sub = ap.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("score", help="记录表 + 规则 → 平时分依据表 spec")
    s.add_argument("--records", required=True, help="课堂互动记录表（.xlsx/.csv）")
    s.add_argument("--rules", required=True, help="平时分规则 JSON")
    s.add_argument("--roster", help="学生名单（.xlsx/.csv，含学号/姓名），可选")
    s.add_argument("--outdir", default="平时分", help="输出目录，默认 ./平时分")

    t = sub.add_parser("template", help="生成空白记录表 spec 与默认规则")
    t.add_argument("--out", required=True, help="记录表 spec 输出路径")
    t.add_argument("--rules-out", help="同时输出默认规则 JSON 的路径")
    t.add_argument("--course", default="【待补：课程名】")

    a = ap.parse_args()
    return cmd_score(a) if a.cmd == "score" else cmd_template(a)


if __name__ == "__main__":
    sys.exit(main())
