#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
desc_stats.py — 研究数据的描述统计与交叉表（只读，绝不修改原始数据）

职责边界（很重要）：
  本脚本**只做描述性统计**：频数、缺失、集中与离散趋势、分组交叉表。
  它不做也不该做：推断统计、因果判断、缺失值填补、剔除"异常值"。
  任何填补、剔除都必须在数据字典里写明并由教师本人决定。

输入：.xlsx（openpyxl，第 1 行为表头）或 .csv / .tsv
输出：
  --out-md     Markdown 表格（可直接进报告）
  --out-json   符合 office-layer/xlsx_kit.py build 的 JSON 规格，供出 xlsx

用法示例：
  PY="vendor/office-layer/.venv/bin/python"
  $PY skills/research-assistant/scripts/desc_stats.py \
      --in 数据.xlsx --sheet 数据 \
      --cross 组别,后测成绩 --cross 组别,性别 \
      --missing-codes "-99,-98,-97" \
      --out-md 描述统计.md --out-json 描述统计.json
"""

import argparse
import csv
import json
import math
import os
import sys
from collections import Counter, OrderedDict

# ------------------------------------------------------------------ 读入
def _read_xlsx(path, sheet=None):
    try:
        import openpyxl
    except ImportError:
        sys.exit("缺少 openpyxl：请用 vendor/office-layer/.venv/bin/python 运行本脚本")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
    wb.close()
    return rows, (sheet or "第一个工作表")


def _read_delimited(path):
    delim = "\t" if path.lower().endswith((".tsv", ".tab")) else ","
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        rows = [list(r) for r in csv.reader(f, delimiter=delim)]
    return rows, "CSV"


def load_table(path, sheet=None):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet)
    if ext in (".csv", ".tsv", ".tab", ".txt"):
        return _read_delimited(path)
    sys.exit(f"暂不支持的文件类型：{ext}（请另存为 .xlsx 或 .csv）")


# ------------------------------------------------------------------ 工具
def to_float(v, missing_codes):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        f = float(v)
        if f in missing_codes:
            return None
        return f if math.isfinite(f) else None
    s = str(v).strip()
    if s == "" or s.lower() in ("na", "n/a", "nan", "null", "none", "."):
        return None
    try:
        f = float(s.replace(",", ""))
    except ValueError:
        return None
    return None if f in missing_codes else f


def pct(n, d):
    return 0.0 if d == 0 else round(100.0 * n / d, 1)


def _cat_key(v, missing_codes):
    """分类取值归一化；缺失（空白或缺失编码）返回 None。"""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        f = float(v)
        if f in missing_codes:
            return None
        return _fmt_key(f)
    s = str(v).strip()
    if s == "" or s.lower() in ("na", "n/a", "nan", "null", "none", "."):
        return None
    try:
        f = float(s.replace(",", ""))
    except ValueError:
        return s
    if f in missing_codes or not math.isfinite(f):
        return None
    return _fmt_key(f)


def _parses_as_number(v):
    """非空白取值能否解析为数值（缺失编码也算能解析）。"""
    if v is None:
        return False
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return math.isfinite(float(v))
    s = str(v).strip()
    if s == "" or s.lower() in ("na", "n/a", "nan", "null", "none", "."):
        return False
    try:
        return math.isfinite(float(s.replace(",", "")))
    except ValueError:
        return False


def _fmt_key(f):
    """数值型分类取值不要显示成 1.0 这种多余小数。"""
    return str(int(f)) if float(f).is_integer() else str(f)


def quantile(sorted_vals, q):
    """线性插值分位数（不引入额外依赖）。"""
    if not sorted_vals:
        return None
    if len(sorted_vals) == 1:
        return sorted_vals[0]
    pos = (len(sorted_vals) - 1) * q
    lo = int(math.floor(pos))
    hi = int(math.ceil(pos))
    if lo == hi:
        return sorted_vals[lo]
    return sorted_vals[lo] + (sorted_vals[hi] - sorted_vals[lo]) * (pos - lo)


# ------------------------------------------------------------------ 主逻辑
_ID_HINT = ("号", "编号", "序号", "id", "ID", "code", "代码", "班级", "学校", "组别", "性别")
_SCORE_HINT = ("成绩", "分数", "得分", "总分", "score", "Score")


def _looks_like_id(name):
    return any(h in name for h in _ID_HINT)


def _looks_like_score(name):
    return any(h in name for h in _SCORE_HINT)


def profile(header, data_rows, missing_codes, cat_max_unique=12,
            force_num=(), force_cat=()):
    """
    返回 (numeric, categorical, lowcard)。

    判定规则（写进报告，不做黑箱猜测）：
      1. 变量名像标识/分组（学号、编号、班级、性别、组别…）→ 分类
      2. 变量名像成绩/分数 且 取值可解析 → 连续
      3. 非空白取值 ≥95% 可解析为数值 且 唯一值 > cat_max_unique → 连续
      4. 其余 → 分类
      --num / --cat 优先级最高，可覆盖以上全部规则。
    """
    numeric, categorical, lowcard = [], [], []
    for j, name in enumerate(header):
        raw = [r[j] if j < len(r) else None for r in data_rows]
        nonblank = [v for v in raw if not (v is None or str(v).strip() == "")]
        parsed = [v for v in (to_float(v, missing_codes) for v in raw) if v is not None]
        n_parseable = sum(1 for v in raw if _parses_as_number(v))
        n_missing = len(raw) - len(parsed)
        uniq = len(set(parsed))
        mostly_numeric = len(nonblank) > 0 and n_parseable / len(nonblank) >= 0.95

        if name in force_cat:
            is_num = False
        elif name in force_num:
            is_num = True
        elif _looks_like_id(name):
            is_num = False
        elif _looks_like_score(name) and mostly_numeric and parsed:
            is_num = True
        else:
            is_num = mostly_numeric and uniq > cat_max_unique

        if name in force_num and not parsed:
            print(f"[警告] --num 指定的变量「{name}」无有效数值，已跳过", file=sys.stderr)
            continue

        rec = OrderedDict(变量=name, N有效=len(parsed), 缺失=n_missing,
                          缺失率=f"{pct(n_missing, len(raw))}%", 唯一值=uniq)
        if is_num:
            sv = sorted(parsed)
            mean = sum(sv) / len(sv)
            var = sum((v - mean) ** 2 for v in sv) / (len(sv) - 1) if len(sv) > 1 else 0.0
            rec.update(均值=round(mean, 3), 标准差=round(math.sqrt(var), 3),
                       最小=round(sv[0], 3), P25=round(quantile(sv, 0.25), 3),
                       中位数=round(quantile(sv, 0.5), 3), P75=round(quantile(sv, 0.75), 3),
                       最大=round(sv[-1], 3))
            numeric.append(rec)
        else:
            cnt = Counter()
            for v in raw:
                key = _cat_key(v, missing_codes)
                if key is None:
                    continue
                cnt[key] += 1
            n_valid_cat = sum(cnt.values())
            rec["N有效"] = n_valid_cat
            rec["缺失"] = len(raw) - n_valid_cat
            rec["缺失率"] = f"{pct(rec['缺失'], len(raw))}%"
            rec["类别数"] = len(cnt)
            rec["唯一值"] = len(cnt)
            rec["最高频取值"] = "；".join(
                f"{k} ({v}, {pct(v, n_valid_cat)}%)" for k, v in cnt.most_common(8)) or "—"
            categorical.append((rec, cnt, rec["缺失"]))
            if mostly_numeric and name not in force_cat and not _looks_like_id(name) \
                    and not _looks_like_score(name):
                lowcard.append((name, len(cnt)))
    return numeric, categorical, lowcard


def crosstab(header, data_rows, a, b, missing_codes):
    if a not in header or b not in header:
        return None
    ia, ib = header.index(a), header.index(b)
    cells = Counter()
    row_tot = Counter()
    n_used = 0
    for r in data_rows:
        va = r[ia] if ia < len(r) else None
        vb = r[ib] if ib < len(r) else None
        ka = _cat_key(va, missing_codes)
        kb = _cat_key(vb, missing_codes)
        if ka is None or kb is None:
            continue
        cells[(ka, kb)] += 1
        row_tot[ka] += 1
        n_used += 1
    if n_used == 0:
        return None
    cols = sorted({kb for _, kb in cells})
    rows = sorted({ka for ka, _ in cells})
    return {"a": a, "b": b, "rows": rows, "cols": cols, "cells": cells,
            "row_tot": row_tot, "n": n_used}


# ------------------------------------------------------------------ 输出
def build_md(src, sheetname, n_rows, numeric, categorical, crosses, missing_codes,
             note_source=None, lowcard=()):
    L = []
    L.append(f"# 描述统计与交叉表：{os.path.basename(src)}")
    L.append("")
    L.append(f"- 数据来源文件：`{os.path.abspath(src)}`（工作表：{sheetname}）")
    L.append(f"- 有效数据行：{n_rows}（另有 {note_source or 0} 行全空被跳过）")
    L.append(f"- 被视为缺失的编码：{', '.join(_fmt_key(c) for c in sorted(missing_codes)) or '（无，仅空白视为缺失）'}")
    L.append("- **本表只含描述统计，不含任何推断结论；未做任何缺失填补或数据剔除。**")
    L.append("- 变量类型由脚本按规则判定：变量名像标识/分组（学号、编号、班级、性别…）→ 分类；"
             "变量名像成绩/分数且可解析 → 连续；≥95% 取值可解析为数值且唯一值 > 阈值 → 连续；"
             "其余 → 分类。**请以数据字典为准**，可用 `--num` / `--cat` 覆盖。")
    if lowcard:
        names = "、".join(f"`{n}`（{k} 类）" for n, k in lowcard)
        L.append(f"- ⚠ **待确认**：以下变量是数值型但唯一值较少，已默认按分类呈现：{names}。"
                 "若它们其实是连续测量值（如成绩、量表总分），请加 "
                 "`--num 变量名1,变量名2` 重新运行。")
    L.append("")
    L.append("## 一、连续变量描述统计")
    L.append("")
    L.append("| 变量 | N有效 | 缺失 | 缺失率 | 均值 | 标准差 | 最小 | P25 | 中位数 | P75 | 最大 |")
    L.append("|---|---|---|---|---|---|---|---|---|---|---|")
    for r in numeric:
        L.append("| {变量} | {N有效} | {缺失} | {缺失率} | {均值} | {标准差} | {最小} | {P25} | {中位数} | {P75} | {最大} |".format(**r))
    if not numeric:
        L.append("| （无连续变量被识别） | | | | | | | | | | |")
    L.append("")
    L.append("## 二、分类变量频数")
    L.append("")
    L.append("| 变量 | 类别数 | N有效 | 缺失 | 缺失率 | 最高频取值（频数, 占比） |")
    L.append("|---|---|---|---|---|---|")
    for rec, _cnt, nmiss in categorical:
        L.append(f"| {rec['变量']} | {rec['类别数']} | {rec['N有效']} | {rec['缺失']} | "
                 f"{rec['缺失率']} | {rec.get('最高频取值', '—')} |")
    if not categorical:
        L.append("| （无分类变量被识别） | | | | | |")
    L.append("")
    L.append("### 分类变量取值明细（前 8 类）")
    L.append("")
    for rec, cnt, _nmiss in categorical:
        L.append(f"**{rec['变量']}**（N有效 {rec['N有效']}）")
        L.append("")
        L.append("| 取值 | 频数 | 占比 |")
        L.append("|---|---|---|")
        for k, v in cnt.most_common(8):
            L.append(f"| {k} | {v} | {pct(v, rec['N有效'])}% |")
        L.append("")
    if crosses:
        L.append("## 三、交叉表")
        L.append("")
        for ct in crosses:
            L.append(f"### {ct['a']} × {ct['b']}（可用样本 N = {ct['n']}）")
            L.append("")
            L.append("| " + ct["a"] + " \\ " + ct["b"] + " | " + " | ".join(ct["cols"]) + " | 合计 |")
            L.append("|" + "---|" * (len(ct["cols"]) + 2))
            for r in ct["rows"]:
                cells = [str(ct["cells"].get((r, c), 0)) for c in ct["cols"]]
                L.append(f"| {r} | " + " | ".join(cells) + f" | {ct['row_tot'][r]} |")
            L.append("")
            L.append("*行内百分比（占该行合计）*")
            L.append("")
            L.append("| " + ct["a"] + " \\ " + ct["b"] + " | " + " | ".join(ct["cols"]) + " | 合计 |")
            L.append("|" + "---|" * (len(ct["cols"]) + 2))
            for r in ct["rows"]:
                tot = ct["row_tot"][r] or 1
                cells = [f"{pct(ct['cells'].get((r, c), 0), tot)}%" for c in ct["cols"]]
                L.append(f"| {r} | " + " | ".join(cells) + f" | {tot} |")
            L.append("")
        L.append("> 交叉表只呈现列联频数。**是否做卡方检验、以及任何因果解释，须由教师判断**；")
        L.append("> 组间差异的统计检验不在本脚本职责内。")
        L.append("")
    L.append("---")
    L.append("")
    L.append("## 生成信息（可复现性）")
    L.append("")
    L.append(f"- 脚本：`skills/research-assistant/scripts/desc_stats.py`")
    L.append(f"- 源文件：`{os.path.abspath(src)}` 工作表 `{sheetname}`")
    L.append(f"- 缺失编码：{list(sorted(missing_codes))}")
    L.append(f"- 数据行数：{n_rows}")
    return "\n".join(L)


def build_spec(numeric, categorical, crosses, meta_lines):
    sheets = []
    sheets.append({
        "name": "说明",
        "header": ["项目", "内容"],
        "rows": [[k, v] for k, v in meta_lines],
        "widths": [22, 90],
    })
    if numeric:
        cols = ["变量", "N有效", "缺失", "缺失率", "均值", "标准差", "最小", "P25", "中位数", "P75", "最大"]
        sheets.append({"name": "连续变量", "header": cols,
                       "rows": [[r.get(c, "") for c in cols] for r in numeric]})
    if categorical:
        cols = ["变量", "类别数", "N有效", "缺失", "缺失率", "最高频取值"]
        sheets.append({"name": "分类变量", "header": cols,
                       "rows": [[rec.get(c, "") for c in cols] for rec, _cnt, _nm in categorical]})
    for ct in crosses:
        header = [f"{ct['a']} \\ {ct['b']}"] + ct["cols"] + ["合计"]
        rows = []
        for r in ct["rows"]:
            tot = ct["row_tot"][r] or 1
            cells = [f"{ct['cells'].get((r, c), 0)} ({pct(ct['cells'].get((r, c), 0), tot)}%)"
                     for c in ct["cols"]]
            rows.append([r] + cells + [tot])
        sheets.append({"name": f"交叉_{ct['a']}x{ct['b']}"[:31], "header": header, "rows": rows})
    return {"sheets": sheets}


def _ensure_parent(path):
    d = os.path.dirname(os.path.abspath(path))
    if d and not os.path.isdir(d):
        os.makedirs(d, exist_ok=True)


def main():
    ap = argparse.ArgumentParser(prog="desc_stats",
                                 description="研究数据描述统计与交叉表（只读，不填补不剔除）")
    ap.add_argument("--in", dest="src", required=True, help="数据文件 .xlsx / .csv / .tsv")
    ap.add_argument("--sheet", help="xlsx 工作表名（缺省取第一个）")
    ap.add_argument("--out-md", help="Markdown 输出路径")
    ap.add_argument("--out-json", help="xlsx_kit build 规格 JSON 输出路径")
    ap.add_argument("--cross", action="append", default=[],
                    help="交叉表变量对，格式 A,B（可重复多次）")
    ap.add_argument("--missing-codes", default="-99,-98,-97",
                    help="视为缺失的数值编码，逗号分隔（默认 -99,-98,-97）")
    ap.add_argument("--cat-max-unique", type=int, default=12,
                    help="唯一值不超过该数的数值变量视为分类（默认 12）")
    ap.add_argument("--num", default="", help="强制按连续变量处理的变量名，逗号分隔")
    ap.add_argument("--cat", default="", help="强制按分类变量处理的变量名，逗号分隔")
    args = ap.parse_args()

    rows, sheetname = load_table(args.src, args.sheet)
    if not rows:
        sys.exit("输入为空")
    header = [str(h).strip() if h is not None and str(h).strip() else f"列{i+1}"
              for i, h in enumerate(rows[0])]
    # 去掉全空行
    body, skipped = [], 0
    for r in rows[1:]:
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            skipped += 1
            continue
        body.append(r)

    try:
        missing_codes = {float(x) for x in str(args.missing_codes).split(",") if str(x).strip()}
    except ValueError:
        sys.exit("--missing-codes 必须是逗号分隔的数字，如 -99,-98,-97")

    force_num = [s.strip() for s in args.num.split(",") if s.strip()]
    force_cat = [s.strip() for s in args.cat.split(",") if s.strip()]
    for n in force_num + force_cat:
        if n not in header:
            sys.exit(f"--num/--cat 指定的变量不存在于表头：{n}\n表头为：{header}")

    numeric, categorical, lowcard = profile(header, body, missing_codes,
                                            args.cat_max_unique, force_num, force_cat)
    crosses = []
    for pair in args.cross:
        if "," not in pair:
            sys.exit(f"--cross 参数格式应为 A,B，收到：{pair}")
        a, b = [s.strip() for s in pair.split(",", 1)]
        ct = crosstab(header, body, a, b, missing_codes)
        if ct is None:
            print(f"[警告] 交叉表跳过：变量名不存在或无有效配对（{pair}）；"
                  f"表头为：{header}", file=sys.stderr)
            continue
        crosses.append(ct)

    meta_lines = [
        ["数据来源", os.path.abspath(args.src)],
        ["工作表", sheetname],
        ["数据行数", len(body)],
        ["缺失编码", ", ".join(str(c) for c in sorted(missing_codes)) or "无"],
        ["变量数", len(header)],
        ["连续变量数", len(numeric)],
        ["分类变量数", len(categorical)],
        ["待确认（数值型但唯一值少）", "、".join(n for n, _ in lowcard) or "无"],
        ["声明", "本工作簿仅含描述统计；未做缺失填补、未剔除任何观测；不含推断结论"],
    ]

    if args.out_md:
        md = build_md(args.src, sheetname, len(body), numeric, categorical,
                      crosses, missing_codes, skipped, lowcard)
        _ensure_parent(args.out_md)
        with open(args.out_md, "w", encoding="utf-8") as f:
            f.write(md + "\n")
    if args.out_json:
        _ensure_parent(args.out_json)
        with open(args.out_json, "w", encoding="utf-8") as f:
            json.dump(build_spec(numeric, categorical, crosses, meta_lines),
                      f, ensure_ascii=False, indent=2)

    print(json.dumps({
        "rows": len(body), "variables": len(header),
        "numeric": len(numeric), "categorical": len(categorical),
        "crosstabs": [f"{c['a']}x{c['b']}" for c in crosses],
        "out_md": args.out_md, "out_json": args.out_json,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
